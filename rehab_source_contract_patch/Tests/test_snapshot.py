from datetime import time

import pytest
from openpyxl import Workbook

from rehab_excel import SourceContractError, build_snapshot


def _make_workbook(path, *, mismatch=False, missing_day=False):
    wb = Workbook()
    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    patients.append([1, "A02", "ΑΣΘΕΝΗΣ Α", "Ν", "παρών"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    planner.append([
        "PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση",
        "ΦΘ_Ώρα", "ΦΘ_Ημέρες", "ΦΘ_Θεραπευτής",
    ])
    planner.append([
        1, "A02", "ΑΛΛΟ ΟΝΟΜΑ" if mismatch else "ΑΣΘΕΝΗΣ Α", "Ν", "παρών",
        time(12, 15), None if missing_day else "Δε-Τε-Πα", "Θ1",
    ])

    settings = wb.create_sheet("SETTINGS")
    settings.append([
        "THERAPISTS_FTH", "HOURS_STANDARD", "HOURS_RECLINED", "DAY_COMBINATIONS",
        "THERAPIES", "PATIENT_STATUS", "YES_NO", "ROOMS",
    ])
    settings.append(["Θ1", "12:15", "8:00", "Δε-Τε-Πα", "ΦΘ", "παρών", "Ν", "A02"])
    wb.save(path)


def test_snapshot_reads_only_confirmed_sources(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path)
    snapshot = build_snapshot(path)
    assert snapshot.patient_count == 1
    assert snapshot.base_entry_count == 1
    assert snapshot.patients[0].patient_id == "1"
    assert snapshot.base_schedule[0].therapist_id == "Θ1"
    assert snapshot.authoritative_sheets == {"PATIENTS", "PATIENT_PLANNER", "SETTINGS"}


def test_snapshot_blocks_identity_contract_error(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, mismatch=True)
    with pytest.raises(SourceContractError, match="patients_planner_identity_mismatch"):
        build_snapshot(path)


def test_snapshot_allows_warning_but_does_not_guess_missing_days(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, missing_day=True)
    snapshot = build_snapshot(path)
    assert snapshot.base_entry_count == 0
    codes = {issue.code for issue in snapshot.audit.issues}
    assert "timed_entries_missing_day_pattern" in codes


def test_snapshot_can_be_built_for_diagnostics_even_with_contract_error(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, mismatch=True)
    snapshot = build_snapshot(path, require_clean_contract=False)
    assert snapshot.audit.ok is False
