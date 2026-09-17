from datetime import time

from openpyxl import Workbook

from rehab_excel import audit_workbook, read_base_schedule, read_patients, read_settings


def _make_workbook(path, *, stale_sessions=False, missing_day=False):
    wb = Workbook()
    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    patients.append([1, "A02", "ΑΣΘΕΝΗΣ Α", "Ν", "παρών"])
    patients.append([2, "B01", "ΑΣΘΕΝΗΣ Β", "Ο", "παρών"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    planner.append(
        [
            "PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση",
            "ΦΘ_Ώρα", "ΦΘ_Ημέρες", "ΦΘ_Θεραπευτής",
            "Ρομποτικό_Ώρα", "Ρομποτικό_Ημέρες", "Ρομποτικό_Θεραπευτής",
            "Ανακλινόμενο_Ώρα", "Ανακλινόμενο_Ημέρες",
            "Πισίνα_Ώρα", "Πισίνα_Ημέρες",
            "Εργο_Ώρα", "Εργο_Ημέρες",
            "Λογο_Ώρα", "Λογο_Ημέρες",
            "ΕΦΑ_Ώρα", "ΕΦΑ_Ημέρες", "Σχόλια 1", "Σχόλια 2",
        ]
    )
    planner.append([1, "A02", "ΑΣΘΕΝΗΣ Α", "Ν", "παρών", time(12, 15), "Καθ/να", "Θ1"])
    planner.append([
        2, "B01", "ΑΣΘΕΝΗΣ Β", "Ο", "παρών",
        time(10, 0), None if missing_day else "Τρ-Πε", "Θ2",
        time(9, 15), "Δε-Τε-Πα", "Θ2",
    ])

    settings = wb.create_sheet("SETTINGS")
    settings.append([
        "THERAPISTS_FTH", "HOURS_STANDARD", "HOURS_RECLINED", "DAY_COMBINATIONS",
        "THERAPIES", "PATIENT_STATUS", "YES_NO", "ROOMS",
    ])
    settings.append(["Θ1", "8:30", "8:00", "Καθ/να", "ΦΘ", "παρών", "Ν", "A02"])
    settings.append(["Θ2", "9:15", "8:30", "Τρ-Πε", "Ρομποτικό", "εξιτήριο", "Ο", "B01"])

    sessions = wb.create_sheet("SESSIONS")
    sessions.append(["PatientID", "Ασθενής", "Θεραπεία", "Ώρα", "Ημέρες", "Θεραπευτής", "StartDate", "EndDate"])
    sessions.append([1, "ΛΑΘΟΣ" if stale_sessions else "ΑΣΘΕΝΗΣ Α", "ΦΘ", "12:15", "Καθ/να", "Θ1"])

    wb.save(path)


def test_reads_patients_and_infectious_flag(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path)
    patients = read_patients(path)
    assert len(patients) == 2
    assert patients[0].infectious is True
    assert patients[1].infectious is False
    assert patients[0].room == "A02"


def test_reads_settings_timeslots_and_existing_status_text(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path)
    settings = read_settings(path)
    assert settings.therapist_names == ("Θ1", "Θ2")
    assert settings.standard_timeslots == (time(8, 30), time(9, 15))
    assert settings.patient_statuses == ("παρών", "εξιτήριο")


def test_reads_base_schedule_from_patient_planner_not_sessions(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, stale_sessions=True)
    entries = read_base_schedule(path)
    assert len(entries) == 3
    assert entries[0].patient_id == "1"
    assert entries[0].treatment == "ΦΘ"
    assert entries[2].treatment == "Ρομποτικό"
    assert entries[2].robotic is True


def test_audit_flags_sessions_identity_drift_without_using_it_as_source(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, stale_sessions=True)
    audit = audit_workbook(path)
    issue_codes = {issue.code for issue in audit.issues}
    assert "sessions_identity_drift" in issue_codes
    assert audit.patient_count == 2
    assert audit.base_entry_count == 3


def test_audit_does_not_guess_missing_day_pattern(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_workbook(path, missing_day=True)
    entries = read_base_schedule(path)
    assert len(entries) == 2
    audit = audit_workbook(path)
    issue = next(
        issue for issue in audit.issues if issue.code == "timed_entries_missing_day_pattern"
    )
    assert issue.count == 1
