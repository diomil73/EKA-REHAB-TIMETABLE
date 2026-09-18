from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from hashlib import sha256
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from .source_contract import get_source_rule


class WritebackSafetyError(RuntimeError):
    """Raised when a requested Excel write violates the source contract."""


class WriteIntent(str, Enum):
    VALUE = "value"
    CLEAR = "clear"
    PRESENTATION = "presentation"


@dataclass(frozen=True)
class CellPatch:
    """One future workbook mutation.

    This object only describes an intended change. The dry-run layer never
    applies it to the workbook.
    """

    sheet: str
    cell: str
    value: object | None = None
    intent: WriteIntent = WriteIntent.VALUE
    strike_through: bool = False
    font_role: str | None = None
    fill_role: str | None = None
    source_tag: str | None = None

    def normalized_key(self) -> tuple[str, str]:
        return self.sheet, self.cell.upper()


@dataclass(frozen=True)
class WritePlan:
    workbook_path: str
    patches: tuple[CellPatch, ...]

    @property
    def operation_count(self) -> int:
        return len(self.patches)

    @property
    def touched_sheets(self) -> tuple[str, ...]:
        return tuple(sorted({patch.sheet for patch in self.patches}))


@dataclass(frozen=True)
class DryRunReport:
    workbook_path: str
    operation_count: int
    touched_sheets: tuple[str, ...]
    sha256_before: str
    sha256_after: str
    workbook_unchanged: bool


def _file_sha256(path: str | Path) -> str:
    digest = sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _validate_cell_reference(cell: str) -> None:
    try:
        row, column = coordinate_to_tuple(cell.upper())
    except Exception as exc:  # openpyxl raises different errors by input shape
        raise WritebackSafetyError(f"Invalid Excel cell reference: {cell!r}") from exc
    if row < 1 or column < 1:
        raise WritebackSafetyError(f"Invalid Excel cell reference: {cell!r}")


def validate_write_plan(plan: WritePlan) -> None:
    """Validate a future write plan without modifying the workbook.

    Current safety rule: Python may only target sheets explicitly marked as
    ``future_write_target`` in the source contract. Authoritative input sheets
    and legacy operational sheets remain protected.
    """

    workbook_path = Path(plan.workbook_path)
    if not workbook_path.exists():
        raise WritebackSafetyError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        workbook_sheets = set(wb.sheetnames)
    finally:
        wb.close()

    seen: dict[tuple[str, str], CellPatch] = {}
    for patch in plan.patches:
        _validate_cell_reference(patch.cell)

        if patch.sheet not in workbook_sheets:
            raise WritebackSafetyError(
                f"Write target sheet does not exist in workbook: {patch.sheet}"
            )

        rule = get_source_rule(patch.sheet)
        if rule is None or not rule.future_write_target:
            raise WritebackSafetyError(
                f"Sheet {patch.sheet} is protected from Python write-back"
            )

        key = patch.normalized_key()
        previous = seen.get(key)
        if previous is not None and previous != patch:
            raise WritebackSafetyError(
                f"Conflicting writes planned for {patch.sheet}!{patch.cell.upper()}"
            )
        seen[key] = patch


def build_write_plan(
    workbook_path: str | Path,
    patches: Iterable[CellPatch],
) -> WritePlan:
    plan = WritePlan(str(workbook_path), tuple(patches))
    validate_write_plan(plan)
    return plan


def dry_run_writeback(plan: WritePlan) -> DryRunReport:
    """Validate a write plan and prove the workbook bytes remain unchanged.

    No workbook is opened in write mode and no ``save`` call occurs here.
    This is intentionally the only write-back mode implemented at this stage.
    """

    validate_write_plan(plan)
    before = _file_sha256(plan.workbook_path)

    # Deliberately no mutation. The plan is only inspected/validated.

    after = _file_sha256(plan.workbook_path)
    return DryRunReport(
        workbook_path=plan.workbook_path,
        operation_count=plan.operation_count,
        touched_sheets=plan.touched_sheets,
        sha256_before=before,
        sha256_after=after,
        workbook_unchanged=before == after,
    )
