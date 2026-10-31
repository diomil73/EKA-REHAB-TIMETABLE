from __future__ import annotations

from dataclasses import dataclass
from datetime import time
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from rehab_core.daily_state import DailySessionState, DailySessionStatus
from rehab_core.models import BaseScheduleEntry, Patient, PatientType

from .master_projection import MasterProjectionError, MasterProjectionRow, build_master_projection
from .writeback import CellPatch, WriteIntent, WritePlan, build_write_plan


class MasterDailyOverlayError(RuntimeError):
    """Raised when a daily MASTER overlay cannot be derived safely."""


_TREATMENT_COLUMNS = {
    "φθ": 4,
    "φυσικοθεραπεια": 4,
    "φυσικοθεραπεία": 4,
    "ρομποτικο": 5,
    "ρομποτικό": 5,
    "πισινα": 6,
    "πισίνα": 6,
    "ανακλινομενο": 7,
    "ανακλινόμενο": 7,
    "εργο": 8,
    "έργο": 8,
    "λογο": 9,
    "λόγο": 9,
    "εφα": 10,
}


def _column_letters(number: int) -> str:
    chars: list[str] = []
    while number:
        number, remainder = divmod(number - 1, 26)
        chars.append(chr(ord("A") + remainder))
    return "".join(reversed(chars))


def _base_entry_id_from_session(session_id: str) -> str:
    if "@" not in session_id:
        return session_id
    return session_id.rsplit("@", 1)[0]


def _treatment_column(treatment: str) -> int:
    key = " ".join(str(treatment).strip().casefold().split())
    column = _TREATMENT_COLUMNS.get(key)
    if column is None:
        raise MasterDailyOverlayError(
            f"MASTER daily overlay has no mapped treatment column for {treatment!r}"
        )
    return column


def _display_value(
    *,
    effective_time: time,
    day_pattern: str,
    effective_provider: str,
) -> str:
    parts = [effective_time.strftime("%H:%M")]
    days = str(day_pattern or "").strip()
    if days:
        parts.append(days)
    provider = str(effective_provider or "").strip()
    if provider:
        parts.append(provider)
    return "\n".join(parts)


def _fallback_status_label(state: DailySessionState) -> str:
    if state.status == DailySessionStatus.PATIENT_ABSENT:
        return "ΑΠΩΝ"
    if state.status == DailySessionStatus.CANCELLED:
        return "ΑΚΥΡΩΘΗΚΕ"
    if state.status == DailySessionStatus.THERAPIST_ABSENT:
        return "ΑΝΑΜΟΝΗ ΑΝΑΠΛΗΡΩΣΗΣ"
    raise MasterDailyOverlayError(
        f"Unsupported non-active MASTER status {state.status.value!r}"
    )


def _patient_status_value(states: Iterable[DailySessionState]) -> str:
    values: list[str] = []
    for state in states:
        reason = str(state.reason or "").strip()
        value = reason or _fallback_status_label(state)
        if value not in values:
            values.append(value)
    return " / ".join(values)


def _find_status_column(path: Path) -> int | None:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        ws = wb["MASTER_SCHEDULE"]
        for row in range(1, min(ws.max_row, 12) + 1):
            room = str(ws.cell(row, 2).value or "").strip().casefold()
            patient = str(ws.cell(row, 3).value or "").strip().casefold()
            if room != "θάλαμος".casefold() or patient != "ασθενής".casefold():
                continue
            for col in range(1, min(ws.max_column, 20) + 1):
                if str(ws.cell(row, col).value or "").strip().casefold() == "κατάσταση".casefold():
                    return col
            return None
        return None
    finally:
        wb.close()


@dataclass(frozen=True)
class MasterDailyBinding:
    session_id: str
    patient_id: str
    master_cell: str


@dataclass(frozen=True)
class MasterDailyOverlayPlan:
    write_plan: WritePlan
    bindings: tuple[MasterDailyBinding, ...]


def build_master_daily_overlay_plan(
    workbook_path: str | Path,
    *,
    base_entries: Iterable[BaseScheduleEntry],
    states: Iterable[DailySessionState],
    patients: Iterable[Patient],
    provider_labels: Mapping[str, str] | None = None,
    master_projection_rows: Iterable[MasterProjectionRow] | None = None,
) -> MasterDailyOverlayPlan:
    """Build date-specific MASTER patches from the authoritative daily state.

    MASTER_SCHEDULE stays a derived view. Replacements may replace the displayed
    provider/time for the selected date. Patient absence, cancellation and a
    pending therapist replacement never destroy the recurring display: the
    affected treatment cell is kept in place and receives a red strike-through.
    The patient-status column carries the dated reason when no active session for
    that inpatient remains on the selected date.
    """

    path = Path(workbook_path)
    entries = {_entry.base_entry_id: _entry for _entry in base_entries}
    state_list = tuple(states)
    patient_map = {patient.patient_id: patient for patient in patients}
    labels = dict(provider_labels or {})

    if master_projection_rows is None:
        try:
            projection = build_master_projection(path, patients=patient_map.values())
        except MasterProjectionError as exc:
            raise MasterDailyOverlayError(str(exc)) from exc
    else:
        projection = tuple(master_projection_rows)
    row_by_patient = {item.patient_id: item.target_row for item in projection}
    status_column = _find_status_column(path)

    patches: list[CellPatch] = []
    bindings: list[MasterDailyBinding] = []
    non_active_by_patient: dict[str, list[DailySessionState]] = {}
    all_by_patient: dict[str, list[DailySessionState]] = {}

    for state in state_list:
        all_by_patient.setdefault(state.patient_id, []).append(state)
        if state.status not in {DailySessionStatus.ACTIVE, DailySessionStatus.REPLACED}:
            non_active_by_patient.setdefault(state.patient_id, []).append(state)

        if state.status == DailySessionStatus.ACTIVE:
            continue
        patient = patient_map.get(state.patient_id)
        if patient is None:
            raise MasterDailyOverlayError(
                f"Missing patient {state.patient_id!r} for MASTER daily overlay"
            )
        if patient.patient_type != PatientType.INPATIENT:
            continue
        row = row_by_patient.get(state.patient_id)
        if row is None:
            raise MasterDailyOverlayError(
                f"Inpatient {state.patient_id!r} has no MASTER_SCHEDULE projection row"
            )

        base_id = _base_entry_id_from_session(state.session_id)
        entry = entries.get(base_id)
        if entry is None:
            raise MasterDailyOverlayError(
                f"Missing base entry {base_id!r} for MASTER daily overlay"
            )
        column = _treatment_column(entry.treatment)
        cell = f"{_column_letters(column)}{row}"

        if state.status == DailySessionStatus.REPLACED:
            if state.effective_time is None or state.effective_therapist_id is None:
                raise MasterDailyOverlayError(
                    f"Replacement state {state.session_id!r} lacks effective provider/time"
                )
            provider = labels.get(state.effective_therapist_id, state.effective_therapist_id)
            patches.append(
                CellPatch(
                    sheet="MASTER_SCHEDULE",
                    cell=cell,
                    value=_display_value(
                        effective_time=state.effective_time,
                        day_pattern=entry.day_pattern,
                        effective_provider=provider,
                    ),
                    wrap_text=True,
                    min_font_size=10,
                    source_tag="master_daily_effective_overlay",
                )
            )
        else:
            patches.append(
                CellPatch(
                    sheet="MASTER_SCHEDULE",
                    cell=cell,
                    intent=WriteIntent.PRESENTATION,
                    strike_through=True,
                    source_tag="master_daily_non_active_overlay",
                )
            )

        bindings.append(
            MasterDailyBinding(
                session_id=state.session_id,
                patient_id=state.patient_id,
                master_cell=cell,
            )
        )

    if status_column is not None:
        for patient_id, patient_states in all_by_patient.items():
            patient = patient_map.get(patient_id)
            row = row_by_patient.get(patient_id)
            if patient is None or patient.patient_type != PatientType.INPATIENT or row is None:
                continue
            if any(
                state.status in {DailySessionStatus.ACTIVE, DailySessionStatus.REPLACED}
                for state in patient_states
            ):
                continue
            inactive = non_active_by_patient.get(patient_id, [])
            if not inactive:
                continue
            patches.append(
                CellPatch(
                    sheet="MASTER_SCHEDULE",
                    cell=f"{_column_letters(status_column)}{row}",
                    value=_patient_status_value(inactive),
                    wrap_text=True,
                    min_font_size=10,
                    source_tag="master_daily_patient_status_overlay",
                )
            )

    return MasterDailyOverlayPlan(
        write_plan=build_write_plan(path, patches),
        bindings=tuple(bindings),
    )


def compose_master_daily_overlay(
    base_plan: WritePlan,
    overlay: MasterDailyOverlayPlan,
) -> WritePlan:
    """Merge MASTER daily patches into an existing daily preview plan."""

    if Path(base_plan.workbook_path).resolve() != Path(overlay.write_plan.workbook_path).resolve():
        raise MasterDailyOverlayError("Daily plans refer to different workbooks")

    patches = list(base_plan.patches)
    existing = {patch.normalized_key() for patch in patches}
    for patch in overlay.write_plan.patches:
        key = patch.normalized_key()
        if key in existing:
            raise MasterDailyOverlayError(
                f"Conflicting daily patch for {patch.sheet}!{patch.cell}"
            )
        existing.add(key)
        patches.append(patch)
    return build_write_plan(base_plan.workbook_path, patches)
