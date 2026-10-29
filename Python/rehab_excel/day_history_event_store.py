from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sys
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.day_history_events import DayHistoryEvent, DayHistoryEventType


SHEET_NAME = "DAY_HISTORY_EVENTS"
HEADERS = (
    "EventId",
    "ServiceDate",
    "EventType",
    "RecordedAt",
    "ActorOrSource",
    "TherapistId",
    "PatientId",
    "SessionId",
    "ReasonCode",
    "RelatedEventId",
    "RelatedSnapshotId",
    "Summary",
)


class DayHistoryEventStoreError(RuntimeError):
    """Raised when append-only day-history persistence is invalid or unsafe."""


def _parse_date(value: object, *, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value or "").strip())
    except ValueError as exc:
        raise DayHistoryEventStoreError(
            f"{SHEET_NAME} row {row}: invalid ServiceDate {value!r}"
        ) from exc


def _parse_datetime(value: object, *, row: int) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value or "").strip())
        except ValueError as exc:
            raise DayHistoryEventStoreError(
                f"{SHEET_NAME} row {row}: invalid RecordedAt {value!r}"
            ) from exc
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise DayHistoryEventStoreError(
            f"{SHEET_NAME} row {row}: RecordedAt must include a timezone offset"
        )
    return parsed


def _event_from_values(values: tuple[object, ...], *, row: int) -> DayHistoryEvent:
    event_id = str(values[0] or "").strip()
    actor = str(values[4] or "").strip()
    if not event_id:
        raise DayHistoryEventStoreError(f"{SHEET_NAME} row {row}: EventId is blank")
    if not actor:
        raise DayHistoryEventStoreError(f"{SHEET_NAME} row {row}: ActorOrSource is blank")
    try:
        event_type = DayHistoryEventType(str(values[2] or "").strip().upper())
    except ValueError as exc:
        raise DayHistoryEventStoreError(
            f"{SHEET_NAME} row {row}: unsupported EventType {values[2]!r}"
        ) from exc
    return DayHistoryEvent(
        event_id=event_id,
        service_date=_parse_date(values[1], row=row),
        event_type=event_type,
        recorded_at=_parse_datetime(values[3], row=row),
        actor_or_source=actor,
        therapist_id=str(values[5] or "").strip() or None,
        patient_id=str(values[6] or "").strip() or None,
        session_id=str(values[7] or "").strip() or None,
        reason_code=str(values[8] or "").strip() or None,
        related_event_id=str(values[9] or "").strip() or None,
        related_snapshot_id=str(values[10] or "").strip() or None,
        summary=str(values[11] or "").strip(),
    )


def read_day_history_events(
    workbook_path: str | Path,
    *,
    service_date: date | None = None,
    event_type: DayHistoryEventType | None = None,
) -> tuple[DayHistoryEvent, ...]:
    path = Path(workbook_path)
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if SHEET_NAME not in wb.sheetnames:
            return ()
        ws = wb[SHEET_NAME]
        actual_headers = tuple(
            str(ws.cell(1, col).value or "").strip()
            for col in range(1, len(HEADERS) + 1)
        )
        if actual_headers != HEADERS:
            raise DayHistoryEventStoreError(
                f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
            )
        result: list[DayHistoryEvent] = []
        seen_ids: set[str] = set()
        for row_number, row_values in enumerate(
            ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS), values_only=True),
            start=2,
        ):
            values = tuple(row_values)
            if not any(value is not None and str(value).strip() for value in values):
                continue
            event = _event_from_values(values, row=row_number)
            if event.event_id in seen_ids:
                raise DayHistoryEventStoreError(f"Duplicate EventId {event.event_id!r}")
            seen_ids.add(event.event_id)
            if service_date is not None and event.service_date != service_date:
                continue
            if event_type is not None and event.event_type != event_type:
                continue
            result.append(event)
        return tuple(result)
    finally:
        wb.close()


def _values_for_com(event: DayHistoryEvent) -> tuple[object, ...]:
    return (
        event.event_id,
        event.service_date.isoformat(),
        event.event_type.value,
        event.recorded_at.isoformat(),
        event.actor_or_source,
        event.therapist_id or "",
        event.patient_id or "",
        event.session_id or "",
        event.reason_code or "",
        event.related_event_id or "",
        event.related_snapshot_id or "",
        event.summary,
    )


def append_day_history_events(
    workbook_path: str | Path,
    events: Iterable[DayHistoryEvent],
) -> int:
    events = tuple(events)
    if not events:
        return 0
    ids = [event.event_id for event in events]
    if any(not value.strip() for value in ids):
        raise DayHistoryEventStoreError("EventId cannot be blank")
    if len(set(ids)) != len(ids):
        raise DayHistoryEventStoreError("Duplicate EventId in append batch")
    if any(event.recorded_at.tzinfo is None or event.recorded_at.utcoffset() is None for event in events):
        raise DayHistoryEventStoreError("RecordedAt must include a timezone offset")
    if sys.platform != "win32":
        raise DayHistoryEventStoreError(
            "Day-history event persistence requires Windows + Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DayHistoryEventStoreError("pywin32 is required") from exc

    path = Path(workbook_path).resolve()
    if not path.exists():
        raise DayHistoryEventStoreError(f"Workbook not found: {path}")

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)

        ws = None
        for sheet in workbook.Worksheets:
            if str(sheet.Name).casefold() == SHEET_NAME.casefold():
                ws = sheet
                break
        if ws is None:
            ws = workbook.Worksheets.Add(
                After=workbook.Worksheets(workbook.Worksheets.Count)
            )
            ws.Name = SHEET_NAME
            for col, header in enumerate(HEADERS, start=1):
                ws.Cells(1, col).Value = header
            ws.Rows(1).Font.Bold = True
        else:
            try:
                ws.Unprotect()
            except Exception:
                pass
            actual_headers = tuple(
                str(ws.Cells(1, col).Value or "").strip()
                for col in range(1, len(HEADERS) + 1)
            )
            if actual_headers != HEADERS:
                raise DayHistoryEventStoreError(
                    f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
                )

        last_row = max(int(ws.UsedRange.Rows.Count), 1)
        existing_ids: set[str] = set()
        for row_number in range(2, last_row + 1):
            values = tuple(
                ws.Cells(row_number, col).Value
                for col in range(1, len(HEADERS) + 1)
            )
            if not any(value is not None and str(value).strip() for value in values):
                continue
            existing_ids.add(_event_from_values(values, row=row_number).event_id)
        overlap = set(ids) & existing_ids
        if overlap:
            raise DayHistoryEventStoreError(
                f"Day-history EventId already exists: {sorted(overlap)!r}"
            )

        start_row = max(last_row + 1, 2)
        for offset, event in enumerate(events):
            target_row = start_row + offset
            for col, value in enumerate(_values_for_com(event), start=1):
                ws.Cells(target_row, col).Value2 = value
            readback = tuple(
                ws.Cells(target_row, col).Value
                for col in range(1, len(HEADERS) + 1)
            )
            persisted = _event_from_values(readback, row=target_row)
            if persisted != event:
                raise DayHistoryEventStoreError(
                    f"COM read-back mismatch for event {event.event_id!r}"
                )

        ws.Visible = 2
        ws.Protect()
        workbook.Save()
        return len(events)
    except DayHistoryEventStoreError:
        raise
    except Exception as exc:
        raise DayHistoryEventStoreError(
            f"Could not persist day-history events: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
