from __future__ import annotations

from pathlib import Path
import unicodedata

from openpyxl import load_workbook

from rehab_core.models import Patient, PatientType


class PatientRegistrySourceError(RuntimeError):
    """Raised when PATIENTS contains an unsafe/unknown patient classification."""


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _patient_id(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _is_yes(value: object) -> bool:
    text = _clean(value)
    if text is None:
        return False
    return text.upper() in {"Ν", "N", "YES", "Y", "TRUE", "1"}


def _header_map(ws) -> dict[str, int]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {
        str(value).strip(): index
        for index, value in enumerate(first)
        if value is not None
    }


def _first_header(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        if name in headers:
            return headers[name]
    return None


def _normalize_label(value: str) -> str:
    """Normalize user-facing labels for robust Greek/Unicode comparison.

    Greek casefolding converts final sigma (ς) to sigma (σ), while accented
    characters may be represented as composed or decomposed Unicode. Normalize
    to decomposed form and remove combining marks so labels such as
    ``Εξωτερικός``, ``ΕΞΩΤΕΡΙΚΟΣ`` and ``Εξωτερικος`` compare identically.
    """

    decomposed = unicodedata.normalize("NFD", value.casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _patient_type(value: object, *, row_number: int) -> PatientType:
    text = _clean(value)
    if text is None:
        return PatientType.INPATIENT

    normalized = _normalize_label(text)
    if normalized in {
        "inpatient",
        "internal",
        "εσωτερικοσ",
        "εσωτερικη",
    }:
        return PatientType.INPATIENT
    if normalized in {
        "outpatient",
        "external",
        "εξωτερικοσ",
        "εξωτερικη",
    }:
        return PatientType.OUTPATIENT

    raise PatientRegistrySourceError(
        f"PATIENTS row {row_number} has unknown patient type {text!r}"
    )


def read_patient_registry(path: str | Path) -> list[Patient]:
    """Read PATIENTS with optional patient-type/MRN extension columns.

    Legacy workbooks with the original five columns remain valid and every
    patient defaults to INPATIENT. Enhanced workbooks may add either Greek or
    English patient-type and hospital-MRN headers without changing the original
    five-column order.
    """

    workbook_path = Path(path)
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if "PATIENTS" not in wb.sheetnames:
            raise PatientRegistrySourceError("Workbook does not contain PATIENTS")

        ws = wb["PATIENTS"]
        headers = _header_map(ws)
        type_col = _first_header(
            headers,
            "PatientType",
            "ΤύποςΑσθενή",
            "Τύπος Ασθενή",
        )
        mrn_col = _first_header(
            headers,
            "HospitalMRN",
            "ΑΜ Νοσοκομείου",
            "ΑΜΝοσοκομείου",
        )

        patients: list[Patient] = []
        for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            patient_id = _patient_id(row[0] if len(row) > 0 else None)
            name = _clean(row[2] if len(row) > 2 else None)
            if patient_id is None and name is None:
                continue
            if patient_id is None or name is None:
                continue

            patient_type = _patient_type(
                row[type_col] if type_col is not None and type_col < len(row) else None,
                row_number=excel_row,
            )
            hospital_mrn = _clean(
                row[mrn_col] if mrn_col is not None and mrn_col < len(row) else None
            )

            patients.append(
                Patient(
                    patient_id=patient_id,
                    display_name=name,
                    room=_clean(row[1] if len(row) > 1 else None),
                    infectious=_is_yes(row[3] if len(row) > 3 else None),
                    status=_clean(row[4] if len(row) > 4 else None),
                    patient_type=patient_type,
                    hospital_mrn=hospital_mrn,
                )
            )

        return patients
    finally:
        wb.close()
