from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import shutil
import sys
from typing import Protocol

from openpyxl import load_workbook

from .student_registry import (
    STUDENT_REGISTRY_HEADERS,
    STUDENT_REGISTRY_SHEET,
    StudentRegistryError,
    read_students,
    validate_student_registry_headers,
)


class StudentRegistrySchemaError(RuntimeError):
    """Raised when the STUDENTS schema preview cannot be created safely."""


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _remove_preview(path: Path, *, required: bool) -> None:
    if not path.exists():
        return
    try:
        path.unlink()
    except PermissionError as exc:
        if required:
            raise StudentRegistrySchemaError(
                "Preview workbook is currently open or locked by Excel. "
                f"Close it and retry: {path}"
            ) from exc


def student_registry_schema_present(path: str | Path) -> bool:
    workbook_path = Path(path)
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=False,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if STUDENT_REGISTRY_SHEET not in wb.sheetnames:
            return False
        ws = wb[STUDENT_REGISTRY_SHEET]
        headers = next(
            ws.iter_rows(
                min_row=1,
                max_row=1,
                min_col=1,
                max_col=len(STUDENT_REGISTRY_HEADERS),
                values_only=True,
            )
        )
        validate_student_registry_headers(headers)
        return True
    finally:
        wb.close()


class StudentRegistrySchemaBackend(Protocol):
    def ensure_schema(self, workbook_path: Path) -> bool:
        """Ensure STUDENTS exists. Return True only when a new sheet was created."""
        ...


class Win32ComStudentRegistrySchemaBackend:
    """Create the authoritative STUDENTS source sheet on a preview copy only."""

    def ensure_schema(self, workbook_path: Path) -> bool:
        if sys.platform != "win32":
            raise StudentRegistrySchemaError(
                "Student registry schema preview requires Windows with Microsoft Excel"
            )
        try:
            import win32com.client  # type: ignore[import-not-found]
        except ImportError as exc:
            raise StudentRegistrySchemaError(
                "pywin32 is required for student registry schema preview"
            ) from exc

        excel = None
        workbook = None
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.ScreenUpdating = False
            excel.EnableEvents = False
            workbook = excel.Workbooks.Open(
                str(workbook_path.resolve()),
                UpdateLinks=0,
                ReadOnly=False,
            )

            existing = None
            try:
                existing = workbook.Worksheets(STUDENT_REGISTRY_SHEET)
            except Exception:
                existing = None

            if existing is not None:
                actual = tuple(
                    str(existing.Cells(1, index).Value2 or "").strip()
                    for index in range(1, len(STUDENT_REGISTRY_HEADERS) + 1)
                )
                try:
                    validate_student_registry_headers(actual)
                except StudentRegistryError as exc:
                    raise StudentRegistrySchemaError(str(exc)) from exc
                return False

            last_sheet = workbook.Worksheets(workbook.Worksheets.Count)
            ws = workbook.Worksheets.Add(After=last_sheet)
            ws.Name = STUDENT_REGISTRY_SHEET
            ws.Range("A1:H1").Value = (STUDENT_REGISTRY_HEADERS,)
            ws.Range("A1:H1").Font.Bold = True

            # Do not set NumberFormat on whole date columns during schema
            # creation. Localized Excel installations can reject a format token
            # at the COM boundary even on a brand-new sheet. The schema itself
            # is data structure, not presentation; date-cell formatting belongs
            # to the later student-row writeback where the exact cells are known.
            ws.Columns("A:H").AutoFit()
            workbook.Save()
            return True
        except StudentRegistrySchemaError:
            raise
        except Exception as exc:
            raise StudentRegistrySchemaError(
                f"Excel student registry schema creation failed: {exc}"
            ) from exc
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=False)
                except Exception:
                    pass
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass


@dataclass(frozen=True)
class StudentRegistrySchemaPreviewReport:
    source_path: str
    output_path: str
    sheet_created: bool
    source_unchanged: bool
    schema_verified: bool
    student_count: int


def create_student_registry_schema_preview(
    source_path: str | Path,
    output_path: str | Path,
    *,
    backend: StudentRegistrySchemaBackend | None = None,
    overwrite: bool = False,
) -> StudentRegistrySchemaPreviewReport:
    """Create a copy containing the authoritative STUDENTS registry schema.

    The immutable baseline is never opened for writing. Existing workbooks that
    do not yet contain STUDENTS remain readable; this preview is the explicit
    schema migration step before student registration writeback is enabled.
    """

    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if not source.exists():
        raise StudentRegistrySchemaError(f"Source workbook not found: {source}")
    if source == output:
        raise StudentRegistrySchemaError("Output must be different from source workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise StudentRegistrySchemaError("Source and output must both be .xlsm files")
    if output.exists() and not overwrite:
        raise StudentRegistrySchemaError(f"Output already exists: {output}")

    source_before = _sha256(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        _remove_preview(output, required=True)
    shutil.copy2(source, output)

    selected_backend = backend or Win32ComStudentRegistrySchemaBackend()
    try:
        created = selected_backend.ensure_schema(output)
    except Exception:
        _remove_preview(output, required=False)
        raise

    source_after = _sha256(source)
    if source_before != source_after:
        _remove_preview(output, required=False)
        raise StudentRegistrySchemaError(
            "Source workbook changed during student registry schema preview"
        )

    try:
        schema_verified = student_registry_schema_present(output)
        students = read_students(output)
    except (StudentRegistryError, ValueError) as exc:
        _remove_preview(output, required=False)
        raise StudentRegistrySchemaError(
            f"Student registry schema verification failed: {exc}"
        ) from exc

    if not schema_verified:
        _remove_preview(output, required=False)
        raise StudentRegistrySchemaError(
            "Student registry schema verification failed: STUDENTS sheet missing"
        )

    return StudentRegistrySchemaPreviewReport(
        source_path=str(source),
        output_path=str(output),
        sheet_created=created,
        source_unchanged=True,
        schema_verified=True,
        student_count=len(students),
    )
