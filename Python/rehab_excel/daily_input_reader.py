from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import unicodedata
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.models import (
    AbsenceKind,
    DailyAbsence,
    DailySessionCancellation,
    Patient,
    Session,
    SessionCancellationKind,
)


SHEET_NAME = "DAILY_INPUT"


class DailyInputReadError(RuntimeError):
    """Raised when DAILY_INPUT contains ambiguous or invalid operational data."""


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().upper())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid DAILY_INPUT date: {value!r}")


def _excel_fraction_to_time(value: float) -> time:
    if value < 0:
        raise DailyInputReadError(f"Invalid Excel time value: {value!r}")
    seconds = int(round((value % 1) * 24 * 60 * 60)) % (24 * 60 * 60)
    return time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)


def _parse_time(value: object, *, field: str) -> time | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        return _excel_fraction_to_time(float(value))
    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid {field} time: {value!r}; expected HH:MM")


def _is_yes(value: object) -> bool:
    return _norm(value) in {"ΝΑΙ", "NAI", "YES", "Y", "1", "TRUE"}


def _is_present_status(value: object) -> bool:
    return _norm(value) in {"ΠΑΡΩΝ", "PARON", "PRESENT"}


def _cancellation_kind(value: object) -> SessionCancellationKind | None:
    normalized = _norm(value)
    if normalized in {"ΑΝΑΒΟΛΗ ΤΜΗΜΑΤΟΣ", "DEPARTMENT POSTPONED"}:
        return SessionCancellationKind.DEPARTMENT_POSTPONED
    if normalized in {"ΔΕΝ ΠΡΟΣΗΛΘΕ", "PATIENT NO SHOW", "NO SHOW"}:
        return SessionCancellationKind.PATIENT_NO_SHOW
    if normalized in {"ΑΝΑΣΤΟΛΗ ΘΕΡΑΠΕΙΑΣ", "TREATMENT SUSPENDED"}:
        return SessionCancellationKind.TREATMENT_SUSPENDED
    return None


def _one_minute_after(value: time) -> time:
    stamp = datetime.combine(date(2000, 1, 1), value) + timedelta(minutes=1)
    return stamp.time()


def _find_header_row(ws, expected: tuple[str, ...]) -> int:
    wanted = tuple(_norm(value) for value in expected)
    for row in range(1, min(ws.max_row, 100) + 1):
        current = tuple(_norm(ws.cell(row, col).value) for col in range(1, len(expected) + 1))
        if current == wanted:
            return row
    raise DailyInputReadError(f"Could not locate DAILY_INPUT header: {' | '.join(expected)}")


def _patient_display_name(value: str) -> str:
    text = str(value or "").strip()
    if text.upper().startswith("ΕΞ |"):
        return text.split("|", 1)[1].strip()
    return text


def _unique_patient(patients: Iterable[Patient], display_name: str) -> Patient:
    wanted = _norm(_patient_display_name(display_name))
    matches = [patient for patient in patients if _norm(patient.display_name) == wanted]
    if len(matches) != 1:
        raise DailyInputReadError(
            f"Expected one patient named {display_name!r}; found {len(matches)}"
        )
    return matches[0]


def _provider_id(
    sessions: Iterable[Session], display_name: str, therapist_names: Iterable[str] = ()
) -> str:
    wanted = _norm(display_name)
    providers = {
        str(session.therapist_id).strip()
        for session in sessions
        if str(session.therapist_id).strip()
    }
    providers.update(str(name).strip() for name in therapist_names if str(name).strip())
    matches = sorted(provider for provider in providers if _norm(provider) == wanted)
    if len(matches) != 1:
        raise DailyInputReadError(
            f"Expected one therapist named {display_name!r} in the official registry/date; found {len(matches)}"
        )
    return matches[0]


def _unique_session_for_patient_at(
    sessions: Iterable[Session], *, patient_id: str, target_date: date, target_time: time
) -> Session:
    matches = [
        session
        for session in sessions
        if session.patient_id == patient_id
        and session.session_date == target_date
        and session.start_time == target_time
    ]
    if len(matches) != 1:
        raise DailyInputReadError(
            f"Expected exactly one session for patient {patient_id!r} at {target_time.strftime('%H:%M')} on {target_date.isoformat()}; found {len(matches)}"
        )
    return matches[0]


@dataclass(frozen=True)
class DailyInputSourceChange:
    change_id: str
    source_row: int
    change_type: str
    subject_id: str
    display_name: str
    reason: str
    all_day: bool
    start_time: time | None = None


@dataclass(frozen=True)
class DailyInputReadResult:
    target_date: date
    absences: tuple[DailyAbsence, ...]
    therapist_rows_used: int
    patient_rows_used: int
    warnings: tuple[str, ...] = ()
    cancellations: tuple[DailySessionCancellation, ...] = ()
    source_changes: tuple[DailyInputSourceChange, ...] = ()


def read_daily_input(
    workbook_path: str | Path,
    *,
    patients: Iterable[Patient],
    sessions: Iterable[Session],
    therapist_names: Iterable[str] = (),
    ignore_rows: Iterable[int] = (),
    allow_empty: bool = False,
) -> DailyInputReadResult:
    """Read DAILY_INPUT into reversible daily facts without modifying workbook."""

    workbook_path = Path(workbook_path)
    patient_list = list(patients)
    session_list = list(sessions)
    therapist_name_list = tuple(therapist_names)
    ignored = {int(row) for row in ignore_rows}
    wb = load_workbook(workbook_path, read_only=False, data_only=False, keep_vba=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise DailyInputReadError(f"Workbook has no {SHEET_NAME} sheet")
        ws = wb[SHEET_NAME]
        target_date = _parse_date(ws["B2"].value)
        therapist_header = _find_header_row(
            ws, ("Θεραπευτής", "Όλη ημέρα", "Από", "Έως", "Αιτία", "Σχόλιο")
        )
        patient_header = _find_header_row(
            ws, ("Ασθενής", "Όλη ημέρα", "Ώρα", "Κατάσταση", "Σχόλιο")
        )
        if patient_header <= therapist_header:
            raise DailyInputReadError("Patient section appears before therapist section")

        absences: list[DailyAbsence] = []
        cancellations: list[DailySessionCancellation] = []
        changes: list[DailyInputSourceChange] = []
        warnings: list[str] = []
        therapist_rows_used = 0
        patient_rows_used = 0

        for row in range(therapist_header + 1, patient_header - 1):
            if row in ignored:
                continue
            therapist_name = str(ws.cell(row, 1).value or "").strip()
            if not therapist_name:
                continue
            therapist_rows_used += 1
            all_day = _is_yes(ws.cell(row, 2).value)
            reason = str(ws.cell(row, 5).value or "").strip() or None
            comment = str(ws.cell(row, 6).value or "").strip()
            if comment:
                reason = f"{reason} | {comment}" if reason else comment
            provider_id = _provider_id(session_list, therapist_name, therapist_name_list)
            if all_day:
                start = end = None
            else:
                start = _parse_time(ws.cell(row, 3).value, field="therapist start")
                end = _parse_time(ws.cell(row, 4).value, field="therapist end")
                if start is None:
                    raise DailyInputReadError(f"Therapist row {row}: choose Όλη ημέρα=ΝΑΙ or provide Από")
                if end is None or end == start:
                    end = _one_minute_after(start)
                elif end < start:
                    raise DailyInputReadError(f"Therapist row {row}: Έως cannot be earlier than Από")
            absences.append(
                DailyAbsence(
                    absence_kind=AbsenceKind.THERAPIST,
                    subject_id=provider_id,
                    absence_date=target_date,
                    start_time=start,
                    end_time=end,
                    reason=reason,
                )
            )
            changes.append(
                DailyInputSourceChange(
                    change_id=f"daily-input:{target_date.isoformat()}:therapist:{row}",
                    source_row=row,
                    change_type="therapist_absence",
                    subject_id=provider_id,
                    display_name=therapist_name,
                    reason=reason or "Απουσία θεραπευτή",
                    all_day=all_day,
                    start_time=start,
                )
            )

        for row in range(patient_header + 1, ws.max_row + 1):
            if row in ignored:
                continue
            patient_name = str(ws.cell(row, 1).value or "").strip()
            if not patient_name:
                continue
            patient_rows_used += 1
            all_day = _is_yes(ws.cell(row, 2).value)
            status = str(ws.cell(row, 4).value or "").strip()
            comment = str(ws.cell(row, 5).value or "").strip()
            if not status:
                raise DailyInputReadError(f"Patient row {row}: Κατάσταση is required")
            if _is_present_status(status):
                warnings.append(f"Patient row {row} ({patient_name}) is ΠΑΡΩΝ and was ignored")
                continue
            patient = _unique_patient(patient_list, patient_name)
            cancellation_kind = _cancellation_kind(status)
            if cancellation_kind is not None:
                if all_day:
                    raise DailyInputReadError(f"Patient row {row}: {status} must target one session, not Όλη ημέρα")
                start = _parse_time(ws.cell(row, 3).value, field="patient cancellation")
                if start is None:
                    raise DailyInputReadError(f"Patient row {row}: {status} requires Ώρα")
                session = _unique_session_for_patient_at(
                    session_list, patient_id=patient.patient_id, target_date=target_date, target_time=start
                )
                reason = status if not comment else f"{status} | {comment}"
                cancellations.append(
                    DailySessionCancellation(
                        cancellation_id=f"daily-input:{row}",
                        target_session_id=session.session_id,
                        cancellation_date=target_date,
                        kind=cancellation_kind,
                        reason=reason,
                    )
                )
                changes.append(
                    DailyInputSourceChange(
                        change_id=f"daily-input:{target_date.isoformat()}:cancellation:{row}",
                        source_row=row,
                        change_type="session_cancellation",
                        subject_id=patient.patient_id,
                        display_name=patient.display_name,
                        reason=reason,
                        all_day=False,
                        start_time=start,
                    )
                )
                continue

            if all_day:
                start = end = None
            else:
                start = _parse_time(ws.cell(row, 3).value, field="patient")
                if start is None:
                    raise DailyInputReadError(f"Patient row {row}: choose Όλη ημέρα=ΝΑΙ or provide Ώρα")
                end = _one_minute_after(start)
            reason = status if not comment else f"{status} | {comment}"
            absences.append(
                DailyAbsence(
                    absence_kind=AbsenceKind.PATIENT,
                    subject_id=patient.patient_id,
                    absence_date=target_date,
                    start_time=start,
                    end_time=end,
                    reason=reason,
                )
            )
            changes.append(
                DailyInputSourceChange(
                    change_id=f"daily-input:{target_date.isoformat()}:patient:{row}",
                    source_row=row,
                    change_type="patient_absence",
                    subject_id=patient.patient_id,
                    display_name=patient.display_name,
                    reason=reason,
                    all_day=all_day,
                    start_time=start,
                )
            )

        if not absences and not cancellations and not allow_empty:
            raise DailyInputReadError("DAILY_INPUT contains no active absence/cancellation rows")

        return DailyInputReadResult(
            target_date=target_date,
            absences=tuple(absences),
            therapist_rows_used=therapist_rows_used,
            patient_rows_used=patient_rows_used,
            warnings=tuple(warnings),
            cancellations=tuple(cancellations),
            source_changes=tuple(changes),
        )
    finally:
        wb.close()
