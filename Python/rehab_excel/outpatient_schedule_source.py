from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.day_patterns import parse_day_pattern
from rehab_core.models import BaseScheduleEntry, Patient

from .outpatient_presentation import validate_outpatient_base_entries
from .reader import IMPLICIT_DAILY_PATTERN, read_base_schedule, read_patients


OUTPATIENT_SCHEDULE_SHEET = "OUTPATIENT_SCHEDULE"
OUTPATIENT_REQUIRED_HEADERS = (
    "PatientID",
    "Ασθενής",
    "Θεραπεία",
    "Ώρα",
    "Ημέρες",
    "Θεραπευτής",
)


class OutpatientScheduleSourceError(RuntimeError):
    """Raised when the outpatient recurring source is inconsistent or unsafe."""


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _patient_id(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _time_value(value: object):
    # Reuse the production reader conversion without making it public API.
    from .reader import _as_time

    return _as_time(value)


def _header_map(ws) -> dict[str, int]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {
        str(value).strip(): index
        for index, value in enumerate(first)
        if value is not None
    }


def read_outpatient_schedule(
    path: str | Path,
    *,
    patients: Iterable[Patient] | None = None,
) -> list[BaseScheduleEntry]:
    """Read the optional authoritative outpatient recurring schedule.

    The sheet uses one row per recurring treatment. It is intentionally simpler
    than PATIENT_PLANNER because it is a storage source, not an inpatient-facing
    planning view. Absence of the sheet is backward-compatible and returns [].
    """

    workbook_path = Path(path)
    patient_list = list(patients) if patients is not None else read_patients(path)
    patient_by_id = {patient.patient_id: patient for patient in patient_list}

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if OUTPATIENT_SCHEDULE_SHEET not in wb.sheetnames:
            return []

        ws = wb[OUTPATIENT_SCHEDULE_SHEET]
        headers = _header_map(ws)
        missing = [name for name in OUTPATIENT_REQUIRED_HEADERS if name not in headers]
        if missing:
            raise OutpatientScheduleSourceError(
                "OUTPATIENT_SCHEDULE missing required columns: " + ", ".join(missing)
            )

        entries: list[BaseScheduleEntry] = []
        for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            patient_id = _patient_id(row[headers["PatientID"]])
            name = _clean(row[headers["Ασθενής"]])
            treatment = _clean(row[headers["Θεραπεία"]])
            start_time = _time_value(row[headers["Ώρα"]])
            day_pattern = _clean(row[headers["Ημέρες"]]) or IMPLICIT_DAILY_PATTERN
            therapist_id = _clean(row[headers["Θεραπευτής"]])

            if not any((patient_id, name, treatment, start_time, therapist_id)):
                continue
            if patient_id is None or name is None or treatment is None or start_time is None:
                raise OutpatientScheduleSourceError(
                    f"OUTPATIENT_SCHEDULE row {excel_row} is incomplete"
                )

            patient = patient_by_id.get(patient_id)
            if patient is None:
                raise OutpatientScheduleSourceError(
                    f"OUTPATIENT_SCHEDULE row {excel_row} references unknown PatientID {patient_id!r}"
                )
            if not patient.is_outpatient:
                raise OutpatientScheduleSourceError(
                    f"OUTPATIENT_SCHEDULE row {excel_row} references inpatient {patient_id!r}"
                )
            if patient.display_name.casefold() != name.casefold():
                raise OutpatientScheduleSourceError(
                    f"OUTPATIENT_SCHEDULE row {excel_row} name does not match PATIENTS for {patient_id!r}"
                )

            try:
                parse_day_pattern(day_pattern)
            except ValueError as exc:
                raise OutpatientScheduleSourceError(
                    f"OUTPATIENT_SCHEDULE row {excel_row} has invalid day pattern {day_pattern!r}"
                ) from exc

            entries.append(
                BaseScheduleEntry(
                    base_entry_id=f"outpatient:{excel_row}",
                    patient_id=patient_id,
                    treatment=treatment,
                    start_time=start_time,
                    day_pattern=day_pattern,
                    therapist_id=therapist_id,
                    robotic=False,
                )
            )

        validate_outpatient_base_entries(entries, patient_list)
        return entries
    finally:
        wb.close()


def read_unified_base_schedule(path: str | Path) -> list[BaseScheduleEntry]:
    """Merge inpatient PATIENT_PLANNER and outpatient recurring sources.

    Patient type controls source/view routing only. Downstream scheduling sees
    one BaseScheduleEntry stream and therefore preserves identical workload,
    capacity, replacement, and daily-state semantics for both patient types.
    """

    patients = read_patients(path)
    inpatient_entries = read_base_schedule(path)
    outpatient_entries = read_outpatient_schedule(path, patients=patients)
    return [*inpatient_entries, *outpatient_entries]
