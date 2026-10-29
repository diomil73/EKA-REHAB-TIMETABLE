from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sys
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.therapist_day_history import TherapistDayHistoryRow


SHEET_NAME = "THERAPIST_DAY_HISTORY"
HEADERS = (
    "SnapshotId",
    "ServiceDate",
    "TherapistId",
    "CapacityTimeslots",
    "PlannedOccupiedTimeslots",
    "PlannedSessions",
    "FinalOccupiedTimeslots",
    "FinalActiveSessions",
    "ReplacementsReceived",
    "ReleasedByCancellationOrNoShow",
    "TreatmentSuspensionsAffectingLoad",
    "ApprovedSharedOrOverCapacityEvents",
)


class TherapistDayHistoryStoreError(RuntimeError):
    """Raised when therapist/day history persistence is invalid or unsafe."""


def _required_int(value: object, *, row: int, field: str) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise TherapistDayHistoryStoreError(
            f"{SHEET_NAME} row {row}: invalid {field} {value!r}"
        ) from exc
    if result < 0:
        raise TherapistDayHistoryStoreError(
            f"{SHEET_NAME} row {row}: {field} cannot be negative"
        )
    return result


def _optional_int(value: object, *, row: int, field: str) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    return _required_int(value, row=row, field=field)


def _parse_date(value: object, *, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value or "").strip())
    except ValueError as exc:
        raise TherapistDayHistoryStoreError(
            f"{SHEET_NAME} row {row}: invalid ServiceDate {value!r}"
        ) from exc


def _row_from_values(values: tuple[object, ...], *, row: int) -> TherapistDayHistoryRow:
    snapshot_id = str(values[0] or "").strip()
    therapist_id = str(values[2] or "").strip()
    if not snapshot_id:
        raise TherapistDayHistoryStoreError(f"{SHEET_NAME} row {row}: SnapshotId is blank")
    if not therapist_id:
        raise TherapistDayHistoryStoreError(f"{SHEET_NAME} row {row}: TherapistId is blank")
    return TherapistDayHistoryRow(
        snapshot_id=snapshot_id,
        service_date=_parse_date(values[1], row=row),
        therapist_id=therapist_id,
        capacity_timeslots=_optional_int(values[3], row=row, field="CapacityTimeslots"),
        planned_occupied_timeslots=_required_int(values[4], row=row, field="PlannedOccupiedTimeslots"),
        planned_sessions=_required_int(values[5], row=row, field="PlannedSessions"),
        final_occupied_timeslots=_required_int(values[6], row=row, field="FinalOccupiedTimeslots"),
        final_active_sessions=_required_int(values[7], row=row, field="FinalActiveSessions"),
        replacements_received=_required_int(values[8], row=row, field="ReplacementsReceived"),
        released_by_cancellation_or_no_show=_required_int(
            values[9], row=row, field="ReleasedByCancellationOrNoShow"
        ),
        treatment_suspensions_affecting_load=_required_int(
            values[10], row=row, field="TreatmentSuspensionsAffectingLoad"
        ),
        approved_shared_or_over_capacity_events=_required_int(
            values[11], row=row, field="ApprovedSharedOrOverCapacityEvents"
        ),
    )


def read_therapist_day_history_rows(
    workbook_path: str | Path,
    *,
    snapshot_id: str | None = None,
    service_date: date | None = None,
) -> tuple[TherapistDayHistoryRow, ...]:
    path = Path(workbook_path)
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if SHEET_NAME not in wb.sheetnames:
            return ()
        ws = wb[SHEET_NAME]
        actual_headers = tuple(
            str(ws.cell(1, col).value or "").strip()
            for col in range(1, len(HEADERS) + 1)
        )
        if actual_headers != HEADERS:
            raise TherapistDayHistoryStoreError(
                f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
            )
        result: list[TherapistDayHistoryRow] = []
        seen: set[tuple[str, str]] = set()
        for row_number, row_values in enumerate(
            ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS), values_only=True),
            start=2,
        ):
            values = tuple(row_values)
            if not any(value is not None and str(value).strip() for value in values):
                continue
            parsed = _row_from_values(values, row=row_number)
            key = (parsed.snapshot_id, parsed.therapist_id)
            if key in seen:
                raise TherapistDayHistoryStoreError(
                    f"Duplicate therapist/day history key {key!r}"
                )
            seen.add(key)
            if snapshot_id is not None and parsed.snapshot_id != snapshot_id:
                continue
            if service_date is not None and parsed.service_date != service_date:
                continue
            result.append(parsed)
        return tuple(result)
    finally:
        wb.close()


def _values_for_com(row: TherapistDayHistoryRow) -> tuple[object, ...]:
    return (
        row.snapshot_id,
        row.service_date.isoformat(),
        row.therapist_id,
        "" if row.capacity_timeslots is None else row.capacity_timeslots,
        row.planned_occupied_timeslots,
        row.planned_sessions,
        row.final_occupied_timeslots,
        row.final_active_sessions,
        row.replacements_received,
        row.released_by_cancellation_or_no_show,
        row.treatment_suspensions_affecting_load,
        row.approved_shared_or_over_capacity_events,
    )


def append_therapist_day_history_rows(
    workbook_path: str | Path,
    rows: Iterable[TherapistDayHistoryRow],
) -> int:
    rows = tuple(rows)
    if not rows:
        return 0
    snapshot_ids = {row.snapshot_id for row in rows}
    service_dates = {row.service_date for row in rows}
    keys = {(row.snapshot_id, row.therapist_id) for row in rows}
    if len(snapshot_ids) != 1 or len(service_dates) != 1:
        raise TherapistDayHistoryStoreError(
            "One append batch must belong to exactly one snapshot/date"
        )
    if len(keys) != len(rows):
        raise TherapistDayHistoryStoreError("Duplicate therapist rows in append batch")
    if sys.platform != "win32":
        raise TherapistDayHistoryStoreError(
            "Therapist/day history persistence requires Windows + Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise TherapistDayHistoryStoreError("pywin32 is required") from exc

    path = Path(workbook_path).resolve()
    if not path.exists():
        raise TherapistDayHistoryStoreError(f"Workbook not found: {path}")

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)

        ws = None
        for sheet in workbook.Worksheets:
            if str(sheet.Name).casefold() == SHEET_NAME.casefold():
                ws = sheet
                break
        if ws is None:
            ws = workbook.Worksheets.Add(
                After=workbook.Worksheets(workbook.Worksheets.Count)
            )
            ws.Name = SHEET_NAME
            for col, header in enumerate(HEADERS, start=1):
                ws.Cells(1, col).Value = header
            ws.Rows(1).Font.Bold = True
        else:
            try:
                ws.Unprotect()
            except Exception:
                pass
            actual_headers = tuple(
                str(ws.Cells(1, col).Value or "").strip()
                for col in range(1, len(HEADERS) + 1)
            )
            if actual_headers != HEADERS:
                raise TherapistDayHistoryStoreError(
                    f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
                )

        last_row = max(int(ws.UsedRange.Rows.Count), 1)
        existing_keys: set[tuple[str, str]] = set()
        for row_number in range(2, last_row + 1):
            values = tuple(
                ws.Cells(row_number, col).Value
                for col in range(1, len(HEADERS) + 1)
            )
            if not any(value is not None and str(value).strip() for value in values):
                continue
            existing = _row_from_values(values, row=row_number)
            existing_keys.add((existing.snapshot_id, existing.therapist_id))

        overlap = keys & existing_keys
        if overlap:
            raise TherapistDayHistoryStoreError(
                f"Therapist/day history already exists for: {sorted(overlap)!r}"
            )

        start_row = max(last_row + 1, 2)
        for offset, row in enumerate(rows):
            target_row = start_row + offset
            for col, value in enumerate(_values_for_com(row), start=1):
                ws.Cells(target_row, col).Value2 = value
            readback = tuple(
                ws.Cells(target_row, col).Value
                for col in range(1, len(HEADERS) + 1)
            )
            persisted = _row_from_values(readback, row=target_row)
            if persisted != row:
                raise TherapistDayHistoryStoreError(
                    f"COM read-back mismatch for therapist {row.therapist_id!r}"
                )

        ws.Visible = 2
        ws.Protect()
        workbook.Save()
        return len(rows)
    except TherapistDayHistoryStoreError:
        raise
    except Exception as exc:
        raise TherapistDayHistoryStoreError(
            f"Could not persist therapist/day history: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
