from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from typing import Iterable
import warnings

from openpyxl import load_workbook

from rehab_core.day_patterns import parse_day_pattern
from rehab_core.models import BaseScheduleEntry, Patient

# openpyxl warns about x14 data validation while streaming .xlsm sheets. This
# adapter never saves the workbook, so those structures are not modified.
warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
)


IMPLICIT_DAILY_PATTERN = "Καθ/να"

REQUIRED_SHEETS = {
    "PATIENTS",
    "PATIENT_PLANNER",
    "SETTINGS",
}


@dataclass(frozen=True)
class WorkbookSettings:
    therapist_names: tuple[str, ...]
    standard_timeslots: tuple[time, ...]
    reclined_timeslots: tuple[time, ...]
    day_patterns: tuple[str, ...]
    therapies: tuple[str, ...]
    patient_statuses: tuple[str, ...]
    yes_no_values: tuple[str, ...]
    rooms: tuple[str, ...]


@dataclass(frozen=True)
class AuditIssue:
    code: str
    message: str
    count: int = 1
    severity: str = "warning"


@dataclass(frozen=True)
class WorkbookAudit:
    workbook_path: str
    patient_count: int
    base_entry_count: int
    provider_entry_count: int
    issues: tuple[AuditIssue, ...]

    @property
    def ok(self) -> bool:
        return not any(issue.severity == "error" for issue in self.issues)


def _clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _patient_id(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_time(value: object) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        fraction = float(value) % 1
        minutes = round(fraction * 24 * 60)
        return time((minutes // 60) % 24, minutes % 60)
    text = str(value).strip().replace(".", ":")
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            pass
    raise ValueError(f"Unsupported time value: {value!r}")


def _is_yes(value: object) -> bool:
    text = _clean_text(value)
    if text is None:
        return False
    return text.upper() in {"Ν", "N", "YES", "Y", "TRUE", "1"}


def _load(path: str | Path):
    """Open workbook read-only.

    data_only=True is intentional: PATIENT_PLANNER mirrors PATIENTS with
    formulas and the adapter needs the cached values Excel last calculated.
    The workbook is never saved by this module.
    """

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
        )
        workbook_path = Path(path)
        return load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=True,
            keep_vba=workbook_path.suffix.casefold() == ".xlsm",
        )


def read_patients(path: str | Path) -> list[Patient]:
    wb = _load(path)
    try:
        if "PATIENTS" not in wb.sheetnames:
            raise ValueError("Workbook does not contain PATIENTS")
        ws = wb["PATIENTS"]
        patients: list[Patient] = []
        for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            patient_id = _patient_id(row[0])
            name = _clean_text(row[2])
            if patient_id is None and name is None:
                continue
            if patient_id is None or name is None:
                continue
            patients.append(
                Patient(
                    patient_id=patient_id,
                    display_name=name,
                    room=_clean_text(row[1]),
                    infectious=_is_yes(row[3]),
                    status=_clean_text(row[4]),
                )
            )
        return patients
    finally:
        wb.close()


def _column_values(ws, column_index: int) -> tuple[str, ...]:
    values: list[str] = []
    for row in ws.iter_rows(
        min_row=2,
        min_col=column_index,
        max_col=column_index,
        values_only=True,
    ):
        value = _clean_text(row[0])
        if value is not None:
            values.append(value)
    return tuple(values)


def read_settings(path: str | Path) -> WorkbookSettings:
    wb = _load(path)
    try:
        if "SETTINGS" not in wb.sheetnames:
            raise ValueError("Workbook does not contain SETTINGS")
        ws = wb["SETTINGS"]
        return WorkbookSettings(
            therapist_names=_column_values(ws, 1),
            standard_timeslots=tuple(
                _as_time(v)
                for v in _column_values(ws, 2)
                if _as_time(v) is not None
            ),
            reclined_timeslots=tuple(
                _as_time(v)
                for v in _column_values(ws, 3)
                if _as_time(v) is not None
            ),
            day_patterns=_column_values(ws, 4),
            therapies=_column_values(ws, 5),
            patient_statuses=_column_values(ws, 6),
            yes_no_values=_column_values(ws, 7),
            rooms=_column_values(ws, 8),
        )
    finally:
        wb.close()


def _planner_header_map(ws) -> dict[str, int]:
    first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {
        str(value).strip(): index
        for index, value in enumerate(first_row)
        if value is not None
    }


def _planner_blocks(headers: dict[str, int]) -> list[tuple[str, int, int, int | None]]:
    definitions = [
        ("ΦΘ", "ΦΘ_Ώρα", "ΦΘ_Ημέρες", "ΦΘ_Θεραπευτής"),
        ("Ρομποτικό", "Ρομποτικό_Ώρα", "Ρομποτικό_Ημέρες", "Ρομποτικό_Θεραπευτής"),
        ("Ανακλινόμενο", "Ανακλινόμενο_Ώρα", "Ανακλινόμενο_Ημέρες", None),
        ("Πισίνα", "Πισίνα_Ώρα", "Πισίνα_Ημέρες", None),
        ("Εργο", "Εργο_Ώρα", "Εργο_Ημέρες", None),
        ("Λογο", "Λογο_Ώρα", "Λογο_Ημέρες", None),
        ("ΕΦΑ", "ΕΦΑ_Ώρα", "ΕΦΑ_Ημέρες", None),
    ]
    blocks: list[tuple[str, int, int, int | None]] = []
    for treatment, time_name, days_name, therapist_name in definitions:
        if time_name not in headers or days_name not in headers:
            continue
        blocks.append(
            (
                treatment,
                headers[time_name],
                headers[days_name],
                headers.get(therapist_name) if therapist_name else None,
            )
        )
    return blocks


def read_base_schedule(path: str | Path) -> list[BaseScheduleEntry]:
    """Read recurring programme entries from PATIENT_PLANNER.

    A timed entry with an empty day pattern is interpreted as the confirmed
    operational convention ``Καθ/να`` (Monday-Friday). This rule is explicit
    and audited; unknown non-empty patterns still raise/are reported.
    """

    wb = _load(path)
    try:
        if "PATIENT_PLANNER" not in wb.sheetnames:
            raise ValueError("Workbook does not contain PATIENT_PLANNER")
        ws = wb["PATIENT_PLANNER"]
        headers = _planner_header_map(ws)
        required = {"PatientID", "Ασθενής"}
        if not required.issubset(headers):
            raise ValueError("PATIENT_PLANNER headers are incomplete")
        blocks = _planner_blocks(headers)
        entries: list[BaseScheduleEntry] = []
        for excel_row, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            patient_id = _patient_id(row[headers["PatientID"]])
            name = _clean_text(row[headers["Ασθενής"]])
            if patient_id is None or name is None:
                continue
            for treatment, time_col, days_col, therapist_col in blocks:
                start_time = _as_time(row[time_col])
                if start_time is None:
                    continue
                day_pattern = _clean_text(row[days_col]) or IMPLICIT_DAILY_PATTERN
                therapist_id = (
                    _clean_text(row[therapist_col])
                    if therapist_col is not None
                    else None
                )
                entries.append(
                    BaseScheduleEntry(
                        base_entry_id=f"planner:{excel_row}:{treatment}",
                        patient_id=patient_id,
                        treatment=treatment,
                        start_time=start_time,
                        day_pattern=day_pattern,
                        therapist_id=therapist_id,
                        robotic=treatment == "Ρομποτικό",
                    )
                )
        return entries
    finally:
        wb.close()


def audit_workbook(path: str | Path) -> WorkbookAudit:
    wb = _load(path)
    issues: list[AuditIssue] = []
    try:
        missing = sorted(REQUIRED_SHEETS - set(wb.sheetnames))
        if missing:
            issues.append(
                AuditIssue(
                    code="missing_required_sheets",
                    message=f"Missing required sheets: {', '.join(missing)}",
                    count=len(missing),
                    severity="error",
                )
            )
            return WorkbookAudit(str(path), 0, 0, 0, tuple(issues))

        # PATIENTS source-of-truth identity map.
        patients_ws = wb["PATIENTS"]
        patient_rows = [
            row
            for row in patients_ws.iter_rows(min_row=2, max_col=5, values_only=True)
            if _clean_text(row[2]) is not None
        ]
        patient_pairs = [(_patient_id(row[0]), _clean_text(row[2])) for row in patient_rows]
        patient_map = {pid: name for pid, name in patient_pairs if pid and name}
        duplicate_ids = len(patient_pairs) - len({pid for pid, _ in patient_pairs if pid})
        if duplicate_ids:
            issues.append(
                AuditIssue(
                    code="duplicate_patient_ids",
                    message="PATIENTS contains duplicate PatientID values",
                    count=duplicate_ids,
                    severity="error",
                )
            )

        planner_ws = wb["PATIENT_PLANNER"]
        headers = _planner_header_map(planner_ws)
        blocks = _planner_blocks(headers)
        planner_map: dict[str, str] = {}
        implicit_daily = 0
        missing_provider = 0
        invalid_pattern = 0
        entry_count = 0
        provider_count = 0
        for row in planner_ws.iter_rows(min_row=2, values_only=True):
            pid = _patient_id(row[headers["PatientID"]])
            name = _clean_text(row[headers["Ασθενής"]])
            if pid is None or name is None:
                continue
            planner_map[pid] = name
            for treatment, time_col, days_col, therapist_col in blocks:
                start_time = _as_time(row[time_col])
                if start_time is None:
                    continue
                day_pattern = _clean_text(row[days_col])
                if day_pattern is None:
                    implicit_daily += 1
                    day_pattern = IMPLICIT_DAILY_PATTERN
                try:
                    parse_day_pattern(day_pattern)
                except ValueError:
                    invalid_pattern += 1
                    continue
                entry_count += 1
                if therapist_col is not None:
                    provider_count += 1
                    if _clean_text(row[therapist_col]) is None:
                        missing_provider += 1

        sync_mismatches = sum(
            1
            for pid in set(patient_map) | set(planner_map)
            if patient_map.get(pid) != planner_map.get(pid)
        )
        if sync_mismatches:
            issues.append(
                AuditIssue(
                    code="patients_planner_identity_mismatch",
                    message="PATIENTS and PATIENT_PLANNER disagree on PatientID/name",
                    count=sync_mismatches,
                    severity="error",
                )
            )
        if implicit_daily:
            issues.append(
                AuditIssue(
                    code="timed_entries_assumed_daily",
                    message=(
                        "Timed PATIENT_PLANNER entries with an empty day pattern "
                        "are interpreted as Καθ/να (Monday-Friday) by confirmed rule"
                    ),
                    count=implicit_daily,
                    severity="info",
                )
            )
        if invalid_pattern:
            issues.append(
                AuditIssue(
                    code="invalid_day_patterns",
                    message="PATIENT_PLANNER contains unsupported day patterns",
                    count=invalid_pattern,
                    severity="error",
                )
            )
        if missing_provider:
            issues.append(
                AuditIssue(
                    code="provider_entries_missing_therapist",
                    message="ΦΘ/Ρομποτικό entries have no therapist",
                    count=missing_provider,
                )
            )

        # SESSIONS is audited but deliberately not used as the import source.
        if "SESSIONS" in wb.sheetnames:
            sessions_ws = wb["SESSIONS"]
            session_names_by_id: dict[str, set[str]] = {}
            for row in sessions_ws.iter_rows(min_row=2, max_col=8, values_only=True):
                pid = _patient_id(row[0])
                name = _clean_text(row[1])
                if pid is None or name is None:
                    continue
                session_names_by_id.setdefault(pid, set()).add(name)
            stale = sum(
                1
                for pid, names in session_names_by_id.items()
                if patient_map.get(pid) not in names
            )
            if stale:
                issues.append(
                    AuditIssue(
                        code="sessions_identity_drift",
                        message=(
                            "SESSIONS PatientID/name values disagree with PATIENTS; "
                            "SESSIONS is not used as the Python import source"
                        ),
                        count=stale,
                    )
                )

        return WorkbookAudit(
            workbook_path=str(path),
            patient_count=len(patient_map),
            base_entry_count=entry_count,
            provider_entry_count=provider_count,
            issues=tuple(issues),
        )
    finally:
        wb.close()
