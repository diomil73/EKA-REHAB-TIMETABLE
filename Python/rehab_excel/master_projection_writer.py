from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import shutil
import sys
from typing import Protocol

from openpyxl import load_workbook

from .master_projection import (
    MasterProjectionRow,
    build_master_projection,
    remap_master_formula,
)
from .native_excel import NativeStylePalette, excel_rgb


class MasterProjectionWriteError(RuntimeError):
    """Raised when a safe MASTER_SCHEDULE projection preview cannot be produced."""


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class MasterProjectionBackend(Protocol):
    def apply(
        self,
        workbook_path: Path,
        projection: tuple[MasterProjectionRow, ...],
    ) -> None:
        ...


class Win32ComMasterProjectionBackend:
    """Rebuild MASTER_SCHEDULE as a sorted projection of PATIENT_PLANNER."""

    _CANONICAL_SOURCE_COLUMNS = {
        1: "D",
        2: "B",
        3: "C",
        4: "F",
        5: "I",
        6: "N",
        7: "L",
        8: "P",
        9: "R",
        10: "T",
    }

    def __init__(self, palette: NativeStylePalette | None = None) -> None:
        self.palette = palette or NativeStylePalette()

    @classmethod
    def _schema(cls, ws) -> tuple[int, int]:
        header_row = cls._header_row(ws)
        k = str(ws.Cells(header_row, 11).Value or "").strip().casefold()
        l = str(ws.Cells(header_row, 12).Value or "").strip().casefold()
        m = str(ws.Cells(header_row, 13).Value or "").strip().casefold()
        if k == "κατάσταση".casefold():
            return 11, 11
        if (
            k == "ψυχολογοι".casefold()
            and l == "απογευματινο προγραμμα".casefold()
            and m == "κατάσταση".casefold()
        ):
            return 13, 13
        raise MasterProjectionWriteError("Unsupported MASTER_SCHEDULE column schema")

    @classmethod
    def _last_value_row(cls, ws) -> int:
        _total_cols, status_col = cls._schema(ws)
        rows = []
        for col in (1, 2, 3, status_col, 14, 15, 16):
            try:
                rows.append(int(ws.Cells(ws.Rows.Count, col).End(-4162).Row))
            except Exception:
                pass
        return max(rows or [1])

    @staticmethod
    def _header_row(ws) -> int:
        for row in range(1, 11):
            room = str(ws.Cells(row, 2).Value or "").strip().casefold()
            patient = str(ws.Cells(row, 3).Value or "").strip().casefold()
            if room == "θάλαμος".casefold() and patient == "ασθενής".casefold():
                return row
        raise MasterProjectionWriteError("MASTER_SCHEDULE header row was not found")

    @classmethod
    def _canonical_formula_templates(cls, status_col: int) -> dict[int, str]:
        formulas = {
            col: f"=PATIENT_PLANNER!{source_col}2"
            for col, source_col in cls._CANONICAL_SOURCE_COLUMNS.items()
        }
        formulas[status_col] = "=PATIENT_PLANNER!E2"
        return formulas

    @classmethod
    def _template_formulas(cls, ws) -> tuple[int, dict[int, str]]:
        header_row = cls._header_row(ws)
        template_row = header_row + 1
        total_cols, status_col = cls._schema(ws)
        required = set(range(1, 11)) | {status_col}
        formulas: dict[int, str] = {}

        last_row = max(template_row, cls._last_value_row(ws))
        for row in range(template_row, last_row + 1):
            for col in range(1, total_cols + 1):
                if col in formulas:
                    continue
                value = ws.Cells(row, col).Formula
                text = str(value or "")
                if text.startswith("=") and "PATIENT_PLANNER!" in text.upper():
                    formulas[col] = text
            if required.issubset(formulas):
                break

        canonical = cls._canonical_formula_templates(status_col)
        for col in required.difference(formulas):
            formulas[col] = canonical[col]

        missing = sorted(required.difference(formulas))
        if missing:
            raise MasterProjectionWriteError(
                f"MASTER_SCHEDULE could not resolve formula templates in columns: {missing}"
            )
        return template_row, formulas

    def _apply_infectious_style(self, row_range) -> None:
        row_range.Interior.Color = self.palette.infectious_yellow
        for edge in (7, 8, 9, 10):
            border = row_range.Borders(edge)
            border.LineStyle = 1
            border.Weight = 2
            border.Color = self.palette.infectious_yellow

    @staticmethod
    def _render_patient_identity(cell, item: MasterProjectionRow) -> None:
        doctor = str(item.responsible_doctor or "").strip()
        if not doctor:
            return

        cross = "✚"
        text = f"{item.display_name}\n{cross} {doctor}"
        cell.Value = text
        cell.WrapText = True
        cell.HorizontalAlignment = -4131
        cell.VerticalAlignment = -4108
        cell.Font.Color = excel_rgb(0, 0, 0)
        try:
            start = len(item.display_name) + 2
            cell.Characters(Start=start, Length=1).Font.Color = excel_rgb(220, 0, 0)
            cell.Characters(Start=start, Length=1).Font.Bold = True
        except Exception:
            pass

    def apply(
        self,
        workbook_path: Path,
        projection: tuple[MasterProjectionRow, ...],
    ) -> None:
        if sys.platform != "win32":
            raise MasterProjectionWriteError(
                "MASTER projection write-back requires Windows with Microsoft Excel"
            )
        try:
            import win32com.client  # type: ignore[import-not-found]
        except ImportError as exc:
            raise MasterProjectionWriteError(
                "pywin32 is required for MASTER projection write-back"
            ) from exc

        excel = None
        workbook = None
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.ScreenUpdating = False
            excel.EnableEvents = False

            workbook = excel.Workbooks.Open(
                str(workbook_path.resolve()),
                UpdateLinks=0,
                ReadOnly=False,
            )
            try:
                ws = workbook.Worksheets("MASTER_SCHEDULE")
            except Exception as exc:
                raise MasterProjectionWriteError(
                    "Workbook has no MASTER_SCHEDULE sheet"
                ) from exc

            template_row, templates = self._template_formulas(ws)
            total_cols, _status_col = self._schema(ws)
            old_last = self._last_value_row(ws)
            new_last = max(
                template_row - 1,
                max((item.target_row for item in projection), default=template_row - 1),
            )
            clear_last = max(old_last, new_last)

            for row in range(template_row, clear_last + 1):
                try:
                    if row > old_last:
                        end_col = "M" if total_cols == 13 else "K"
                        ws.Range(f"A{template_row}:{end_col}{template_row}").Copy()
                        ws.Range(f"A{row}:{end_col}{row}").PasteSpecial(Paste=-4122)
                    else:
                        # Preserve the established visual shell on existing MASTER
                        # rows. Only B:C need a clean baseline because infectious
                        # highlighting is reapplied after the projection is rebuilt.
                        ws.Range(f"B{template_row}:C{template_row}").Copy()
                        ws.Range(f"B{row}:C{row}").PasteSpecial(Paste=-4122)
                except Exception:
                    pass
                ws.Range(f"A{row}:P{row}").ClearContents()

            for item in projection:
                row = item.target_row
                for col, template in templates.items():
                    ws.Cells(row, col).Formula = remap_master_formula(
                        template,
                        item.planner_row,
                    )

                self._render_patient_identity(ws.Cells(row, 3), item)

                if item.infectious:
                    self._apply_infectious_style(ws.Range(f"B{row}:C{row}"))

            workbook.Application.CutCopyMode = False
            # Verification checks the written formulas and identities directly;
            # recalculating the entire workbook here adds substantial latency but
            # does not strengthen the preview safety guarantees.
            workbook.Save()
        except MasterProjectionWriteError:
            raise
        except Exception as exc:
            raise MasterProjectionWriteError(
                f"Excel MASTER projection failed: {exc}"
            ) from exc
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=False)
                except Exception:
                    pass
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass


@dataclass(frozen=True)
class MasterProjectionPreviewReport:
    source_path: str
    output_path: str
    projected_rows: int
    source_unchanged: bool
    verified_in_output: bool


def _verify_projection(
    output: Path,
    projection: tuple[MasterProjectionRow, ...],
) -> bool:
    wb = load_workbook(
        output,
        read_only=False,
        data_only=False,
        keep_vba=True,
    )
    try:
        if "MASTER_SCHEDULE" not in wb.sheetnames:
            return False
        ws = wb["MASTER_SCHEDULE"]
        header_row = 1
        for candidate in range(1, min(ws.max_row, 10) + 1):
            if (
                str(ws.cell(candidate, 2).value or "").strip().casefold() == "θάλαμος".casefold()
                and str(ws.cell(candidate, 3).value or "").strip().casefold() == "ασθενής".casefold()
            ):
                header_row = candidate
                break
        status_header = str(ws.cell(header_row, 13).value or "").strip().casefold()
        status_col = 13 if status_header == "κατάσταση".casefold() else 11
        total_cols = 13 if status_col == 13 else 11
        for item in projection:
            row = item.target_row
            expected = {
                1: f"=PATIENT_PLANNER!D{item.planner_row}",
                2: f"=PATIENT_PLANNER!B{item.planner_row}",
                status_col: f"=PATIENT_PLANNER!E{item.planner_row}",
            }
            if item.responsible_doctor:
                if str(ws.cell(row, 3).value or "") != (
                    f"{item.display_name}\n✚ {item.responsible_doctor}"
                ):
                    return False
            else:
                expected[3] = f"=PATIENT_PLANNER!C{item.planner_row}"

            for col, value in expected.items():
                if str(ws.cell(row, col).value or "") != value:
                    return False
            for col in (14, 15, 16):
                if ws.cell(row, col).value not in (None, ""):
                    return False

        tail_row = max((item.target_row for item in projection), default=header_row) + 1
        if tail_row <= ws.max_row:
            for col in range(1, total_cols + 1):
                if ws.cell(tail_row, col).value not in (None, ""):
                    return False
        return True
    finally:
        wb.close()


def create_master_projection_preview(
    source_path: str | Path,
    output_path: str | Path,
    *,
    backend: MasterProjectionBackend | None = None,
    overwrite: bool = False,
) -> MasterProjectionPreviewReport:
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()

    if not source.exists():
        raise MasterProjectionWriteError(f"Source workbook not found: {source}")
    if source == output:
        raise MasterProjectionWriteError("Output must differ from source workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise MasterProjectionWriteError("Source and output must both be .xlsm")
    if output.exists() and not overwrite:
        raise MasterProjectionWriteError(f"Output already exists: {output}")

    projection = build_master_projection(source)
    source_before = _sha256(source)

    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        try:
            output.unlink()
        except PermissionError as exc:
            raise MasterProjectionWriteError(
                f"Preview workbook is open or locked: {output}"
            ) from exc
    shutil.copy2(source, output)

    selected = backend or Win32ComMasterProjectionBackend()
    try:
        selected.apply(output, projection)
    except Exception:
        try:
            output.unlink()
        except OSError:
            pass
        raise

    source_after = _sha256(source)
    if source_before != source_after:
        try:
            output.unlink()
        except OSError:
            pass
        raise MasterProjectionWriteError(
            "Source workbook changed during MASTER projection preview"
        )

    verified = _verify_projection(output, projection)
    if not verified:
        try:
            output.unlink()
        except OSError:
            pass
        raise MasterProjectionWriteError(
            "MASTER projection verification failed after write-back"
        )

    return MasterProjectionPreviewReport(
        source_path=str(source),
        output_path=str(output),
        projected_rows=len(projection),
        source_unchanged=True,
        verified_in_output=True,
    )


def refresh_master_projection_in_place(
    workbook_path: str | Path,
    *,
    backend: MasterProjectionBackend | None = None,
) -> int:
    """Rebuild MASTER_SCHEDULE inside an already-created working/preview workbook."""

    path = Path(workbook_path).resolve()
    if not path.exists():
        raise MasterProjectionWriteError(f"Workbook not found: {path}")
    if path.suffix.casefold() != ".xlsm":
        raise MasterProjectionWriteError("MASTER projection refresh requires .xlsm")

    projection = build_master_projection(path)
    selected = backend or Win32ComMasterProjectionBackend()
    selected.apply(path, projection)

    if not _verify_projection(path, projection):
        raise MasterProjectionWriteError(
            "MASTER projection verification failed after in-place refresh"
        )
    return len(projection)
