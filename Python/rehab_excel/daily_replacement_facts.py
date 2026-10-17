from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
import sys
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.models import ReplacementAssignment, ReplacementProviderKind, Session


SHEET_NAME = "DAILY_REPLACEMENT_FACTS"
ACTIVE_STATUS = "ACTIVE"
HEADERS = (
    "ReplacementId",
    "TargetSessionId",
    "PatientId",
    "OriginalTherapistId",
    "ReplacementTherapistId",
    "ReplacementDate",
    "ReplacementTime",
    "Reason",
    "ProviderKind",
    "Status",
)


class DailyReplacementFactError(RuntimeError):
    """Raised when persisted daily replacement facts are invalid or conflicting."""


@dataclass(frozen=True)
class DailyReplacementFactRow:
    replacement_id: str
    target_session_id: str
    patient_id: str
    original_therapist_id: str
    replacement_therapist_id: str
    replacement_date: date
    replacement_time: time
    reason: str | None
    provider_kind: ReplacementProviderKind
    status: str = ACTIVE_STATUS

    @property
    def target_key(self) -> tuple[str, date]:
        return self.target_session_id, self.replacement_date

    @property
    def semantic_key(self) -> tuple[object, ...]:
        return (
            self.target_session_id,
            self.patient_id,
            self.original_therapist_id,
            self.replacement_therapist_id,
            self.replacement_date,
            self.replacement_time,
            self.provider_kind,
            self.status,
        )


def _normalized_time(value: time) -> time:
    return value.replace(second=0, microsecond=0)


def _parse_date(value: object, *, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            # COM can persist local midnight as a timezone-shifted Excel serial
            # (for example 2026-10-19 -> 46313.875). ReplacementDate is a
            # calendar-only field, so normalize a legacy numeric value to the
            # nearest whole Excel day before converting it back to a date.
            serial_day = int(float(value) + 0.5)
            return (datetime(1899, 12, 30) + timedelta(days=serial_day)).date()
        except (OverflowError, TypeError, ValueError):
            pass
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyReplacementFactError(
        f"{SHEET_NAME} row {row}: invalid ReplacementDate {value!r}"
    )


def _parse_time(value: object, *, row: int) -> time:
    if isinstance(value, datetime):
        return _normalized_time(value.time())
    if isinstance(value, time):
        return _normalized_time(value)
    if isinstance(value, (int, float)):
        seconds = int(round((float(value) % 1) * 24 * 60 * 60)) % (24 * 60 * 60)
        return time(seconds // 3600, (seconds % 3600) // 60)
    text = str(value or "").strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return _normalized_time(datetime.strptime(text, fmt).time())
        except ValueError:
            pass
    raise DailyReplacementFactError(
        f"{SHEET_NAME} row {row}: invalid ReplacementTime {value!r}"
    )


def _provider_kind(value: object, *, row: int) -> ReplacementProviderKind:
    text = str(value or "").strip().casefold()
    if not text:
        return ReplacementProviderKind.THERAPIST
    try:
        return ReplacementProviderKind(text)
    except ValueError as exc:
        raise DailyReplacementFactError(
            f"{SHEET_NAME} row {row}: invalid ProviderKind {value!r}"
        ) from exc


def _validate_row_against_session(
    fact: DailyReplacementFactRow,
    session_by_id: dict[str, Session],
) -> None:
    session = session_by_id.get(fact.target_session_id)
    if session is None:
        raise DailyReplacementFactError(
            f"Unknown target session {fact.target_session_id!r} in {SHEET_NAME}"
        )
    if fact.patient_id != session.patient_id:
        raise DailyReplacementFactError(
            f"Replacement {fact.replacement_id!r}: patient {fact.patient_id!r} does not "
            f"match session patient {session.patient_id!r}"
        )
    if fact.original_therapist_id != session.therapist_id:
        raise DailyReplacementFactError(
            f"Replacement {fact.replacement_id!r}: original therapist "
            f"{fact.original_therapist_id!r} does not match session therapist "
            f"{session.therapist_id!r}"
        )
    if fact.replacement_date != session.session_date:
        raise DailyReplacementFactError(
            f"Replacement {fact.replacement_id!r}: date {fact.replacement_date} does not "
            f"match session date {session.session_date}"
        )
    if not fact.replacement_therapist_id.strip():
        raise DailyReplacementFactError(
            f"Replacement {fact.replacement_id!r}: replacement provider is blank"
        )
    if fact.status != ACTIVE_STATUS:
        raise DailyReplacementFactError(
            f"Replacement {fact.replacement_id!r}: unsupported status {fact.status!r}"
        )


def build_daily_replacement_fact_rows(
    replacements: Iterable[ReplacementAssignment],
    *,
    sessions: Iterable[Session],
) -> tuple[DailyReplacementFactRow, ...]:
    session_by_id = {session.session_id: session for session in sessions}
    rows: list[DailyReplacementFactRow] = []
    seen_targets: dict[tuple[str, date], DailyReplacementFactRow] = {}
    seen_ids: set[str] = set()

    for replacement in replacements:
        row = DailyReplacementFactRow(
            replacement_id=str(replacement.replacement_id).strip(),
            target_session_id=str(replacement.target_session_id).strip(),
            patient_id=str(replacement.patient_id).strip(),
            original_therapist_id=str(replacement.original_therapist_id).strip(),
            replacement_therapist_id=str(replacement.replacement_therapist_id).strip(),
            replacement_date=replacement.replacement_date,
            replacement_time=_normalized_time(replacement.replacement_time),
            reason=replacement.reason,
            provider_kind=replacement.replacement_provider_kind,
        )
        if not row.replacement_id:
            raise DailyReplacementFactError("ReplacementId cannot be blank")
        if row.replacement_id in seen_ids:
            raise DailyReplacementFactError(
                f"Duplicate ReplacementId {row.replacement_id!r}"
            )
        seen_ids.add(row.replacement_id)
        _validate_row_against_session(row, session_by_id)
        existing = seen_targets.get(row.target_key)
        if existing is not None:
            raise DailyReplacementFactError(
                f"Multiple active replacement facts for target session/date {row.target_key!r}"
            )
        seen_targets[row.target_key] = row
        rows.append(row)

    return tuple(rows)


def _read_rows_from_sheet(ws, *, target_date: date | None) -> tuple[DailyReplacementFactRow, ...]:
    actual_headers = tuple(str(ws.cell(1, col).value or "").strip() for col in range(1, len(HEADERS) + 1))
    if actual_headers != HEADERS:
        raise DailyReplacementFactError(
            f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
        )

    rows: list[DailyReplacementFactRow] = []
    seen_targets: dict[tuple[str, date], DailyReplacementFactRow] = {}
    seen_ids: set[str] = set()
    for row_number in range(2, ws.max_row + 1):
        values = [ws.cell(row_number, col).value for col in range(1, len(HEADERS) + 1)]
        if not any(value is not None and str(value).strip() for value in values):
            continue
        replacement_id = str(values[0] or "").strip()
        target_session_id = str(values[1] or "").strip()
        patient_id = str(values[2] or "").strip()
        original_therapist_id = str(values[3] or "").strip()
        replacement_therapist_id = str(values[4] or "").strip()
        replacement_date = _parse_date(values[5], row=row_number)
        if target_date is not None and replacement_date != target_date:
            continue
        replacement_time = _parse_time(values[6], row=row_number)
        reason_text = str(values[7] or "").strip()
        provider_kind = _provider_kind(values[8], row=row_number)
        status = str(values[9] or ACTIVE_STATUS).strip().upper()
        fact = DailyReplacementFactRow(
            replacement_id=replacement_id,
            target_session_id=target_session_id,
            patient_id=patient_id,
            original_therapist_id=original_therapist_id,
            replacement_therapist_id=replacement_therapist_id,
            replacement_date=replacement_date,
            replacement_time=replacement_time,
            reason=reason_text or None,
            provider_kind=provider_kind,
            status=status,
        )
        if not fact.replacement_id:
            raise DailyReplacementFactError(
                f"{SHEET_NAME} row {row_number}: ReplacementId is blank"
            )
        if fact.replacement_id in seen_ids:
            raise DailyReplacementFactError(
                f"Duplicate ReplacementId {fact.replacement_id!r}"
            )
        seen_ids.add(fact.replacement_id)
        existing = seen_targets.get(fact.target_key)
        if existing is not None:
            raise DailyReplacementFactError(
                f"Multiple active replacement facts for target session/date {fact.target_key!r}"
            )
        seen_targets[fact.target_key] = fact
        rows.append(fact)
    return tuple(rows)


def read_daily_replacement_facts(
    workbook_path: str | Path,
    *,
    sessions: Iterable[Session],
    target_date: date | None = None,
) -> tuple[ReplacementAssignment, ...]:
    """Read structured date-specific replacement overlays from a workbook.

    Missing sheet means no persisted replacement facts, which keeps older
    workbooks backward-compatible. The function never saves the workbook.
    """

    path = Path(workbook_path)
    session_list = tuple(sessions)
    session_by_id = {session.session_id: session for session in session_list}
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if SHEET_NAME not in wb.sheetnames:
            return ()
        facts = _read_rows_from_sheet(wb[SHEET_NAME], target_date=target_date)
    finally:
        wb.close()

    assignments: list[ReplacementAssignment] = []
    for fact in facts:
        _validate_row_against_session(fact, session_by_id)
        assignments.append(
            ReplacementAssignment(
                replacement_id=fact.replacement_id,
                target_session_id=fact.target_session_id,
                patient_id=fact.patient_id,
                original_therapist_id=fact.original_therapist_id,
                replacement_therapist_id=fact.replacement_therapist_id,
                replacement_date=fact.replacement_date,
                replacement_time=fact.replacement_time,
                reason=fact.reason,
                replacement_provider_kind=fact.provider_kind,
            )
        )
    return tuple(assignments)


def _time_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):
        seconds = int(round((float(value) % 1) * 24 * 60 * 60)) % (24 * 60 * 60)
        return time(seconds // 3600, (seconds % 3600) // 60).strftime("%H:%M")
    return str(value or "").strip()[:5]


def _existing_rows_com(ws) -> dict[tuple[str, date], DailyReplacementFactRow]:
    actual_headers = tuple(str(ws.Cells(1, col).Value or "").strip() for col in range(1, len(HEADERS) + 1))
    if actual_headers != HEADERS:
        raise DailyReplacementFactError(
            f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
        )
    existing: dict[tuple[str, date], DailyReplacementFactRow] = {}
    last_row = max(int(ws.UsedRange.Rows.Count), 1)
    seen_ids: set[str] = set()
    for row_number in range(2, last_row + 1):
        replacement_id = str(ws.Cells(row_number, 1).Value or "").strip()
        if not replacement_id:
            continue
        replacement_date = _parse_date(ws.Cells(row_number, 6).Value, row=row_number)
        fact = DailyReplacementFactRow(
            replacement_id=replacement_id,
            target_session_id=str(ws.Cells(row_number, 2).Value or "").strip(),
            patient_id=str(ws.Cells(row_number, 3).Value or "").strip(),
            original_therapist_id=str(ws.Cells(row_number, 4).Value or "").strip(),
            replacement_therapist_id=str(ws.Cells(row_number, 5).Value or "").strip(),
            replacement_date=replacement_date,
            replacement_time=_parse_time(_time_text(ws.Cells(row_number, 7).Value), row=row_number),
            reason=str(ws.Cells(row_number, 8).Value or "").strip() or None,
            provider_kind=_provider_kind(ws.Cells(row_number, 9).Value, row=row_number),
            status=str(ws.Cells(row_number, 10).Value or ACTIVE_STATUS).strip().upper(),
        )
        if replacement_id in seen_ids:
            raise DailyReplacementFactError(f"Duplicate ReplacementId {replacement_id!r}")
        seen_ids.add(replacement_id)
        if fact.target_key in existing:
            raise DailyReplacementFactError(
                f"Multiple active replacement facts for target session/date {fact.target_key!r}"
            )
        existing[fact.target_key] = fact
    return existing


def append_daily_replacement_facts(
    workbook_path: str | Path,
    replacements: Iterable[ReplacementAssignment],
    *,
    sessions: Iterable[Session],
) -> int:
    """Append validated replacement overlays to the output workbook.

    Existing semantically identical facts are idempotent. A second different
    active replacement for the same target session/date is rejected.
    """

    rows = build_daily_replacement_fact_rows(replacements, sessions=sessions)
    if not rows:
        return 0
    if sys.platform != "win32":
        raise DailyReplacementFactError(
            "Daily replacement persistence requires Windows + Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DailyReplacementFactError("pywin32 is required") from exc

    path = Path(workbook_path).resolve()
    excel = None
    workbook = None
    written = 0
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)

        ws = None
        for sheet in workbook.Worksheets:
            if str(sheet.Name).casefold() == SHEET_NAME.casefold():
                ws = sheet
                break
        if ws is None:
            ws = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
            ws.Name = SHEET_NAME
            for col, header in enumerate(HEADERS, start=1):
                ws.Cells(1, col).Value = header
            ws.Visible = 2  # xlSheetVeryHidden
            existing: dict[tuple[str, date], DailyReplacementFactRow] = {}
        else:
            existing = _existing_rows_com(ws)
            ws.Visible = 2

        next_row = max(int(ws.UsedRange.Rows.Count) + 1, 2)
        for fact in rows:
            old = existing.get(fact.target_key)
            if old is not None:
                if old.semantic_key == fact.semantic_key:
                    continue
                raise DailyReplacementFactError(
                    f"Conflicting active replacement fact for target session/date {fact.target_key!r}"
                )
            ws.Cells(next_row, 1).Value = fact.replacement_id
            ws.Cells(next_row, 2).Value = fact.target_session_id
            ws.Cells(next_row, 3).Value = fact.patient_id
            ws.Cells(next_row, 4).Value = fact.original_therapist_id
            ws.Cells(next_row, 5).Value = fact.replacement_therapist_id
            ws.Cells(next_row, 6).NumberFormat = "@"
            ws.Cells(next_row, 6).Value2 = fact.replacement_date.isoformat()
            ws.Cells(next_row, 7).NumberFormat = "@"
            ws.Cells(next_row, 7).Value2 = fact.replacement_time.strftime("%H:%M")
            ws.Cells(next_row, 8).Value = fact.reason or ""
            ws.Cells(next_row, 9).Value = fact.provider_kind.value
            ws.Cells(next_row, 10).Value = fact.status
            existing[fact.target_key] = fact
            next_row += 1
            written += 1
        workbook.Save()
    except DailyReplacementFactError:
        raise
    except Exception as exc:
        raise DailyReplacementFactError(
            f"Could not persist daily replacement facts: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass

    return written
