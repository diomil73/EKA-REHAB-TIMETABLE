from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook

from rehab_core.base_schedule import materialize_sessions_for_date
from rehab_core.daily_state import build_daily_session_states

from . import patient_registration_vertical_slice as _vertical
from .native_excel import apply_write_plan_to_copy
from .patient_registration_form_vba import PATIENT_FORM_NAME, install_patient_form
from .patient_scheduling_form_vba import install_patient_scheduling_ui
from .reader import read_base_schedule, read_patients
from .render_plan import build_daily_excel_render_plan
from .render_writeback import build_daily_write_plan

PatientRegistrationBridgeVbaError = _vertical.PatientRegistrationBridgeVbaError
_BASE_VERTICAL_FORM_CODE = _vertical._vertical_form_code


def _scheduling_form_code() -> str:
    code = _BASE_VERTICAL_FORM_CODE()
    old = '''    Unload Me
    GoToMaster
    MsgBox "Ο ασθενής καταχωρήθηκε επιτυχώς." & vbCrLf & _
           "Patient ID: " & patientId, _
           vbInformation, "Νέος ασθενής"
    Exit Sub'''
    new = '''    MsgBox "Ο ασθενής καταχωρήθηκε επιτυχώς." & vbCrLf & _
           "Patient ID: " & patientId & vbCrLf & vbCrLf & _
           "Συνεχίζουμε στον προγραμματισμό θεραπείας.", _
           vbInformation, "Νέος ασθενής"
    Unload Me
    OpenPatientScheduling patientId
    GoToMaster
    Exit Sub'''
    if old not in code:
        raise RuntimeError("Patient scheduling continuation marker missing")
    return code.replace(old, new, 1)


def _position_control(control, left: float, top: float) -> None:
    try:
        control.Left = left
    except Exception:
        pass
    try:
        control.Top = top
    except Exception:
        pass


def _has_vb_component(vbproject, component_name: str) -> bool:
    """Return True when a named VBA component already exists.

    VBComponents(name) raises COM errors when the component is absent, so this
    helper deliberately treats lookup failures as "missing".
    """

    try:
        return vbproject.VBComponents(component_name) is not None
    except Exception:
        return False


def _install_base_patient_form(workbook_path: Path) -> None:
    """Ensure frmNewPatient exists on a private prerequisite copy.

    Clean repository baselines do not contain the form, while versioned smoke
    previews may already contain it. Preserve an existing form so rebuilding
    from a mature preview is idempotent and does not ask the VBE to remove and
    recreate an already-loaded UserForm.
    """

    if sys.platform != "win32":
        raise PatientRegistrationBridgeVbaError(
            "Patient form installation requires Windows with Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise PatientRegistrationBridgeVbaError("pywin32 is required") from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(workbook_path.resolve()), UpdateLinks=0, ReadOnly=False)

        if not _has_vb_component(workbook.VBProject, PATIENT_FORM_NAME):
            install_patient_form(
                workbook.VBProject,
                position_control=_position_control,
            )
            workbook.Save()
    except PatientRegistrationBridgeVbaError:
        raise
    except Exception as exc:
        raise PatientRegistrationBridgeVbaError(
            f"Patient form prerequisite installation failed: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def _install_scheduling_ui(workbook_path: Path) -> None:
    if sys.platform != "win32":
        raise PatientRegistrationBridgeVbaError(
            "Patient scheduling UI installation requires Windows with Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise PatientRegistrationBridgeVbaError("pywin32 is required") from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(workbook_path.resolve()), UpdateLinks=0, ReadOnly=False)
        install_patient_scheduling_ui(
            workbook.VBProject,
            position_control=_position_control,
        )
        workbook.Save()
    except PatientRegistrationBridgeVbaError:
        raise
    except Exception as exc:
        raise PatientRegistrationBridgeVbaError(
            f"Patient scheduling UI installation failed: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def _daily_target_date(workbook_path: Path) -> date | None:
    """Read the dated operational view from DAILY_INPUT!B2 when available."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_vba=True)
    try:
        if "DAILY_INPUT" not in workbook.sheetnames:
            return None
        value = workbook["DAILY_INPUT"]["B2"].value
    finally:
        workbook.close()

    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise PatientRegistrationBridgeVbaError(
        f"DAILY_INPUT!B2 has unsupported date value: {value!r}"
    )


def _refresh_daily_visuals(workbook_path: Path) -> None:
    """Reapply current THERAPIST_DAILY rich-text roles on a mature preview.

    Versioned previews already contain the daily sheet text. Rebuilding a new
    scheduling slice must not carry stale line formatting forward, so when a
    DAILY_INPUT date is available we regenerate the semantic daily render and
    apply it through the native Excel rich-text writer on a private copy.
    """

    target_date = _daily_target_date(workbook_path)
    if target_date is None:
        return

    daily_copy = workbook_path.with_name(f".{workbook_path.stem}.daily-render.xlsm")
    daily_copy.unlink(missing_ok=True)
    try:
        patients = read_patients(workbook_path)
        base_entries = read_base_schedule(workbook_path)
        sessions = materialize_sessions_for_date(base_entries, target_date)
        states = build_daily_session_states(
            sessions,
            absences=(),
            replacements=(),
            cancellations=(),
            target_date=target_date,
        )
        render_plan = build_daily_excel_render_plan(
            workbook_path,
            states,
            sessions,
            patients,
        )
        write_plan = build_daily_write_plan(render_plan)
        apply_write_plan_to_copy(write_plan, daily_copy, overwrite=True)
        daily_copy.replace(workbook_path)
    except PatientRegistrationBridgeVbaError:
        raise
    except Exception as exc:
        raise PatientRegistrationBridgeVbaError(
            f"THERAPIST_DAILY visual refresh failed: {exc}"
        ) from exc
    finally:
        daily_copy.unlink(missing_ok=True)


def create_patient_registration_bridge_preview(
    source_path: str | Path,
    output_path: str | Path,
    *,
    overwrite: bool = False,
):
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    prerequisite = output.with_name(f".{output.stem}.patient-form-prerequisite.xlsm")

    if not source.exists():
        raise PatientRegistrationBridgeVbaError(f"Source workbook not found: {source}")
    if source == output:
        raise PatientRegistrationBridgeVbaError("Output must be different from source workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise PatientRegistrationBridgeVbaError("Source and output must both be .xlsm files")
    if output.exists() and not overwrite:
        raise PatientRegistrationBridgeVbaError(f"Output already exists: {output}")

    prerequisite.unlink(missing_ok=True)
    shutil.copy2(source, prerequisite)

    try:
        _install_base_patient_form(prerequisite)

        original_form_builder = _vertical._vertical_form_code
        try:
            _vertical._vertical_form_code = _scheduling_form_code
            report = _vertical.create_patient_registration_bridge_preview(
                prerequisite,
                output,
                overwrite=overwrite,
            )
        finally:
            _vertical._vertical_form_code = original_form_builder

        _install_scheduling_ui(output)
        _refresh_daily_visuals(output)
        return report
    finally:
        prerequisite.unlink(missing_ok=True)
