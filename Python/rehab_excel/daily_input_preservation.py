from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class DailyInputSnapshot:
    cells: tuple[tuple[str, object], ...]


class DailyInputPreservationError(RuntimeError):
    pass


def snapshot_daily_input(path: str | Path) -> DailyInputSnapshot:
    workbook_path = Path(path)
    wb = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputPreservationError("Workbook has no DAILY_INPUT sheet")
        ws = wb["DAILY_INPUT"]
        cells: list[tuple[str, object]] = []
        for row in ws.iter_rows(
            min_row=1,
            max_row=max(1, ws.max_row),
            min_col=1,
            max_col=max(1, ws.max_column),
        ):
            for cell in row:
                cells.append((cell.coordinate, cell.value))
        return DailyInputSnapshot(tuple(cells))
    finally:
        wb.close()


def assert_daily_input_preserved(
    expected: DailyInputSnapshot,
    path: str | Path,
) -> None:
    actual = snapshot_daily_input(path)
    expected_map = dict(expected.cells)
    actual_map = dict(actual.cells)

    addresses = sorted(set(expected_map) | set(actual_map))
    mismatches = [
        address
        for address in addresses
        if expected_map.get(address) != actual_map.get(address)
    ]
    if mismatches:
        sample = ", ".join(mismatches[:10])
        raise DailyInputPreservationError(
            "DAILY_INPUT changed during operational write-back at: " + sample
        )
