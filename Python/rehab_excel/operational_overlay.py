from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Mapping
import unicodedata

from openpyxl import load_workbook

from rehab_core.daily_state import DailySessionState, DailySessionStatus
from rehab_core.models import Patient, ReplacementProviderKind, Session, Student
from rehab_core.students import student_display_state

from .layout_contract import LayoutSafetyError, resolve_therapist_daily_cell
from .writeback import CellPatch, TextRunPatch, WritePlan, build_write_plan


class OperationalOverlayError(RuntimeError):
    """Raised when a daily overlay cannot be applied without guessing."""


@dataclass(frozen=True)
class OverlayBinding:
    session_id: str
    patient_id: str
    status: DailySessionStatus
    original_cell: str
    effective_cell: str | None


@dataclass(frozen=True)
class OperationalOverlayPlan:
    write_plan: WritePlan
    bindings: tuple[OverlayBinding, ...]


@dataclass
class _CellDraft:
    text: str
    text_runs: list[TextRunPatch]
    fill_role: str | None = None
    border_role: str | None = None


_INFECTIOUS_RGB = {"FFFFFF43", "FFFF43", "FFFFFF00", "FFFF00"}
_ROBOTIC_RGB = {"FFF8CBAD", "F8CBAD"}


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _line_spans(text: str) -> list[tuple[int, int, str]]:
    """Return 1-based start, visible length, visible line text."""

    spans: list[tuple[int, int, str]] = []
    cursor = 0
    for raw in text.splitlines(keepends=True):
        visible = raw.rstrip("\r\n")
        if visible:
            spans.append((cursor + 1, len(visible), visible))
        cursor += len(raw)
    if text and not spans:
        spans.append((1, len(text), text))
    return spans


def _matching_patient_span(text: str, patient_name: str) -> tuple[int, int, str]:
    wanted = _norm(patient_name)
    matches: list[tuple[int, int, str]] = []
    for span in _line_spans(text):
        line_norm = _norm(span[2])
        if line_norm == wanted or line_norm.startswith(wanted + " ") or line_norm.startswith(wanted + "["):
            matches.append(span)
    if len(matches) != 1:
        raise OperationalOverlayError(
            f"Expected exactly one line for patient {patient_name!r}; found {len(matches)} in {text!r}"
        )
    return matches[0]


def _append_unique_line(text: str, line: str) -> tuple[str, tuple[int, int, str]]:
    existing = _line_spans(text)
    wanted = _norm(line)
    for span in existing:
        if _norm(span[2]) == wanted:
            return text, span
    new_text = f"{text}\n{line}" if text else line
    return new_text, _line_spans(new_text)[-1]


def _rgb(cell) -> str | None:
    color = cell.fill.fgColor
    if color is None or color.type != "rgb":
        return None
    return (color.rgb or "").upper()


def _existing_roles(cell) -> tuple[str | None, str | None]:
    rgb = _rgb(cell)
    if rgb in _ROBOTIC_RGB:
        return "robotic_pink", None
    if rgb in _INFECTIOUS_RGB:
        return "infectious_yellow", "infectious_yellow"
    return None, None


def _merge_roles(
    fill_role: str | None,
    border_role: str | None,
    *,
    robotic: bool,
    infectious: bool,
) -> tuple[str | None, str | None]:
    if robotic:
        if fill_role == "infectious_yellow":
            border_role = "infectious_yellow"
        fill_role = "robotic_pink"
    elif infectious and fill_role is None:
        fill_role = "infectious_yellow"
    if infectious:
        border_role = "infectious_yellow"
    return fill_role, border_role


def _student_aliases(student: Student, target_date: date) -> tuple[str, ...]:
    display = student_display_state(student, target_date)
    return (
        display.label,
        f"φοιτ {student.student_number}",
        f"φοιτητής {student.student_number}",
    )


def _provider_display(
    provider_id: str,
    *,
    provider_kind: ReplacementProviderKind,
    labels: Mapping[str, str],
    students: Mapping[str, Student],
    target_date: date,
) -> tuple[str, str]:
    if provider_kind == ReplacementProviderKind.STUDENT:
        student = students.get(provider_id)
        if student is not None:
            state = student_display_state(student, target_date)
            return state.label, "student_active_green" if state.use_green_font else "default"
    return labels.get(provider_id, provider_id), "default"


def _resolve_provider_cell(
    workbook_path: str | Path,
    provider_id: str,
    slot_time,
    *,
    provider_kind: ReplacementProviderKind,
    labels: Mapping[str, str],
    students: Mapping[str, Student],
    target_date: date,
) -> str:
    candidates: list[str] = []
    if provider_kind == ReplacementProviderKind.STUDENT:
        student = students.get(provider_id)
        if student is not None:
            candidates.extend(_student_aliases(student, target_date))
    candidates.extend((labels.get(provider_id, provider_id), provider_id))

    last_error: Exception | None = None
    seen: set[str] = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        try:
            return resolve_therapist_daily_cell(workbook_path, candidate, slot_time)
        except LayoutSafetyError as exc:
            last_error = exc
    if last_error is not None:
        raise OperationalOverlayError(str(last_error)) from last_error
    raise OperationalOverlayError(f"No provider cell found for {provider_id!r}")


def _read_daily_cells(workbook_path: str | Path):
    wb = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_vba=True,
    )
    try:
        ws = wb["THERAPIST_DAILY"]
        values: dict[str, str] = {}
        roles: dict[str, tuple[str | None, str | None]] = {}
        for row in ws.iter_rows():
            for cell in row:
                values[cell.coordinate] = "" if cell.value is None else str(cell.value).strip()
                roles[cell.coordinate] = _existing_roles(cell)
        return values, roles
    finally:
        wb.close()


def build_operational_overlay_write_plan(
    workbook_path: str | Path,
    states: Iterable[DailySessionState],
    sessions: Iterable[Session],
    patients: Iterable[Patient],
    *,
    provider_labels: Mapping[str, str] | None = None,
    students: Iterable[Student] = (),
) -> OperationalOverlayPlan:
    """Build a minimal daily overlay on top of the existing THERAPIST_DAILY text.

    Only non-active states are touched. Existing unrelated lines are kept verbatim.
    Replacements keep the original patient visible in muted italics and append a
    provider/time note below it. Absence/cancellation-style states strike the
    original patient line. The base workbook itself is never mutated here.
    """

    workbook_path = Path(workbook_path)
    state_list = tuple(state for state in states if state.status != DailySessionStatus.ACTIVE)
    if not state_list:
        raise OperationalOverlayError("No non-active daily states were provided")

    dates = {state.session_date for state in state_list}
    if len(dates) != 1:
        raise OperationalOverlayError("Operational overlay requires exactly one target date")
    target_date = next(iter(dates))

    session_by_id = {session.session_id: session for session in sessions}
    patient_by_id = {patient.patient_id: patient for patient in patients}
    student_by_id = {student.student_id: student for student in students}
    labels = dict(provider_labels or {})
    values, existing_roles = _read_daily_cells(workbook_path)

    drafts: dict[str, _CellDraft] = {}
    bindings: list[OverlayBinding] = []

    def draft_for(cell: str) -> _CellDraft:
        draft = drafts.get(cell)
        if draft is None:
            fill, border = existing_roles.get(cell, (None, None))
            draft = _CellDraft(values.get(cell, ""), [], fill, border)
            drafts[cell] = draft
        return draft

    for state in state_list:
        session = session_by_id.get(state.session_id)
        if session is None:
            raise OperationalOverlayError(f"Missing Session for state {state.session_id!r}")
        patient = patient_by_id.get(state.patient_id)
        if patient is None:
            raise OperationalOverlayError(f"Missing Patient {state.patient_id!r}")

        original_cell = _resolve_provider_cell(
            workbook_path,
            state.original_therapist_id,
            state.original_time,
            provider_kind=ReplacementProviderKind.THERAPIST,
            labels=labels,
            students=student_by_id,
            target_date=target_date,
        )
        original = draft_for(original_cell)
        patient_span = _matching_patient_span(original.text, patient.display_name)

        # Remove any prior run for the same line in this plan, then apply the
        # state-specific presentation. This makes repeated operations deterministic.
        original.text_runs = [
            run
            for run in original.text_runs
            if not (run.start == patient_span[0] and run.length == patient_span[1])
        ]
        original.text_runs.append(
            TextRunPatch(
                start=patient_span[0],
                length=patient_span[1],
                strike_through=state.status
                in {DailySessionStatus.PATIENT_ABSENT, DailySessionStatus.THERAPIST_ABSENT},
                italic=state.status == DailySessionStatus.REPLACED,
                font_role="muted",
            )
        )
        original.fill_role, original.border_role = _merge_roles(
            original.fill_role,
            original.border_role,
            robotic=session.robotic,
            infectious=patient.infectious,
        )

        effective_cell: str | None = None
        if state.status == DailySessionStatus.REPLACED:
            if state.effective_therapist_id is None or state.effective_time is None:
                raise OperationalOverlayError(
                    f"Replacement state {state.session_id!r} lacks effective provider/time"
                )
            provider_kind = state.effective_provider_kind or ReplacementProviderKind.THERAPIST
            provider_label, provider_font = _provider_display(
                state.effective_therapist_id,
                provider_kind=provider_kind,
                labels=labels,
                students=student_by_id,
                target_date=target_date,
            )
            note = f"→ {provider_label} {state.effective_time.strftime('%H:%M')}"
            original.text, note_span = _append_unique_line(original.text, note)
            original.text_runs.append(
                TextRunPatch(
                    start=note_span[0],
                    length=note_span[1],
                    strike_through=False,
                    italic=False,
                    font_role=provider_font,
                )
            )

            effective_cell = _resolve_provider_cell(
                workbook_path,
                state.effective_therapist_id,
                state.effective_time,
                provider_kind=provider_kind,
                labels=labels,
                students=student_by_id,
                target_date=target_date,
            )
            if effective_cell != original_cell:
                effective = draft_for(effective_cell)
                effective.text, replacement_span = _append_unique_line(
                    effective.text, patient.display_name
                )
                effective.text_runs.append(
                    TextRunPatch(
                        start=replacement_span[0],
                        length=replacement_span[1],
                        strike_through=False,
                        italic=False,
                        font_role="default",
                    )
                )
                effective.fill_role, effective.border_role = _merge_roles(
                    effective.fill_role,
                    effective.border_role,
                    robotic=session.robotic,
                    infectious=patient.infectious,
                )

        bindings.append(
            OverlayBinding(
                session_id=state.session_id,
                patient_id=state.patient_id,
                status=state.status,
                original_cell=original_cell,
                effective_cell=effective_cell,
            )
        )

    patches: list[CellPatch] = []
    for cell, draft in sorted(drafts.items()):
        # TextRunPatch validation requires sorted non-overlapping runs.
        deduped = {
            (run.start, run.length): run
            for run in sorted(draft.text_runs, key=lambda item: (item.start, item.length))
        }
        runs = tuple(sorted(deduped.values(), key=lambda item: (item.start, item.length)))
        patches.append(
            CellPatch(
                sheet="THERAPIST_DAILY",
                cell=cell,
                value=draft.text,
                wrap_text=True,
                min_font_size=10,
                fill_role=draft.fill_role,
                border_role=draft.border_role,
                text_runs=runs,
                source_tag="operational_overlay",
            )
        )

    return OperationalOverlayPlan(
        write_plan=build_write_plan(workbook_path, patches),
        bindings=tuple(bindings),
    )
