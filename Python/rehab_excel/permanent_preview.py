from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, time
from hashlib import sha256
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook


class PermanentPreviewError(RuntimeError):
    pass


def parse_hhmm(value: str) -> time:
    try:
        return datetime.strptime(value.strip(), "%H:%M").time()
    except ValueError as exc:
        raise PermanentPreviewError(f"Invalid time {value!r}; expected HH:MM") from exc


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _norm(value) -> str:
    return "" if value is None else str(value).strip()


def _time_from_cell(value) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        minutes = round((float(value) % 1) * 24 * 60)
        return time((minutes // 60) % 24, minutes % 60)
    text = str(value).strip().replace(".", ":")
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            pass
    return None


@dataclass(frozen=True)
class PlannerTarget:
    row: int
    patient_id: str
    patient_name: str
    day_pattern: str
    treatment: str
    time_column: int
    days_column: int
    therapist_column: int


def locate_physio_assignment(
    workbook_path: str | Path,
    *,
    patient_name: str,
    from_provider: str,
    from_time: time,
) -> PlannerTarget:
    path = Path(workbook_path)
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    try:
        if "PATIENT_PLANNER" not in wb.sheetnames:
            raise PermanentPreviewError("PATIENT_PLANNER sheet not found")
        ws = wb["PATIENT_PLANNER"]
        headers = {
            _norm(cell.value): index
            for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1)
            if _norm(cell.value)
        }
        required = ("PatientID", "Ασθενής", "ΦΘ_Ώρα", "ΦΘ_Ημέρες", "ΦΘ_Θεραπευτής")
        missing = [name for name in required if name not in headers]
        if missing:
            raise PermanentPreviewError(
                "PATIENT_PLANNER headers missing: " + ", ".join(missing)
            )

        matches = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            name = _norm(row[headers["Ασθενής"] - 1])
            provider = _norm(row[headers["ΦΘ_Θεραπευτής"] - 1])
            start = _time_from_cell(row[headers["ΦΘ_Ώρα"] - 1])
            if (
                name.casefold() == patient_name.strip().casefold()
                and provider.casefold() == from_provider.strip().casefold()
                and start == from_time
            ):
                day_pattern = _norm(row[headers["ΦΘ_Ημέρες"] - 1]) or "Καθ/να"
                matches.append(
                    PlannerTarget(
                        row=row_idx,
                        patient_id=_norm(row[headers["PatientID"] - 1]),
                        patient_name=name,
                        day_pattern=day_pattern,
                        treatment="ΦΘ",
                        time_column=headers["ΦΘ_Ώρα"],
                        days_column=headers["ΦΘ_Ημέρες"],
                        therapist_column=headers["ΦΘ_Θεραπευτής"],
                    )
                )

        if not matches:
            raise PermanentPreviewError(
                f"No PATIENT_PLANNER ΦΘ assignment matched "
                f"{patient_name} / {from_provider} / {from_time.strftime('%H:%M')}"
            )
        if len(matches) != 1:
            raise PermanentPreviewError(
                f"Expected one matching ΦΘ assignment, found {len(matches)}"
            )
        return matches[0]
    finally:
        wb.close()


def excel_column_name(index: int) -> str:
    if index < 1:
        raise ValueError("Excel column index must be >= 1")
    result = ""
    n = index
    while n:
        n, rem = divmod(n - 1, 26)
        result = chr(65 + rem) + result
    return result


@dataclass(frozen=True)
class PermanentPreviewReport:
    source_path: str
    output_path: str
    row: int
    patient_name: str
    old_provider: str
    old_time: str
    new_provider: str
    new_time: str
    day_pattern: str
    source_unchanged: bool
    vba_preserved: bool


def create_permanent_assignment_preview(
    *,
    source_path: str | Path,
    output_path: str | Path,
    patient_name: str,
    from_provider: str,
    from_time: time,
    to_provider: str,
    to_time: time,
    overwrite: bool = False,
) -> PermanentPreviewReport:
    """Create a NEW .xlsm copy with one PATIENT_PLANNER ΦΘ assignment changed.

    This intentionally never edits the source workbook. Weekly capacity/conflict
    validation should be performed before calling this function.
    """

    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise PermanentPreviewError("Output must be different from source workbook")
    if source.suffix.lower() != ".xlsm" or output.suffix.lower() != ".xlsm":
        raise PermanentPreviewError("Source and output must both be .xlsm")
    if not source.exists():
        raise PermanentPreviewError(f"Source workbook not found: {source}")
    if output.exists() and not overwrite:
        raise PermanentPreviewError(f"Output already exists: {output}")

    target = locate_physio_assignment(
        source,
        patient_name=patient_name,
        from_provider=from_provider,
        from_time=from_time,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    source_hash_before = _sha256(source)
    if output.exists():
        try:
            output.unlink()
        except PermissionError as exc:
            raise PermanentPreviewError(
                f"Preview is open or locked: {output}. Close Excel and try again."
            ) from exc
    shutil.copy2(source, output)

    if sys.platform != "win32":
        output.unlink(missing_ok=True)
        raise PermanentPreviewError("Native preview requires Windows + Microsoft Excel")

    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        output.unlink(missing_ok=True)
        raise PermanentPreviewError("pywin32 is required") from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(
            str(output), UpdateLinks=0, ReadOnly=False
        )
        ws = workbook.Worksheets("PATIENT_PLANNER")

        time_cell = ws.Cells(target.row, target.time_column)
        provider_cell = ws.Cells(target.row, target.therapist_column)

        # Excel stores time as a fraction of a day. The existing cell already
        # carries the workbook's native time number format, so no localized
        # NumberFormat call is needed.
        to_fraction = (to_time.hour * 60 + to_time.minute) / (24 * 60)
        time_cell.Value = to_fraction
        provider_cell.Value = to_provider

        workbook.Save()
    except Exception as exc:
        output.unlink(missing_ok=True)
        raise PermanentPreviewError(f"Excel permanent preview write failed: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass

    source_hash_after = _sha256(source)
    if source_hash_after != source_hash_before:
        output.unlink(missing_ok=True)
        raise PermanentPreviewError("Source workbook changed unexpectedly")

    # .xlsm VBA project presence check
    import zipfile
    with zipfile.ZipFile(output, "r") as zf:
        vba_preserved = "xl/vbaProject.bin" in zf.namelist()
    if not vba_preserved:
        output.unlink(missing_ok=True)
        raise PermanentPreviewError("VBA project was not preserved in preview")

    return PermanentPreviewReport(
        source_path=str(source),
        output_path=str(output),
        row=target.row,
        patient_name=target.patient_name,
        old_provider=from_provider,
        old_time=from_time.strftime("%H:%M"),
        new_provider=to_provider,
        new_time=to_time.strftime("%H:%M"),
        day_pattern=target.day_pattern,
        source_unchanged=True,
        vba_preserved=True,
    )
