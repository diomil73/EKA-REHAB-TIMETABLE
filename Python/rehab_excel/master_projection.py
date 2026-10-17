from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata
from typing import Iterable, Sequence

from openpyxl import load_workbook

from rehab_core.models import Patient, PatientType

from .patient_registry_source import read_patient_registry


class MasterProjectionError(RuntimeError):
    """Raised when MASTER_SCHEDULE cannot be projected safely."""


@dataclass(frozen=True)
class MasterProjectionRow:
    target_row: int
    planner_row: int
    patient_id: str
    room: str
    display_name: str
    infectious: bool
    status: str | None
    responsible_doctor: str | None = None


def _sort_text(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _clinic_key(room: str | None) -> str:
    text = str(room or "").strip().casefold()
    if not text:
        return ""
    if text[0] in {"a", "α"}:
        return "A"
    if text[0] in {"b", "β"}:
        return "B"
    return ""


def _room_number(room: str | None) -> int | None:
    text = str(room or "").strip()
    match = re.search(r"(\d+)", text)
    if match is None:
        return None
    return int(match.group(1))


def _projection_metadata(
    path: str | Path,
    *,
    room_order: Sequence[str] | None,
) -> tuple[tuple[str, ...], dict[str, int], int]:
    """Read all workbook metadata needed for MASTER in one openpyxl session."""

    workbook_path = Path(path)
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if "PATIENTS" not in wb.sheetnames:
            raise MasterProjectionError("Workbook has no PATIENTS sheet")

        patient_ws = wb["PATIENTS"]
        row_by_id: dict[str, int] = {}
        for row in range(2, patient_ws.max_row + 1):
            raw = patient_ws.cell(row, 1).value
            if raw is None:
                continue
            if isinstance(raw, float) and raw.is_integer():
                patient_id = str(int(raw))
            else:
                patient_id = str(raw).strip()
            if not patient_id:
                continue
            if patient_id in row_by_id:
                raise MasterProjectionError(
                    f"Duplicate PatientID {patient_id!r} in PATIENTS"
                )
            row_by_id[patient_id] = row

        if room_order is None:
            if "SETTINGS" not in wb.sheetnames:
                raise MasterProjectionError("Workbook has no SETTINGS sheet")
            settings_ws = wb["SETTINGS"]
            rooms = tuple(
                str(settings_ws.cell(row, 8).value).strip()
                for row in range(2, settings_ws.max_row + 1)
                if settings_ws.cell(row, 8).value is not None
                and str(settings_ws.cell(row, 8).value).strip()
            )
        else:
            rooms = tuple(room_order)

        header_row = 1
        if "MASTER_SCHEDULE" in wb.sheetnames:
            master_ws = wb["MASTER_SCHEDULE"]
            for row in range(1, min(master_ws.max_row, 10) + 1):
                room = str(master_ws.cell(row, 2).value or "").strip().casefold()
                patient = str(master_ws.cell(row, 3).value or "").strip().casefold()
                if room == "θάλαμος".casefold() and patient == "ασθενής".casefold():
                    header_row = row
                    break

        return rooms, row_by_id, header_row
    finally:
        wb.close()


def build_master_projection(
    workbook_path: str | Path,
    *,
    patients: Iterable[Patient] | None = None,
    room_order: Sequence[str] | None = None,
) -> tuple[MasterProjectionRow, ...]:
    path = Path(workbook_path)
    patient_list = tuple(patients) if patients is not None else tuple(read_patient_registry(path))
    rooms, row_by_id, master_header_row = _projection_metadata(
        path,
        room_order=room_order,
    )

    room_rank = {_sort_text(room): index for index, room in enumerate(rooms)}

    inpatients = [
        patient
        for patient in patient_list
        if patient.patient_type == PatientType.INPATIENT
    ]

    def sort_key(patient: Patient):
        room = patient.room or ""
        normalized_room = _sort_text(room)
        clinic = _clinic_key(room)
        number = _room_number(room)

        if clinic in {"A", "B"} and number is not None:
            return (
                0 if clinic == "A" else 1,
                number,
                normalized_room,
                _sort_text(patient.display_name),
                _sort_text(patient.patient_id),
            )

        return (
            2,
            room_rank.get(normalized_room, len(room_rank)),
            normalized_room,
            _sort_text(patient.display_name),
            _sort_text(patient.patient_id),
        )

    ordered = sorted(inpatients, key=sort_key)
    projection: list[MasterProjectionRow] = []
    next_target_row = master_header_row + 1
    previous_clinic = ""

    for patient in ordered:
        clinic = _clinic_key(patient.room)
        if clinic and clinic != previous_clinic:
            next_target_row += 1
            previous_clinic = clinic

        planner_row = row_by_id.get(patient.patient_id)
        if planner_row is None:
            raise MasterProjectionError(
                f"PatientID {patient.patient_id!r} has no PATIENTS row"
            )
        projection.append(
            MasterProjectionRow(
                target_row=next_target_row,
                planner_row=planner_row,
                patient_id=patient.patient_id,
                room=patient.room or "",
                display_name=patient.display_name,
                infectious=patient.infectious,
                status=patient.status,
                responsible_doctor=patient.responsible_doctor,
            )
        )
        next_target_row += 1

    return tuple(projection)


_PLANNER_REF_RE = re.compile(r"(PATIENT_PLANNER!\$?[A-Z]+\$?)(\d+)", re.IGNORECASE)


def remap_master_formula(formula: str, planner_row: int) -> str:
    text = str(formula or "")
    if not text.startswith("="):
        raise MasterProjectionError(f"MASTER template cell is not a formula: {text!r}")
    if planner_row < 2:
        raise MasterProjectionError(f"Invalid PATIENT_PLANNER row: {planner_row}")

    replaced, count = _PLANNER_REF_RE.subn(
        lambda match: f"{match.group(1)}{planner_row}",
        text,
    )
    if count < 1:
        raise MasterProjectionError(
            f"MASTER formula has no PATIENT_PLANNER reference: {text!r}"
        )
    return replaced
