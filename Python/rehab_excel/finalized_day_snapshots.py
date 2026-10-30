from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sys
from typing import Iterable

from openpyxl import load_workbook

from rehab_core.day_finalization import FinalizedDaySnapshot, FinalizedSnapshotStatus


SHEET_NAME = "FINALIZED_DAY_SNAPSHOTS"
HEADERS = (
    "SnapshotId",
    "ServiceDate",
    "Version",
    "Status",
    "FinalizedAt",
    "SupersedesSnapshotId",
    "UnresolvedSessions",
    "PlannedSessionCount",
    "FinalSessionCount",
    "PatientAbsenceCount",
    "CancellationCount",
    "NoShowCount",
    "ReplacementCount",
    "TreatmentSuspensionCount",
    "ApprovedExceptionCount",
)


class FinalizedDaySnapshotStoreError(RuntimeError):
    """Raised when finalized-day history storage is invalid or unsafe."""


def _parse_date(value: object, *, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: invalid ServiceDate {value!r}"
        ) from exc


def _parse_datetime(value: object, *, row: int) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value or "").strip()
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError as exc:
            raise FinalizedDaySnapshotStoreError(
                f"{SHEET_NAME} row {row}: invalid FinalizedAt {value!r}"
            ) from exc
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: FinalizedAt must include a timezone offset"
        )
    return parsed


def _required_int(value: object, *, row: int, field: str) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: invalid {field} {value!r}"
        ) from exc
    if result < 0:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: {field} cannot be negative"
        )
    return result


def _optional_int(value: object, *, row: int, field: str) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    return _required_int(value, row=row, field=field)


def _parse_status(value: object, *, row: int) -> FinalizedSnapshotStatus:
    text = str(value or "").strip().upper()
    try:
        return FinalizedSnapshotStatus(text)
    except ValueError as exc:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: unsupported Status {value!r}"
        ) from exc


def _snapshot_from_values(values: tuple[object, ...], *, row: int) -> FinalizedDaySnapshot:
    snapshot_id = str(values[0] or "").strip()
    if not snapshot_id:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: SnapshotId is blank"
        )
    version = _required_int(values[2], row=row, field="Version")
    if version < 1:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: Version must be at least 1"
        )
    snapshot = FinalizedDaySnapshot(
        snapshot_id=snapshot_id,
        service_date=_parse_date(values[1], row=row),
        version=version,
        status=_parse_status(values[3], row=row),
        finalized_at=_parse_datetime(values[4], row=row),
        supersedes_snapshot_id=str(values[5] or "").strip() or None,
        unresolved_sessions=_required_int(values[6], row=row, field="UnresolvedSessions"),
        planned_session_count=_optional_int(values[7], row=row, field="PlannedSessionCount"),
        final_session_count=_required_int(values[8], row=row, field="FinalSessionCount"),
        patient_absence_count=_required_int(values[9], row=row, field="PatientAbsenceCount"),
        cancellation_count=_required_int(values[10], row=row, field="CancellationCount"),
        no_show_count=_required_int(values[11], row=row, field="NoShowCount"),
        replacement_count=_required_int(values[12], row=row, field="ReplacementCount"),
        treatment_suspension_count=_required_int(values[13], row=row, field="TreatmentSuspensionCount"),
        approved_exception_count=_required_int(values[14], row=row, field="ApprovedExceptionCount"),
    )
    if snapshot.unresolved_sessions != 0:
        raise FinalizedDaySnapshotStoreError(
            f"{SHEET_NAME} row {row}: finalized snapshot has unresolved sessions"
        )
    return snapshot


def read_finalized_day_snapshots(
    workbook_path: str | Path,
    *,
    service_date: date | None = None,
    current_only: bool = False,
) -> tuple[FinalizedDaySnapshot, ...]:
    """Read durable finalized snapshot headers without mutating the workbook."""

    path = Path(workbook_path)
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if SHEET_NAME not in wb.sheetnames:
            return ()
        ws = wb[SHEET_NAME]
        actual_headers = tuple(
            str(ws.cell(1, col).value or "").strip()
            for col in range(1, len(HEADERS) + 1)
        )
        if actual_headers != HEADERS:
            raise FinalizedDaySnapshotStoreError(
                f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
            )

        result: list[FinalizedDaySnapshot] = []
        seen_ids: set[str] = set()
        current_by_date: dict[date, str] = {}
        for row_number, row_values in enumerate(
            ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS), values_only=True),
            start=2,
        ):
            values = tuple(row_values)
            if not any(value is not None and str(value).strip() for value in values):
                continue
            snapshot = _snapshot_from_values(values, row=row_number)
            if snapshot.snapshot_id in seen_ids:
                raise FinalizedDaySnapshotStoreError(
                    f"Duplicate SnapshotId {snapshot.snapshot_id!r}"
                )
            seen_ids.add(snapshot.snapshot_id)
            if snapshot.status == FinalizedSnapshotStatus.CURRENT_FINAL:
                prior_id = current_by_date.get(snapshot.service_date)
                if prior_id is not None:
                    raise FinalizedDaySnapshotStoreError(
                        "Multiple CURRENT_FINAL snapshots for "
                        f"{snapshot.service_date.isoformat()}: {prior_id!r}, "
                        f"{snapshot.snapshot_id!r}"
                    )
                current_by_date[snapshot.service_date] = snapshot.snapshot_id
            if service_date is not None and snapshot.service_date != service_date:
                continue
            if current_only and snapshot.status != FinalizedSnapshotStatus.CURRENT_FINAL:
                continue
            result.append(snapshot)
        return tuple(result)
    finally:
        wb.close()


def _semantic_key(snapshot: FinalizedDaySnapshot) -> tuple[object, ...]:
    return (
        snapshot.snapshot_id,
        snapshot.service_date,
        snapshot.version,
        snapshot.status,
        snapshot.finalized_at.isoformat(),
        snapshot.supersedes_snapshot_id,
        snapshot.unresolved_sessions,
        snapshot.planned_session_count,
        snapshot.final_session_count,
        snapshot.patient_absence_count,
        snapshot.cancellation_count,
        snapshot.no_show_count,
        snapshot.replacement_count,
        snapshot.treatment_suspension_count,
        snapshot.approved_exception_count,
    )


def _validate_snapshot_for_append(snapshot: FinalizedDaySnapshot) -> None:
    if not snapshot.snapshot_id.strip():
        raise FinalizedDaySnapshotStoreError("SnapshotId cannot be blank")
    if snapshot.version < 1:
        raise FinalizedDaySnapshotStoreError("Snapshot version must be at least 1")
    if snapshot.status != FinalizedSnapshotStatus.CURRENT_FINAL:
        raise FinalizedDaySnapshotStoreError(
            "New finalized snapshot append must have CURRENT_FINAL status"
        )
    if snapshot.unresolved_sessions != 0:
        raise FinalizedDaySnapshotStoreError(
            "Cannot persist a finalized snapshot with unresolved sessions"
        )
    if snapshot.finalized_at.tzinfo is None or snapshot.finalized_at.utcoffset() is None:
        raise FinalizedDaySnapshotStoreError(
            "FinalizedAt must include a timezone offset"
        )


def _values_for_com(snapshot: FinalizedDaySnapshot) -> tuple[object, ...]:
    return (
        snapshot.snapshot_id,
        snapshot.service_date.isoformat(),
        snapshot.version,
        snapshot.status.value,
        snapshot.finalized_at.isoformat(),
        snapshot.supersedes_snapshot_id or "",
        snapshot.unresolved_sessions,
        "" if snapshot.planned_session_count is None else snapshot.planned_session_count,
        snapshot.final_session_count,
        snapshot.patient_absence_count,
        snapshot.cancellation_count,
        snapshot.no_show_count,
        snapshot.replacement_count,
        snapshot.treatment_suspension_count,
        snapshot.approved_exception_count,
    )


def append_finalized_day_snapshot(
    workbook_path: str | Path,
    snapshot: FinalizedDaySnapshot,
) -> bool:
    """Append one CURRENT_FINAL snapshot using native Excel.

    This operation is idempotent for the exact same SnapshotId + content. It
    refuses a second CURRENT_FINAL for the same service date. Amendment/version
    chain transitions are intentionally a separate operation so ordinary day
    close can never silently rewrite history.
    """

    _validate_snapshot_for_append(snapshot)
    if sys.platform != "win32":
        raise FinalizedDaySnapshotStoreError(
            "Finalized snapshot persistence requires Windows + Microsoft Excel"
        )
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise FinalizedDaySnapshotStoreError("pywin32 is required") from exc

    path = Path(workbook_path).resolve()
    if not path.exists():
        raise FinalizedDaySnapshotStoreError(f"Workbook not found: {path}")

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)

        ws = None
        for sheet in workbook.Worksheets:
            if str(sheet.Name).casefold() == SHEET_NAME.casefold():
                ws = sheet
                break
        if ws is None:
            ws = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
            ws.Name = SHEET_NAME
            for col, header in enumerate(HEADERS, start=1):
                ws.Cells(1, col).Value = header
            ws.Rows(1).Font.Bold = True
        else:
            try:
                ws.Unprotect()
            except Exception:
                pass
            actual_headers = tuple(
                str(ws.Cells(1, col).Value or "").strip()
                for col in range(1, len(HEADERS) + 1)
            )
            if actual_headers != HEADERS:
                raise FinalizedDaySnapshotStoreError(
                    f"{SHEET_NAME} headers do not match authoritative schema: {actual_headers!r}"
                )

        last_row = max(int(ws.UsedRange.Rows.Count), 1)
        existing_id: FinalizedDaySnapshot | None = None
        current_for_date: FinalizedDaySnapshot | None = None
        for row_number in range(2, last_row + 1):
            values = tuple(ws.Cells(row_number, col).Value for col in range(1, len(HEADERS) + 1))
            if not any(value is not None and str(value).strip() for value in values):
                continue
            existing = _snapshot_from_values(values, row=row_number)
            if existing.snapshot_id == snapshot.snapshot_id:
                existing_id = existing
            if (
                existing.service_date == snapshot.service_date
                and existing.status == FinalizedSnapshotStatus.CURRENT_FINAL
            ):
                current_for_date = existing

        if existing_id is not None:
            if _semantic_key(existing_id) == _semantic_key(snapshot):
                ws.Visible = 2
                ws.Protect()
                return False
            raise FinalizedDaySnapshotStoreError(
                f"SnapshotId {snapshot.snapshot_id!r} already exists with different content"
            )
        if current_for_date is not None:
            raise FinalizedDaySnapshotStoreError(
                "Service date already has CURRENT_FINAL snapshot "
                f"{current_for_date.snapshot_id!r}; use the explicit amendment workflow"
            )

        next_row = max(last_row + 1, 2)
        values = _values_for_com(snapshot)
        for col, value in enumerate(values, start=1):
            ws.Cells(next_row, col).Value2 = value

        readback = tuple(ws.Cells(next_row, col).Value for col in range(1, len(HEADERS) + 1))
        persisted = _snapshot_from_values(readback, row=next_row)
        if _semantic_key(persisted) != _semantic_key(snapshot):
            raise FinalizedDaySnapshotStoreError(
                "Finalized snapshot COM read-back did not match requested snapshot"
            )

        ws.Visible = 2
        ws.Protect()
        workbook.Save()
        return True
    except FinalizedDaySnapshotStoreError:
        raise
    except Exception as exc:
        raise FinalizedDaySnapshotStoreError(
            f"Could not persist finalized day snapshot: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
