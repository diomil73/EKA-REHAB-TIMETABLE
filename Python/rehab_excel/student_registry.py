from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from rehab_core.models import Student


STUDENT_REGISTRY_SHEET = "STUDENTS"
STUDENT_REGISTRY_HEADERS = (
    "StudentID",
    "Φοιτητής",
    "StudentNumber",
    "PlacementStart",
    "PlacementEnd",
    "SupervisorTherapist",
    "ReplacementCapable",
    "RoboticCapable",
)


class StudentRegistryError(ValueError):
    """Raised when the authoritative STUDENTS registry is malformed."""


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _as_date(value: object, *, field: str, row: int, epoch) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch=epoch)
        except (TypeError, ValueError, OverflowError) as exc:
            raise StudentRegistryError(
                f"STUDENTS row {row}: invalid {field} Excel date {value!r}"
            ) from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
        raise StudentRegistryError(
            f"STUDENTS row {row}: invalid {field} Excel date {value!r}"
        )
    text = _clean(value)
    if text is None:
        raise StudentRegistryError(f"STUDENTS row {row}: {field} is required")
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise StudentRegistryError(
        f"STUDENTS row {row}: invalid {field} date {value!r}"
    )


def _as_positive_int(value: object, *, field: str, row: int) -> int:
    if isinstance(value, bool):
        raise StudentRegistryError(f"STUDENTS row {row}: {field} must be a positive integer")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise StudentRegistryError(
            f"STUDENTS row {row}: {field} must be a positive integer"
        ) from exc
    if number < 1:
        raise StudentRegistryError(f"STUDENTS row {row}: {field} must be positive")
    return number


def _as_bool(value: object, *, default: bool, field: str, row: int) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    text = str(value).strip().casefold()
    if text in {"ναι", "ν", "yes", "y", "true", "1"}:
        return True
    if text in {"οχι", "όχι", "oχι", "no", "n", "false", "0"}:
        return False
    raise StudentRegistryError(
        f"STUDENTS row {row}: {field} must be a yes/no value"
    )


def _header_tuple(ws) -> tuple[str, ...]:
    values = next(
        ws.iter_rows(
            min_row=1,
            max_row=1,
            min_col=1,
            max_col=len(STUDENT_REGISTRY_HEADERS),
            values_only=True,
        )
    )
    return tuple(str(value).strip() if value is not None else "" for value in values)


def validate_student_registry_headers(headers: Iterable[object]) -> None:
    actual = tuple(str(value).strip() if value is not None else "" for value in headers)
    if actual != STUDENT_REGISTRY_HEADERS:
        raise StudentRegistryError(
            "STUDENTS headers do not match the authoritative schema: "
            + ", ".join(STUDENT_REGISTRY_HEADERS)
        )


def read_students(path: str | Path) -> list[Student]:
    """Read the authoritative student registry.

    Old workbooks remain compatible: when ``STUDENTS`` does not exist yet,
    an empty list is returned. Once the sheet exists its first eight columns
    must match ``STUDENT_REGISTRY_HEADERS`` exactly.
    """

    workbook_path = Path(path)
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if STUDENT_REGISTRY_SHEET not in wb.sheetnames:
            return []

        ws = wb[STUDENT_REGISTRY_SHEET]
        headers = _header_tuple(ws)
        validate_student_registry_headers(headers)

        students: list[Student] = []
        seen_ids: set[str] = set()
        seen_numbers: set[int] = set()

        for excel_row, values in enumerate(
            ws.iter_rows(
                min_row=2,
                min_col=1,
                max_col=len(STUDENT_REGISTRY_HEADERS),
                values_only=True,
            ),
            start=2,
        ):
            if all(value is None or str(value).strip() == "" for value in values):
                continue

            student_id = _clean(values[0])
            display_name = _clean(values[1])
            if student_id is None:
                raise StudentRegistryError(f"STUDENTS row {excel_row}: StudentID is required")
            if display_name is None:
                raise StudentRegistryError(f"STUDENTS row {excel_row}: Φοιτητής is required")

            student_number = _as_positive_int(
                values[2], field="StudentNumber", row=excel_row
            )
            placement_start = _as_date(
                values[3], field="PlacementStart", row=excel_row, epoch=wb.epoch
            )
            placement_end = _as_date(
                values[4], field="PlacementEnd", row=excel_row, epoch=wb.epoch
            )
            if placement_end < placement_start:
                raise StudentRegistryError(
                    f"STUDENTS row {excel_row}: PlacementEnd cannot be before PlacementStart"
                )

            id_key = student_id.casefold()
            if id_key in seen_ids:
                raise StudentRegistryError(
                    f"STUDENTS row {excel_row}: duplicate StudentID {student_id!r}"
                )
            if student_number in seen_numbers:
                raise StudentRegistryError(
                    f"STUDENTS row {excel_row}: duplicate StudentNumber {student_number}"
                )
            seen_ids.add(id_key)
            seen_numbers.add(student_number)

            students.append(
                Student(
                    student_id=student_id,
                    display_name=display_name,
                    student_number=student_number,
                    placement_start=placement_start,
                    placement_end=placement_end,
                    supervisor_therapist_id=_clean(values[5]),
                    replacement_capable=_as_bool(
                        values[6],
                        default=True,
                        field="ReplacementCapable",
                        row=excel_row,
                    ),
                    robotic_capable=_as_bool(
                        values[7],
                        default=False,
                        field="RoboticCapable",
                        row=excel_row,
                    ),
                )
            )

        return students
    finally:
        wb.close()
