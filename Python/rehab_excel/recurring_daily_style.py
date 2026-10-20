from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import unicodedata

from openpyxl import load_workbook

from rehab_core.models import BaseScheduleEntry, Patient

from .layout_contract import LayoutSafetyError, resolve_therapist_daily_cell
from .writeback import CellPatch, TextRunPatch, WriteIntent, WritePlan, build_write_plan


class RecurringDailyStyleError(RuntimeError):
    """Raised when recurring THERAPIST_DAILY styling cannot be mapped safely."""


def _key(value: object) -> str:
    text = " ".join(str(value or "").strip().casefold().split())
    decomposed = unicodedata.normalize("NFD", text)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _line_starts(text: str) -> list[tuple[str, int]]:
    lines = text.split("\n")
    result: list[tuple[str, int]] = []
    cursor = 1
    for index, line in enumerate(lines):
        result.append((line, cursor))
        cursor += len(line)
        if index < len(lines) - 1:
            cursor += 1
    return result


def _roles_for_lines(
    text: str,
    entries: list[BaseScheduleEntry],
    patients: dict[str, Patient],
) -> tuple[TextRunPatch, ...]:
    """Map existing recurring display lines back to authoritative assignments.

    The workbook text is preserved verbatim. A line is matched by patient name
    plus day pattern when available, which disambiguates same-patient split
    schedules such as ΦΘ on Τρ-Πε and Ρομποτικό on Δε-Τε-Πα. Outpatients use
    neutral black text; their outpatient identity is communicated by the
    dedicated cell presentation layer. Robotic assignments keep orange
    precedence.
    """

    runs: list[TextRunPatch] = []
    for line, start in _line_starts(text):
        if not line:
            continue
        line_key = _key(line)
        matches: list[BaseScheduleEntry] = []
        for entry in entries:
            patient = patients.get(entry.patient_id)
            if patient is None:
                continue
            if _key(patient.display_name) not in line_key:
                continue
            if _key(entry.day_pattern) and _key(entry.day_pattern) not in line_key:
                continue
            matches.append(entry)

        if not matches:
            # Some single-assignment legacy cells omit the day pattern. Only
            # fall back to patient-name matching when that is unambiguous.
            name_matches = [
                entry
                for entry in entries
                if (patient := patients.get(entry.patient_id)) is not None
                and _key(patient.display_name) in line_key
            ]
            if len(name_matches) == 1:
                matches = name_matches

        if len(matches) != 1:
            continue
        entry = matches[0]
        patient = patients.get(entry.patient_id)
        robotic = bool(entry.robotic) or _key(entry.treatment) == _key("Ρομποτικό")
        if robotic:
            font_role = "robotic_orange"
        else:
            # Both ordinary and outpatient recurring lines start from neutral
            # black. Outpatient semantics are carried by the dedicated fill,
            # avoiding stale inherited green/orange font styling after edits.
            font_role = "default"
        runs.append(
            TextRunPatch(
                start=start,
                length=len(line),
                font_role=font_role,
            )
        )
    return tuple(runs)


def build_recurring_therapist_daily_style_plan(
    workbook_path: str | Path,
    entries: list[BaseScheduleEntry],
    patients: list[Patient],
) -> WritePlan:
    """Style recurring THERAPIST_DAILY lines without changing their text.

    Cells containing robotic or outpatient recurring assignments are targeted.
    Existing cell text, line order, fills, borders and layout remain untouched;
    rich-text font roles are refreshed from PATIENT_PLANNER + PATIENTS data.
    """

    workbook_path = Path(workbook_path)
    patient_map = {patient.patient_id: patient for patient in patients}
    grouped: dict[tuple[str, object], list[BaseScheduleEntry]] = defaultdict(list)
    for entry in entries:
        if entry.therapist_id is None:
            continue
        grouped[(entry.therapist_id, entry.start_time)].append(entry)

    wb = load_workbook(workbook_path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "THERAPIST_DAILY" not in wb.sheetnames:
            raise RecurringDailyStyleError("Workbook does not contain THERAPIST_DAILY")
        ws = wb["THERAPIST_DAILY"]
        patches: list[CellPatch] = []
        for (therapist_id, start_time), slot_entries in grouped.items():
            has_special_style = any(
                bool(entry.robotic)
                or _key(entry.treatment) == _key("Ρομποτικό")
                or (
                    (patient := patient_map.get(entry.patient_id)) is not None
                    and patient.is_outpatient
                )
                for entry in slot_entries
            )
            if not has_special_style:
                continue
            try:
                cell_address = resolve_therapist_daily_cell(
                    workbook_path,
                    therapist_id,
                    start_time,
                )
            except LayoutSafetyError as exc:
                raise RecurringDailyStyleError(str(exc)) from exc

            value = ws[cell_address].value
            if value in (None, ""):
                continue
            text = str(value)
            runs = _roles_for_lines(text, slot_entries, patient_map)
            if not runs:
                continue
            patches.append(
                CellPatch(
                    sheet="THERAPIST_DAILY",
                    cell=cell_address,
                    value=text,
                    intent=WriteIntent.PRESENTATION,
                    text_runs=runs,
                    source_tag="recurring_daily_style",
                )
            )
    finally:
        wb.close()

    return build_write_plan(workbook_path, patches)
