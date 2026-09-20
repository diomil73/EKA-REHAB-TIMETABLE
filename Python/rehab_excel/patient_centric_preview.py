from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date, time
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from rehab_core.base_schedule import pattern_applies_on_date
from rehab_core.daily_state import DailySessionState, DailySessionStatus
from rehab_core.models import BaseScheduleEntry, Patient, ReplacementProviderKind, Session, Student
from rehab_core.students import student_display_state

from .layout_contract import LayoutSafetyError, resolve_therapist_daily_cell
from .writeback import CellPatch, TextRunPatch, WritePlan, build_write_plan


class PatientCentricPreviewError(RuntimeError):
    """Raised when a patient-centric preview cannot be built without guessing."""


@dataclass(frozen=True)
class VisualLine:
    text: str
    patient_id: str | None = None
    strike_through: bool = False
    italic: bool = False
    font_role: str | None = None


@dataclass(frozen=True)
class SlotRender:
    therapist_id: str
    start_time: time
    lines: tuple[VisualLine, ...]

    @property
    def text(self) -> str:
        return "\n".join(line.text for line in self.lines)


@dataclass(frozen=True)
class PreviewBinding:
    session_id: str
    patient_id: str
    status: DailySessionStatus
    original_cell: str
    effective_cell: str | None


@dataclass(frozen=True)
class PatientCentricPreviewPlan:
    write_plan: WritePlan
    bindings: tuple[PreviewBinding, ...]
    rebuilt_group_cells: tuple[str, ...]


_TREATMENT_LABELS = {
    "Ρομποτικό": "ΡΟΜΠ",
    "ΡΟΜΠΟΤΙΚΟ": "ΡΟΜΠ",
    "ΦΘ": "ΦΘ",
    "Ανακλινόμενο": "ΑΝΑΚ",
    "Πισίνα": "ΠΙΣ",
    "Εργο": "ΕΡΓΟ",
    "Λογο": "ΛΟΓΟ",
    "ΕΦΑ": "ΕΦΑ",
}

_WEEKDAY_LABELS = ("Δε", "Τρ", "Τε", "Πε", "Πα", "Σα", "Κυ")

_INFECTIOUS_RGB = {"FFFFFF43", "FFFF43", "FFFFFF00", "FFFF00"}
_ROBOTIC_RGB = {"FFF8CBAD", "F8CBAD"}


def _treatment_label(value: str) -> str:
    return _TREATMENT_LABELS.get(value, value.upper())


def _temporary_date_label(target_date: date) -> str:
    return f"{_WEEKDAY_LABELS[target_date.weekday()]} {target_date.strftime('%d/%m')}"


def _group_entries(
    entries: Iterable[BaseScheduleEntry],
) -> dict[tuple[str, time], list[BaseScheduleEntry]]:
    grouped: dict[tuple[str, time], list[BaseScheduleEntry]] = {}
    for entry in entries:
        if entry.therapist_id is None:
            continue
        grouped.setdefault((entry.therapist_id, entry.start_time), []).append(entry)
    return grouped


def _ordered_patient_entries(entries: Iterable[BaseScheduleEntry]) -> list[tuple[str, list[BaseScheduleEntry]]]:
    ordered: list[str] = []
    by_patient: dict[str, list[BaseScheduleEntry]] = {}
    for entry in entries:
        if entry.patient_id not in by_patient:
            ordered.append(entry.patient_id)
            by_patient[entry.patient_id] = []
        by_patient[entry.patient_id].append(entry)
    return [(patient_id, by_patient[patient_id]) for patient_id in ordered]


def render_base_slot(
    therapist_id: str,
    start_time: time,
    entries: Iterable[BaseScheduleEntry],
    patients: Mapping[str, Patient],
) -> SlotRender:
    """Render a therapist/time cell from patient-owned assignments.

    Different patients sharing the slot get their own line with day pattern.
    Multiple assignments belonging to the SAME patient are compacted into one
    patient block: the name appears once and the assignment details appear on
    one following line. This avoids the misleading duplicate-patient display.
    """

    slot_entries = list(entries)
    patient_groups = _ordered_patient_entries(slot_entries)
    distinct_patient_count = len(patient_groups)
    lines: list[VisualLine] = []

    for patient_id, patient_entries in patient_groups:
        patient = patients.get(patient_id)
        if patient is None:
            raise PatientCentricPreviewError(f"Missing patient {patient_id!r}")

        if len(patient_entries) == 1:
            entry = patient_entries[0]
            label = patient.display_name
            if distinct_patient_count > 1:
                label += f" [{entry.day_pattern}]"
            lines.append(VisualLine(label, patient_id=patient_id))
            continue

        # Same patient, same therapist/time, multiple recurring assignments.
        # Show the name once and summarize the independent assignments below.
        lines.append(VisualLine(patient.display_name, patient_id=patient_id))
        details = " | ".join(
            f"{entry.day_pattern} {_treatment_label(entry.treatment)}"
            for entry in patient_entries
        )
        lines.append(VisualLine(details, patient_id=patient_id, font_role="muted"))

    return SlotRender(therapist_id, start_time, tuple(lines))


def _line_runs(lines: Iterable[VisualLine]) -> tuple[TextRunPatch, ...]:
    runs: list[TextRunPatch] = []
    cursor = 0
    for index, line in enumerate(lines):
        if line.text and (line.strike_through or line.italic or line.font_role):
            runs.append(
                TextRunPatch(
                    start=cursor + 1,
                    length=len(line.text),
                    strike_through=line.strike_through,
                    italic=line.italic,
                    font_role=line.font_role,
                )
            )
        cursor += len(line.text)
        if index >= 0:
            cursor += 1  # newline between visual lines; harmless after final line
    return tuple(runs)


def _style_patient(
    render: SlotRender,
    patient_id: str,
    *,
    strike: bool = False,
    italic: bool = False,
    font_role: str | None = None,
) -> SlotRender:
    found = False
    updated: list[VisualLine] = []
    for line in render.lines:
        if line.patient_id == patient_id:
            found = True
            updated.append(
                replace(
                    line,
                    strike_through=strike,
                    italic=italic,
                    font_role=font_role,
                )
            )
        else:
            updated.append(line)
    if not found:
        raise PatientCentricPreviewError(
            f"Patient {patient_id!r} is not present in {render.therapist_id}@{render.start_time}"
        )
    return replace(render, lines=tuple(updated))


def _append_line(render: SlotRender, line: VisualLine) -> SlotRender:
    if any(existing.text == line.text for existing in render.lines):
        return render
    return replace(render, lines=render.lines + (line,))


def _rgb(cell) -> str | None:
    color = cell.fill.fgColor
    if color is None or color.type != "rgb":
        return None
    return (color.rgb or "").upper()


def _existing_roles(workbook_path: Path) -> dict[str, tuple[str | None, str | None]]:
    wb = load_workbook(workbook_path, read_only=False, data_only=False, keep_vba=True)
    try:
        ws = wb["THERAPIST_DAILY"]
        roles: dict[str, tuple[str | None, str | None]] = {}
        for row in ws.iter_rows():
            for cell in row:
                rgb = _rgb(cell)
                if rgb in _ROBOTIC_RGB:
                    roles[cell.coordinate] = ("robotic_pink", None)
                elif rgb in _INFECTIOUS_RGB:
                    roles[cell.coordinate] = ("infectious_yellow", "infectious_yellow")
                else:
                    roles[cell.coordinate] = (None, None)
        return roles
    finally:
        wb.close()


def _merge_roles(
    current: tuple[str | None, str | None],
    *,
    robotic: bool,
    infectious: bool,
) -> tuple[str | None, str | None]:
    fill_role, border_role = current
    if robotic:
        if fill_role == "infectious_yellow":
            border_role = "infectious_yellow"
        fill_role = "robotic_pink"
    elif infectious and fill_role is None:
        fill_role = "infectious_yellow"
    if infectious:
        border_role = "infectious_yellow"
    return fill_role, border_role


def _provider_display(
    provider_id: str,
    *,
    provider_kind: ReplacementProviderKind,
    provider_labels: Mapping[str, str],
    students: Mapping[str, Student],
    target_date: date,
) -> tuple[str, str | None]:
    if provider_kind == ReplacementProviderKind.STUDENT:
        student = students.get(provider_id)
        if student is not None:
            state = student_display_state(student, target_date)
            return state.label, "student_active_green" if state.use_green_font else None
    return provider_labels.get(provider_id, provider_id), None


def _provider_cell(
    workbook_path: Path,
    provider_id: str,
    slot_time: time,
    *,
    provider_kind: ReplacementProviderKind,
    provider_labels: Mapping[str, str],
    students: Mapping[str, Student],
    target_date: date,
) -> str:
    candidates: list[str] = []
    if provider_kind == ReplacementProviderKind.STUDENT:
        student = students.get(provider_id)
        if student is not None:
            state = student_display_state(student, target_date)
            candidates.extend(
                (
                    state.label,
                    f"φοιτ {student.student_number}",
                    f"φοιτητής {student.student_number}",
                )
            )
    candidates.extend((provider_labels.get(provider_id, provider_id), provider_id))

    last_error: Exception | None = None
    seen: set[str] = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        try:
            return resolve_therapist_daily_cell(workbook_path, candidate, slot_time)
        except LayoutSafetyError as exc:
            last_error = exc
    raise PatientCentricPreviewError(str(last_error or f"No cell for provider {provider_id!r}"))


def _base_entry_id_from_session(session_id: str) -> str:
    if "@" not in session_id:
        return session_id
    return session_id.rsplit("@", 1)[0]


def build_patient_centric_preview_plan(
    workbook_path: str | Path,
    *,
    base_entries: Iterable[BaseScheduleEntry],
    states: Iterable[DailySessionState],
    sessions: Iterable[Session],
    patients: Iterable[Patient],
    provider_labels: Mapping[str, str] | None = None,
    students: Iterable[Student] = (),
    rebuild_multi_member_groups: bool = True,
) -> PatientCentricPreviewPlan:
    """Build a patient-centric THERAPIST_DAILY preview.

    This is deliberately a preview-layer experiment. It rebuilds only affected
    cells plus, optionally, existing multi-assignment cells so paired patients
    and same-patient split schedules can be inspected without rewriting the
    whole daily sheet.
    """

    workbook_path = Path(workbook_path)
    entry_list = list(base_entries)
    state_list = [state for state in states if state.status != DailySessionStatus.ACTIVE]
    patient_map = {patient.patient_id: patient for patient in patients}
    student_map = {student.student_id: student for student in students}
    labels = dict(provider_labels or {})
    session_map = {session.session_id: session for session in sessions}
    entry_by_id = {entry.base_entry_id: entry for entry in entry_list}
    grouped = _group_entries(entry_list)

    dates = {state.session_date for state in state_list}
    if len(dates) > 1:
        raise PatientCentricPreviewError("Preview states must belong to one date")
    if not dates:
        raise PatientCentricPreviewError("At least one non-active daily state is required")
    target_date = next(iter(dates))

    renders: dict[tuple[str, time], SlotRender] = {}
    if rebuild_multi_member_groups:
        for key, entries in grouped.items():
            if len(entries) > 1:
                renders[key] = render_base_slot(key[0], key[1], entries, patient_map)

    bindings: list[PreviewBinding] = []
    roles = _existing_roles(workbook_path)
    cell_roles: dict[str, tuple[str | None, str | None]] = {}

    def ensure_render(key: tuple[str, time]) -> SlotRender:
        render = renders.get(key)
        if render is None:
            render = render_base_slot(key[0], key[1], grouped.get(key, ()), patient_map)
            renders[key] = render
        return render

    for state in state_list:
        session = session_map.get(state.session_id)
        if session is None:
            raise PatientCentricPreviewError(f"Missing session {state.session_id!r}")
        patient = patient_map.get(state.patient_id)
        if patient is None:
            raise PatientCentricPreviewError(f"Missing patient {state.patient_id!r}")

        source_key = (state.original_therapist_id, state.original_time)
        source_render = ensure_render(source_key)

        if state.status == DailySessionStatus.REPLACED:
            source_render = _style_patient(
                source_render,
                patient.patient_id,
                italic=True,
                font_role="muted",
            )
            if state.effective_therapist_id is None or state.effective_time is None:
                raise PatientCentricPreviewError("Replacement lacks effective provider/time")
            provider_kind = state.effective_provider_kind or ReplacementProviderKind.THERAPIST
            provider_name, provider_font = _provider_display(
                state.effective_therapist_id,
                provider_kind=provider_kind,
                provider_labels=labels,
                students=student_map,
                target_date=target_date,
            )
            date_label = _temporary_date_label(target_date)
            # THERAPIST_DAILY is already a dated operational view.  Repeating the
            # date on the source replacement note makes narrow cells wrap into
            # three or four visual lines when several replacements are present.
            # Keep the source overlay compact; retain the date on the destination
            # assignment where it is needed to distinguish a temporary insertion
            # from the recurring base programme.
            source_render = _append_line(
                source_render,
                VisualLine(
                    f"→ {provider_name} {state.effective_time.strftime('%H:%M')}",
                    font_role=provider_font,
                ),
            )
            renders[source_key] = source_render

            destination_key = (state.effective_therapist_id, state.effective_time)
            destination_render = ensure_render(destination_key)
            if destination_key != source_key:
                active_entries = [
                    entry
                    for entry in grouped.get(destination_key, ())
                    if pattern_applies_on_date(entry.day_pattern, target_date)
                ]
                if active_entries:
                    occupied = ", ".join(
                        patient_map[entry.patient_id].display_name
                        for entry in active_entries
                        if entry.patient_id in patient_map
                    ) or "άλλη ενεργή συνεδρία"
                    raise PatientCentricPreviewError(
                        f"Destination {state.effective_therapist_id} "
                        f"{state.effective_time.strftime('%H:%M')} is already active on "
                        f"{target_date.isoformat()} for {occupied}. Choose another timeslot."
                    )
                destination_render = _append_line(
                    destination_render,
                    VisualLine(
                        f"{patient.display_name} [{date_label}]",
                        patient_id=patient.patient_id,
                    ),
                )
                renders[destination_key] = destination_render

            effective_cell = _provider_cell(
                workbook_path,
                state.effective_therapist_id,
                state.effective_time,
                provider_kind=provider_kind,
                provider_labels=labels,
                students=student_map,
                target_date=target_date,
            )
        else:
            source_render = _style_patient(
                source_render,
                patient.patient_id,
                strike=True,
                font_role="muted",
            )
            renders[source_key] = source_render
            effective_cell = None

        original_cell = _provider_cell(
            workbook_path,
            state.original_therapist_id,
            state.original_time,
            provider_kind=ReplacementProviderKind.THERAPIST,
            provider_labels=labels,
            students=student_map,
            target_date=target_date,
        )
        base_id = _base_entry_id_from_session(state.session_id)
        base_entry = entry_by_id.get(base_id)
        if base_entry is not None:
            current = cell_roles.get(original_cell, roles.get(original_cell, (None, None)))
            cell_roles[original_cell] = _merge_roles(
                current,
                robotic=base_entry.robotic,
                infectious=patient.infectious,
            )
        bindings.append(
            PreviewBinding(
                session_id=state.session_id,
                patient_id=state.patient_id,
                status=state.status,
                original_cell=original_cell,
                effective_cell=effective_cell,
            )
        )

    patches: list[CellPatch] = []
    rebuilt_cells: list[str] = []
    for (therapist_id, slot_time), render in sorted(
        renders.items(), key=lambda item: (item[0][1], item[0][0].casefold())
    ):
        cell = _provider_cell(
            workbook_path,
            therapist_id,
            slot_time,
            provider_kind=ReplacementProviderKind.THERAPIST,
            provider_labels=labels,
            students=student_map,
            target_date=target_date,
        )
        if len(grouped.get((therapist_id, slot_time), ())) > 1:
            rebuilt_cells.append(cell)
        fill_role, border_role = cell_roles.get(cell, roles.get(cell, (None, None)))
        patches.append(
            CellPatch(
                sheet="THERAPIST_DAILY",
                cell=cell,
                value=render.text,
                wrap_text=True,
                min_font_size=10,
                fill_role=fill_role,
                border_role=border_role,
                text_runs=_line_runs(render.lines),
                source_tag="patient_centric_preview",
            )
        )

    return PatientCentricPreviewPlan(
        write_plan=build_write_plan(workbook_path, patches),
        bindings=tuple(bindings),
        rebuilt_group_cells=tuple(sorted(set(rebuilt_cells))),
    )
