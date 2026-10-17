from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.models import SessionCancellationKind  # noqa: E402
from rehab_excel.daily_input_reader import read_daily_input  # noqa: E402
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402


STATS_SHEET = "DAILY_STATISTICS"


def read_statistics_values(path: Path) -> dict[str, object]:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if STATS_SHEET not in wb.sheetnames:
            raise RuntimeError(f"Workbook has no {STATS_SHEET} sheet")
        ws = wb[STATS_SHEET]
        result: dict[str, object] = {}
        for row in range(4, ws.max_row + 1):
            key = str(ws.cell(row, 1).value or "").strip()
            if key:
                result[key] = ws.cell(row, 3).value
        return result
    finally:
        wb.close()


def verify_persistence(
    workbook_path: str | Path,
    *,
    expected_suspended: int | None = None,
    expected_unresolved: int | None = None,
    expected_finalized: str | None = None,
) -> tuple[int, dict[str, object]]:
    path = Path(workbook_path).resolve()
    if not path.exists():
        raise FileNotFoundError(path)

    patients = read_patient_registry(path)
    base_entries = read_unified_base_schedule(path)

    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        raw_date = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()

    from datetime import date, datetime

    if isinstance(raw_date, datetime):
        target_date = raw_date.date()
    elif isinstance(raw_date, date):
        target_date = raw_date
    else:
        target_date = datetime.strptime(str(raw_date).strip(), "%d/%m/%Y").date()

    sessions = materialize_sessions_for_date(base_entries, target_date)
    daily = read_daily_input(path, patients=patients, sessions=sessions)
    suspended = sum(
        1
        for cancellation in daily.cancellations
        if cancellation.kind == SessionCancellationKind.TREATMENT_SUSPENDED
    )
    stats = read_statistics_values(path)

    if expected_suspended is not None and suspended != expected_suspended:
        raise RuntimeError(
            f"Suspension reread mismatch: expected {expected_suspended}, got {suspended}"
        )
    if expected_suspended is not None and int(stats.get("treatment_suspended", -1)) != expected_suspended:
        raise RuntimeError(
            "DAILY_STATISTICS treatment_suspended mismatch: "
            f"expected {expected_suspended}, got {stats.get('treatment_suspended')!r}"
        )
    if expected_unresolved is not None and int(stats.get("unresolved", -1)) != expected_unresolved:
        raise RuntimeError(
            f"DAILY_STATISTICS unresolved mismatch: expected {expected_unresolved}, "
            f"got {stats.get('unresolved')!r}"
        )
    if expected_finalized is not None:
        actual = str(stats.get("finalized", "")).strip().upper()
        wanted = str(expected_finalized).strip().upper()
        if actual != wanted:
            raise RuntimeError(
                f"DAILY_STATISTICS finalized mismatch: expected {wanted!r}, got {actual!r}"
            )

    return suspended, stats


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Re-read persisted treatment suspensions from DAILY_INPUT and compare them "
            "with the stored DAILY_STATISTICS sheet."
        )
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--expected-suspended", type=int)
    parser.add_argument("--expected-unresolved", type=int)
    parser.add_argument("--expected-finalized")
    args = parser.parse_args()

    try:
        suspended, stats = verify_persistence(
            args.workbook,
            expected_suspended=args.expected_suspended,
            expected_unresolved=args.expected_unresolved,
            expected_finalized=args.expected_finalized,
        )
    except Exception as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("ORPHAN RESOLUTION PERSISTENCE OK")
    print(f"Workbook: {args.workbook.resolve()}")
    print(f"Treatment suspensions reread: {suspended}")
    print(f"Statistics treatment_suspended: {stats.get('treatment_suspended')}")
    print(f"Statistics unresolved: {stats.get('unresolved')}")
    print(f"Statistics finalized: {stats.get('finalized')}")
    print(
        "NOTE: this verifies suspension persistence and stored statistics only. "
        "Replacement placements are not yet authoritative reread facts."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
