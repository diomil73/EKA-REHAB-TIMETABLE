from __future__ import annotations

import argparse
import sys
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from openpyxl import load_workbook  # noqa: E402

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.models import AbsenceKind, Therapist  # noqa: E402
from rehab_core.therapist_absence_queue import (  # noqa: E402
    build_therapist_absence_replacement_queue,
)
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.reader import read_base_schedule, read_patients, read_settings  # noqa: E402


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _is_legacy_student_name(name: str) -> bool:
    compact = _norm(name).replace(" ", "")
    return compact.startswith("ΦΟΙΤ") or compact.startswith("FOIT")


def _read_daily_input_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputReadError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid DAILY_INPUT date: {raw!r}")


def _fmt_times(slots: tuple[time, ...]) -> str:
    return ", ".join(slot.strftime("%H:%M") for slot in slots)


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Read therapist absences from DAILY_INPUT and show the affected patients "
            "with ranked replacement suggestions. Read-only: no workbook is written."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "DAILY_INPUT_PREVIEW.xlsm",
    )
    parser.add_argument("--limit", type=int, default=5)
    args = parser.parse_args()

    input_book = args.input.resolve()
    if not input_book.exists():
        print(f"SAFETY STOP: input workbook not found: {input_book}")
        return 2

    try:
        target_date = _read_daily_input_date(input_book)
        patients = read_patients(input_book)
        settings = read_settings(input_book)
        base_entries = read_base_schedule(input_book)
        sessions = materialize_sessions_for_date(base_entries, target_date)
        daily = read_daily_input(input_book, patients=patients, sessions=sessions)

        therapist_absences = tuple(
            absence
            for absence in daily.absences
            if absence.absence_kind == AbsenceKind.THERAPIST
        )
        if not therapist_absences:
            raise DailyInputReadError(
                "DAILY_INPUT has no therapist absence rows. Add one therapist row and save the workbook."
            )

        therapists = tuple(
            Therapist(therapist_id=name, display_name=name)
            for name in settings.therapist_names
            if not _is_legacy_student_name(name)
        )
        queue = build_therapist_absence_replacement_queue(
            sessions=sessions,
            absences=daily.absences,
            therapists=therapists,
            patients=patients,
            timeslots=settings.standard_timeslots,
        )
    except (DailyInputReadError, ValueError, KeyError) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("THERAPIST ABSENCE REPLACEMENT PLAN OK")
    print(f"Date: {target_date.isoformat()}")
    print(f"Input: {input_book}")
    print(f"Therapist rows read: {daily.therapist_rows_used}")
    print(f"Patient rows read: {daily.patient_rows_used}")
    print(f"Affected sessions needing replacement: {queue.affected_sessions}")
    if queue.skipped_patient_absent:
        print(
            "Patient-absent sessions excluded from replacement queue: "
            f"{queue.skipped_patient_absent}"
        )
    for warning in daily.warnings:
        print(f"WARNING: {warning}")

    if not queue.items:
        print("No session currently needs a replacement for the declared therapist absence(s).")
        return 0

    limit = max(args.limit, 0)
    for item_index, item in enumerate(queue.items, start=1):
        treatment = f" | {item.treatment}" if item.treatment else ""
        print(
            f"\n{item_index}. {item.patient_name} | "
            f"{item.original_therapist_id} {item.original_time.strftime('%H:%M')}{treatment}"
        )
        if item.robotic:
            print("   ROBOTIC SESSION: capability data must be authoritative before accepting a candidate.")
        if not item.options:
            print("   No feasible replacement provider/timeslot found.")
            continue
        for rank, option in enumerate(item.options[:limit], start=1):
            exact = "exact" if option.exact_time_available else "alternative"
            print(
                f"   {rank}. {option.display_name} | load {option.active_sessions} | "
                f"recommended {option.recommended_time.strftime('%H:%M')} ({exact}) | "
                f"free: {_fmt_times(option.available_timeslots)}"
            )

    print(
        "\nRULE: patient absence wins over therapist absence; only sessions that will "
        "actually take place are placed in the replacement queue."
    )
    print(
        "RULE: candidates are load-first; exact time is preferred at equal load, "
        "otherwise another common free timeslot is suggested."
    )
    print("READ-ONLY: no Excel file was modified.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
