from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Inspect PATIENT_PLANNER headers and identity rows without modifying the workbook."
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--rows", type=int, default=12)
    args = parser.parse_args()

    path = args.workbook.resolve()
    if not path.exists():
        print(f"SAFETY STOP: workbook not found: {path}")
        return 2

    wb_values = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    wb_formulas = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if "PATIENT_PLANNER" not in wb_values.sheetnames:
            print("SAFETY STOP: workbook has no PATIENT_PLANNER sheet")
            return 2

        ws_values = wb_values["PATIENT_PLANNER"]
        ws_formulas = wb_formulas["PATIENT_PLANNER"]

        print("PATIENT_PLANNER HEADERS")
        headers = []
        for col in range(1, ws_values.max_column + 1):
            value = _text(ws_values.cell(1, col).value)
            if value:
                headers.append((col, value))
                print(f"{col}: {value}")

        print()
        print("PATIENT_PLANNER SAMPLE ROWS")
        max_rows = min(ws_values.max_row, max(2, args.rows + 1))
        for row in range(2, max_rows + 1):
            pid_value = ws_values.cell(row, 1).value
            name_value = ws_values.cell(row, 2).value
            pid_formula = ws_formulas.cell(row, 1).value
            name_formula = ws_formulas.cell(row, 2).value
            if not any((_text(pid_value), _text(name_value), _text(pid_formula), _text(name_formula))):
                continue
            print(
                f"row {row} | "
                f"A value={pid_value!r} formula={pid_formula!r} | "
                f"B value={name_value!r} formula={name_formula!r}"
            )

        print()
        print(f"Rows: {ws_values.max_row}")
        print(f"Columns: {ws_values.max_column}")
        print("READ ONLY: workbook was not modified.")
        return 0
    finally:
        wb_values.close()
        wb_formulas.close()


if __name__ == "__main__":
    raise SystemExit(main())
