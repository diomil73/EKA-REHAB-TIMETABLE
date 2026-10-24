from __future__ import annotations

import argparse
import hashlib
import shutil
import sys
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.daily_statistics import calculate_daily_treatment_statistics  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_preview_composer import (  # noqa: E402
    DailyPreviewCompositionError,
    compose_outpatient_daily_plan,
)
from rehab_excel.daily_replacement_facts import (  # noqa: E402
    DailyReplacementFactError,
    read_daily_replacement_facts,
    restore_daily_replacement_fact,
)
from rehab_excel.daily_statistics_writer import (  # noqa: E402
    DailyStatisticsWriteError,
    write_daily_statistics,
)
from rehab_excel.master_daily_overlay import (  # noqa: E402
    MasterDailyOverlayError,
    build_master_daily_overlay_plan,
    compose_master_daily_overlay,
)
from rehab_excel.master_projection_writer import (  # noqa: E402
    MasterProjectionWriteError,
    refresh_master_projection_in_place,
)
from rehab_excel.native_excel import apply_write_plan_to_copy  # noqa: E402
from rehab_excel.outpatient_presentation import OutpatientPresentationError  # noqa: E402
from rehab_excel.outpatient_schedule_source import (  # noqa: E402
    OutpatientScheduleSourceError,
    read_unified_base_schedule,
)
from rehab_excel.patient_centric_preview import (  # noqa: E402
    PatientCentricPreviewError,
    build_patient_centric_preview_plan,
)
from rehab_excel.patient_registry_source import (  # noqa: E402
    PatientRegistrySourceError,
    read_patient_registry,
)


LEDGER_SHEET = "DAILY_CHANGE_LEDGER"
STATUS_RESTORED = "RESTORED"


class AuthoritativeDailyRestoreError(RuntimeError):
    pass


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _has_vba(path: Path) -> bool:
    with ZipFile(path) as archive:
        return "xl/vbaProject.bin" in archive.namelist()


def _read_daily_input_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise AuthoritativeDailyRestoreError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise AuthoritativeDailyRestoreError(f"Invalid DAILY_INPUT date: {raw!r}")


def _mark_ledger_restored(workbook_path: Path, replacement_id: str) -> None:
    if sys.platform != "win32":
        raise AuthoritativeDailyRestoreError("Ledger refresh requires Windows + Microsoft Excel")
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise AuthoritativeDailyRestoreError("pywin32 is required") from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(str(workbook_path.resolve()), UpdateLinks=0, ReadOnly=False)
        try:
            ws = workbook.Worksheets(LEDGER_SHEET)
        except Exception:
            workbook.Save()
            return

        last_row = int(ws.Cells(ws.Rows.Count, 1).End(-4162).Row)
        matches: list[int] = []
        for row in range(2, last_row + 1):
            if str(ws.Cells(row, 1).Value or "").strip() == replacement_id:
                matches.append(row)
        if len(matches) > 1:
            raise AuthoritativeDailyRestoreError(
                f"Duplicate ledger ChangeId {replacement_id!r}"
            )
        if matches:
            row = matches[0]
            ws.Cells(row, 12).Value = STATUS_RESTORED
            ws.Cells(row, 17).Value = datetime.now()
        workbook.Save()
        try:
            excel.Run(f"'{workbook.Name}'!RefreshDailyChangeLedgerView")
            workbook.Save()
        except Exception:
            # The authoritative fact and ledger history are already saved. A
            # missing/blocked render macro must not reverse the operational fact.
            pass
    except AuthoritativeDailyRestoreError:
        raise
    except Exception as exc:
        raise AuthoritativeDailyRestoreError(
            f"Could not update daily change ledger after authoritative restore: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def restore_replacement_and_recompute(
    source_path: str | Path,
    output_path: str | Path,
    *,
    replacement_id: str,
    overwrite: bool = False,
) -> tuple[int, int, bool]:
    """Restore one exact replacement fact and rebuild the effective day.

    Removing a placement while its therapist absence still exists can make that
    session orphan/unresolved again. That is intentional: authoritative restore
    removes the overlay and recomputes truth from persisted daily facts.
    """

    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if not source.exists():
        raise AuthoritativeDailyRestoreError(f"Source workbook not found: {source}")
    if source == output:
        raise AuthoritativeDailyRestoreError("Output must differ from source workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise AuthoritativeDailyRestoreError("Source and output must both be .xlsm")
    if output.exists() and not overwrite:
        raise AuthoritativeDailyRestoreError(f"Output already exists: {output}")

    source_before = _sha256(source)
    stage = output.with_name(f".{output.stem}.restore-stage{output.suffix}")
    if stage.exists():
        stage.unlink()
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, stage)

    try:
        changed = restore_daily_replacement_fact(stage, replacement_id)
        if not changed:
            raise AuthoritativeDailyRestoreError(
                f"ReplacementId {replacement_id!r} is already restored"
            )

        target_date = _read_daily_input_date(stage)
        patients = read_patient_registry(stage)
        base_entries = read_unified_base_schedule(stage)
        sessions = materialize_sessions_for_date(base_entries, target_date)
        daily = read_daily_input(stage, patients=patients, sessions=sessions)
        replacements = read_daily_replacement_facts(
            stage,
            sessions=sessions,
            target_date=target_date,
        )
        states = build_daily_session_states(
            sessions,
            absences=daily.absences,
            replacements=replacements,
            cancellations=daily.cancellations,
            target_date=target_date,
        )
        statistics = calculate_daily_treatment_statistics(states)

        # MASTER must return to recurring truth before surviving daily overlays
        # are reapplied, otherwise a removed replacement can leave a stale cell.
        refresh_master_projection_in_place(stage)

        preview = build_patient_centric_preview_plan(
            stage,
            base_entries=base_entries,
            states=states,
            sessions=sessions,
            patients=patients,
            rebuild_multi_member_groups=True,
        )
        composed_plan = compose_outpatient_daily_plan(
            preview.write_plan,
            states=states,
            patients=patients,
        )
        master_overlay = build_master_daily_overlay_plan(
            stage,
            base_entries=base_entries,
            states=states,
            patients=patients,
        )
        composed_plan = compose_master_daily_overlay(composed_plan, master_overlay)
        report = apply_write_plan_to_copy(
            composed_plan,
            output,
            overwrite=overwrite,
        )
        write_daily_statistics(output, target_date, statistics)
        _mark_ledger_restored(output, replacement_id)

        source_after = _sha256(source)
        if source_before != source_after:
            raise AuthoritativeDailyRestoreError("Source workbook changed during restore preview")
        vba_preserved = _has_vba(source) == _has_vba(output) and _has_vba(output)
        if not report.source_unchanged or not vba_preserved:
            raise AuthoritativeDailyRestoreError(
                "Restore preview failed source/VBA preservation checks"
            )
        return len(replacements), statistics.unresolved_sessions, statistics.is_finalized
    finally:
        try:
            stage.unlink()
        except OSError:
            pass


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Restore one authoritative daily replacement fact and recompute the saved day."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--replacement-id", required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    try:
        replacements, unresolved, finalized = restore_replacement_and_recompute(
            args.input,
            args.output,
            replacement_id=args.replacement_id,
            overwrite=args.overwrite,
        )
    except (
        AuthoritativeDailyRestoreError,
        DailyInputReadError,
        DailyReplacementFactError,
        DailyStatisticsWriteError,
        PatientCentricPreviewError,
        PatientRegistrySourceError,
        OutpatientScheduleSourceError,
        OutpatientPresentationError,
        DailyPreviewCompositionError,
        MasterDailyOverlayError,
        MasterProjectionWriteError,
        ValueError,
        KeyError,
    ) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("AUTHORITATIVE DAILY RESTORE OK")
    print(f"Input: {args.input.resolve()}")
    print(f"Output: {args.output.resolve()}")
    print(f"Restored replacement fact: {args.replacement_id}")
    print(f"Active replacements after restore: {replacements}")
    print(f"Unresolved sessions after recompute: {unresolved}")
    print(f"Finalized after recompute: {'ΝΑΙ' if finalized else 'ΟΧΙ'}")
    print("Source unchanged: True")
    print("VBA preserved: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
