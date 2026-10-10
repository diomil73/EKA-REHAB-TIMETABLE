from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Inspect MASTER_SCHEDULE formulas and identity mapping without modifying the workbook."
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--rows", type=int, default=12)
    args = parser.parse_args()

    path = args.workbook.resolve()
    if not path.exists():
        print(f"SAFETY STOP: workbook not found: {path}")
        return 2

    wb_values = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    wb_formulas = load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if "MASTER_SCHEDULE" not in wb_values.sheetnames:
            print("SAFETY STOP: workbook has no MASTER_SCHEDULE sheet")
            return 2

        ws_values = wb_values["MASTER_SCHEDULE"]
        ws_formulas = wb_formulas["MASTER_SCHEDULE"]

        print("MASTER_SCHEDULE FORMULA SAMPLE")
        print("row | A infectious | B room | C patient | K status")
        max_rows = min(ws_values.max_row, max(2, args.rows + 1))
        for row in range(2, max_rows + 1):
            parts = []
            for col, label in ((1, "A"), (2, "B"), (3, "C"), (11, "K")):
                value = ws_values.cell(row, col).value
                formula = ws_formulas.cell(row, col).value
                parts.append(f"{label} value={value!r} formula={formula!r}")
            if any(_text(ws_values.cell(row, col).value) for col in (1, 2, 3, 11)):
                print(f"row {row} | " + " | ".join(parts))

        print()
        for target_row in (26, 27, 28, 99, 100, 101):
            if target_row > ws_values.max_row:
                continue
            print(f"TARGET ROW {target_row}")
            for col, label in ((1, "A"), (2, "B"), (3, "C"), (4, "D"), (11, "K")):
                print(
                    f"{label}: value={ws_values.cell(target_row, col).value!r} "
                    f"formula={ws_formulas.cell(target_row, col).value!r}"
                )

        print()
        print(f"Rows: {ws_values.max_row}")
        print(f"Columns: {ws_values.max_column}")
        print("READ ONLY: workbook was not modified.")
        return 0
    finally:
        wb_values.close()
        wb_formulas.close()


if __name__ == "__main__":
    raise SystemExit(main())
