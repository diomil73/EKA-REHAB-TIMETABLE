from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import json
from pathlib import Path
import sys
import unicodedata

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_config.provider_policy_store import load_policy_book  # noqa: E402
from rehab_config.service_hours_store import load_service_hours  # noqa: E402
from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.models import AbsenceKind, Therapist  # noqa: E402
from rehab_core.replacement_policy import find_policy_replacement_options  # noqa: E402
from rehab_core.service_replacement_options import build_service_id_by_session, filter_replacement_options_for_service_hours  # noqa: E402
from rehab_core.therapist_absence_queue import build_therapist_absence_replacement_queue  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_replacement_facts import read_daily_replacement_facts  # noqa: E402
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.student_registry import read_students  # noqa: E402

HEADERS = (
    "target_session_id", "patient_name", "original_therapist", "original_time",
    "provider_id", "provider_name", "replacement_time", "load_text", "free_times",
    "is_global_recommendation",
)


class ReplacementDecisionPreparationError(RuntimeError):
    pass


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _is_legacy_student_name(name: str) -> bool:
    compact = _norm(name).replace(" ", "")
    return compact.startswith("ΦΟΙΤ") or compact.startswith("FOIT")


def _target_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputReadError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid DAILY_INPUT date: {raw!r}")


def _load_text(option) -> str:
    if option.capacity_limit is None or option.capacity_remaining is None:
        return str(option.active_timeslots)
    used = option.capacity_limit - option.capacity_remaining
    return f"{used}/{option.capacity_limit}"


def _free_times(option) -> str:
    return ", ".join(slot.strftime("%H:%M") for slot in option.available_timeslots)


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def prepare_decision(workbook_path: str | Path, rows_path: str | Path, *, limit: int = 20) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    rows_output = Path(rows_path).resolve()
    if not source.exists():
        raise ReplacementDecisionPreparationError(f"Workbook not found: {source}")
    if source.suffix.casefold() != ".xlsm":
        raise ReplacementDecisionPreparationError("Workbook must be an .xlsm file")

    source_hash = file_sha256(source)
    target_date = _target_date(source)
    patients = tuple(read_patient_registry(source))
    patient_by_id = {patient.patient_id: patient for patient in patients}
    students = tuple(read_students(source))
    settings = read_settings(source)
    base_entries = tuple(read_unified_base_schedule(source))
    sessions = tuple(materialize_sessions_for_date(base_entries, target_date))
    session_by_id = {session.session_id: session for session in sessions}
    daily = read_daily_input(source, patients=patients, sessions=sessions, therapist_names=settings.therapist_names)
    if not any(item.absence_kind == AbsenceKind.THERAPIST for item in daily.absences):
        raise DailyInputReadError("DAILY_INPUT has no therapist absence rows")

    therapists = tuple(
        Therapist(therapist_id=name, display_name=name)
        for name in settings.therapist_names
        if not _is_legacy_student_name(name)
    )
    policy_book = load_policy_book(REPO_ROOT / "Config" / "provider_policy.json")
    service_hours = load_service_hours(REPO_ROOT / "Config" / "service_hours.json")
    service_id_by_session = build_service_id_by_session(base_entries, target_date)
    existing = tuple(read_daily_replacement_facts(source, sessions=sessions, target_date=target_date))
    queue = build_therapist_absence_replacement_queue(
        sessions=sessions, absences=daily.absences, cancellations=daily.cancellations,
        replacements=existing, therapists=therapists, patients=patients,
        timeslots=settings.standard_timeslots, students=students, policy_book=policy_book,
        base_entries=base_entries,
    )
    targets = tuple(session_by_id[item.session_id] for item in queue.items)
    if not targets:
        rows_output.parent.mkdir(parents=True, exist_ok=True)
        with rows_output.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.DictWriter(handle, fieldnames=HEADERS, delimiter="\t").writeheader()
        return {
            "ok": True,
            "mode": "complete",
            "date": target_date.isoformat(),
            "pending_count": 0,
            "choice_count": 0,
            "source_unchanged": file_sha256(source) == source_hash,
        }

    options_by_session: dict[str, tuple] = {}
    for target in targets:
        options = find_policy_replacement_options(
            target_session=target, therapists=therapists, sessions=sessions, policy_book=policy_book,
            absences=daily.absences, patients=patients, replacements=existing,
            cancellations=daily.cancellations, requested_time=target.start_time,
            timeslots=settings.standard_timeslots, students=students, base_entries=base_entries,
        )
        options_by_session[target.session_id] = tuple(
            filter_replacement_options_for_service_hours(
                options,
                target_session=target,
                hours_book=service_hours,
                explicit_service_id=service_id_by_session.get(target.session_id),
            )
        )

    next_target = min(
        targets,
        key=lambda target: (
            len(options_by_session[target.session_id]),
            target.start_time,
            patient_by_id[target.patient_id].display_name.casefold(),
            target.session_id,
        ),
    )

    # Manual mode must stay responsive. It intentionally does not run the full
    # global assignment solver merely to decorate one option as a recommendation.
    # The automatic "Yes" route still owns the authoritative global solve.
    patient = patient_by_id[next_target.patient_id]
    choices: list[dict[str, str]] = []
    for option in options_by_session[next_target.session_id]:
        for slot in option.available_timeslots:
            slot_text = slot.strftime("%H:%M")
            choices.append({
                "target_session_id": next_target.session_id,
                "patient_name": patient.display_name,
                "original_therapist": next_target.therapist_id,
                "original_time": next_target.start_time.strftime("%H:%M"),
                "provider_id": option.provider_id,
                "provider_name": option.display_name,
                "replacement_time": slot_text,
                "load_text": _load_text(option),
                "free_times": _free_times(option),
                "is_global_recommendation": "0",
            })

    if limit > 0 and len(choices) > limit:
        choices = choices[:limit]

    rows_output.parent.mkdir(parents=True, exist_ok=True)
    with rows_output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(choices)

    if file_sha256(source) != source_hash:
        rows_output.unlink(missing_ok=True)
        raise ReplacementDecisionPreparationError("Workbook changed during replacement decision preparation")

    return {
        "ok": True,
        "mode": "decision_required",
        "date": target_date.isoformat(),
        "pending_count": len(targets),
        "target_session_id": next_target.session_id,
        "patient_name": patient.display_name,
        "original_therapist": next_target.therapist_id,
        "original_time": next_target.start_time.strftime("%H:%M"),
        "normal_option_count": len(options_by_session[next_target.session_id]),
        "choice_count": len(choices),
        "recommended_provider_id": "",
        "recommended_time": "",
        "source_unchanged": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare one operator-controlled therapist replacement decision.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--rows", required=True, type=Path)
    parser.add_argument("--response", type=Path)
    parser.add_argument("--limit", type=int, default=20)
    args = parser.parse_args()
    try:
        result = prepare_decision(args.input, args.rows, limit=max(args.limit, 1))
    except Exception as exc:
        args.rows.unlink(missing_ok=True)
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2
    _write_response(args.response, result)
    print("REPLACEMENT OPERATOR DECISION PREP OK")
    print(f"Mode: {result['mode']}")
    print(f"Pending: {result['pending_count']}")
    print(f"Choices: {result['choice_count']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
