from __future__ import annotations

import argparse
from datetime import date, datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_excel.application_contract import assert_application_contract  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_preview_composer import compose_outpatient_daily_plan  # noqa: E402
from rehab_excel.daily_replacement_facts import read_daily_replacement_facts  # noqa: E402
from rehab_excel.master_daily_overlay import (  # noqa: E402
    build_master_daily_overlay_plan,
    compose_master_daily_overlay,
)
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_centric_preview import build_patient_centric_preview_plan  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.writeback import CellPatch, validate_write_plan  # noqa: E402


class DailyChangeRestorePreparationError(RuntimeError):
    pass


SUPPORTED_SOURCE_CHANGE_TYPES = frozenset({"patient_absence", "session_cancellation"})


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def _target_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyChangeRestorePreparationError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyChangeRestorePreparationError(f"Invalid DAILY_INPUT date: {raw!r}")


def _escape(value: object | None) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("\\", "\\\\")
        .replace("\t", "\\t")
        .replace("\r", "\\r")
        .replace("\n", "\\n")
    )


def _bool_token(value: bool | None) -> str:
    if value is None:
        return ""
    return "1" if value else "0"


def _text_runs_token(patch: CellPatch) -> str:
    return ";".join(
        ":".join(
            (
                str(run.start),
                str(run.length),
                "1" if run.strike_through else "0",
                "1" if run.italic else "0",
                run.font_role or "",
            )
        )
        for run in patch.text_runs
    )


def _patch_line(patch: CellPatch) -> str:
    min_font_size = None if patch.text_runs else patch.min_font_size
    fields = (
        patch.sheet,
        patch.cell,
        patch.intent.value,
        patch.value,
        "1" if patch.strike_through else "0",
        patch.font_role,
        patch.fill_role,
        patch.border_role,
        _bool_token(patch.wrap_text),
        min_font_size,
        patch.source_tag,
        _text_runs_token(patch),
    )
    return "\t".join(["PATCH", *(_escape(value) for value in fields)])


def _source_clear_line(source_row: int) -> str:
    return "\t".join(("SOURCE_CLEAR", "DAILY_INPUT", str(source_row), "A", "E"))


def _restore_marker_line(change_id: str) -> str:
    return "\t".join(("RESTORE_CHANGE", _escape(change_id)))


def prepare_daily_change_restore_payload(
    workbook_path: str | Path,
    payload_path: str | Path,
    *,
    change_ids: tuple[str, ...],
) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    payload = Path(payload_path).resolve()
    requested = tuple(dict.fromkeys(str(value or "").strip() for value in change_ids if str(value or "").strip()))
    if not source.exists():
        raise DailyChangeRestorePreparationError(f"Workbook not found: {source}")
    if source.suffix.casefold() != ".xlsm":
        raise DailyChangeRestorePreparationError("Workbook must be an .xlsm file")
    if not requested:
        raise DailyChangeRestorePreparationError("At least one non-empty change_id is required")

    source_hash = file_sha256(source)
    contract = assert_application_contract(source)
    target_date = _target_date(source)
    patients = read_patient_registry(source)
    settings = read_settings(source)
    base_entries = read_unified_base_schedule(source)
    sessions = materialize_sessions_for_date(base_entries, target_date)

    current_daily = read_daily_input(
        source,
        patients=patients,
        sessions=sessions,
        therapist_names=settings.therapist_names,
    )
    by_id = {change.change_id: change for change in current_daily.source_changes}
    missing = tuple(change_id for change_id in requested if change_id not in by_id)
    if missing:
        raise DailyChangeRestorePreparationError(
            f"Active DAILY_INPUT source change(s) not found: {missing!r}"
        )
    changes = tuple(by_id[change_id] for change_id in requested)
    unsupported = tuple(change for change in changes if change.change_type not in SUPPORTED_SOURCE_CHANGE_TYPES)
    if unsupported:
        if any(change.change_type == "therapist_absence" for change in unsupported):
            raise DailyChangeRestorePreparationError(
                "Therapist-absence restore requires coordinated rollback of related replacement facts and is not enabled by this preparer"
            )
        raise DailyChangeRestorePreparationError(
            f"Unsupported DAILY_INPUT source change type(s): {tuple(change.change_type for change in unsupported)!r}"
        )

    ignore_rows = tuple(sorted({change.source_row for change in changes}))
    restored_daily = read_daily_input(
        source,
        patients=patients,
        sessions=sessions,
        therapist_names=settings.therapist_names,
        ignore_rows=ignore_rows,
        allow_empty=True,
    )
    replacements = read_daily_replacement_facts(
        source,
        sessions=sessions,
        target_date=target_date,
    )
    states = build_daily_session_states(
        sessions,
        absences=restored_daily.absences,
        replacements=replacements,
        cancellations=restored_daily.cancellations,
        target_date=target_date,
    )

    preview = build_patient_centric_preview_plan(
        source,
        base_entries=base_entries,
        states=states,
        sessions=sessions,
        patients=patients,
        rebuild_multi_member_groups=True,
    )
    plan = compose_outpatient_daily_plan(
        preview.write_plan,
        states=states,
        patients=patients,
    )
    master_overlay = build_master_daily_overlay_plan(
        source,
        base_entries=base_entries,
        states=states,
        patients=patients,
    )
    plan = compose_master_daily_overlay(plan, master_overlay)
    validate_write_plan(plan)

    lines = [
        *(_source_clear_line(change.source_row) for change in changes),
        *(_restore_marker_line(change.change_id) for change in changes),
        *(_patch_line(patch) for patch in plan.patches),
    ]
    payload.parent.mkdir(parents=True, exist_ok=True)
    payload.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")

    if file_sha256(source) != source_hash:
        try:
            payload.unlink()
        except OSError:
            pass
        raise DailyChangeRestorePreparationError("Workbook changed during daily-change restore validation")

    touched_sheets = sorted({patch.sheet for patch in plan.patches})
    return {
        "ok": True,
        "payload_path": str(payload),
        "source_sha256_before": source_hash,
        "date": target_date.isoformat(),
        "change_ids": list(requested),
        "change_count": len(changes),
        "source_rows": list(ignore_rows),
        "patch_count": len(plan.patches),
        "touched_sheets": touched_sheets,
        "application_contract_components": len(contract.component_names),
        "source_unchanged": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Prepare an atomic in-place restore payload for reversible DAILY_INPUT source changes."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--payload", required=True, type=Path)
    parser.add_argument("--response", type=Path)
    parser.add_argument("--change-id", action="append", default=[])
    args = parser.parse_args()

    try:
        result = prepare_daily_change_restore_payload(
            args.input,
            args.payload,
            change_ids=tuple(args.change_id),
        )
    except Exception as exc:
        try:
            if args.payload.exists():
                args.payload.unlink()
        except OSError:
            pass
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2

    _write_response(args.response, result)
    print("DAILY CHANGE RESTORE IN-PLACE PAYLOAD VERIFIED")
    print(f"Date: {result['date']}")
    print(f"Changes: {result['change_count']}")
    print(f"Source rows: {', '.join(str(value) for value in result['source_rows'])}")
    print(f"Patches: {result['patch_count']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
