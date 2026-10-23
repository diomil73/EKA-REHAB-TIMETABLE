from __future__ import annotations

import argparse
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from openpyxl import load_workbook  # noqa: E402

from rehab_config.provider_policy_store import load_policy_book  # noqa: E402
from rehab_config.service_hours_store import load_service_hours  # noqa: E402
from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.daily_statistics import calculate_daily_treatment_statistics  # noqa: E402
from rehab_core.global_replacement_end_to_end import (  # noqa: E402
    build_global_replacement_assignments,
)
from rehab_core.models import AbsenceKind, Therapist  # noqa: E402
from rehab_core.reclining_adjacency import build_reclining_adjacency_penalty  # noqa: E402
from rehab_core.service_replacement_options import build_service_id_by_session  # noqa: E402
from rehab_core.therapist_absence_queue import (  # noqa: E402
    build_therapist_absence_replacement_queue,
)
from rehab_excel.application_contract import (  # noqa: E402
    ApplicationContractError,
    assert_application_contract,
)
from rehab_excel.daily_change_ledger_vba import (  # noqa: E402
    DailyChangeLedgerPreviewError,
    install_daily_change_ledger_preview,
)
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_preview_composer import (  # noqa: E402
    DailyPreviewCompositionError,
    compose_outpatient_daily_plan,
)
from rehab_excel.daily_replacement_facts import (  # noqa: E402
    DailyReplacementFactError,
    read_daily_replacement_fact_rows,
    read_daily_replacement_facts,
)
from rehab_excel.daily_statistics_writer import (  # noqa: E402
    DailyStatisticsWriteError,
    write_daily_statistics,
)
from rehab_excel.global_replacement_ledger import (  # noqa: E402
    GlobalReplacementLedgerError,
    build_global_replacement_ledger_records,
)
from rehab_excel.global_replacement_lifecycle_bridge import (  # noqa: E402
    GlobalReplacementLifecycleError,
    persist_global_replacement_result,
)
from rehab_excel.master_daily_overlay import (  # noqa: E402
    MasterDailyOverlayError,
    build_master_daily_overlay_plan,
    compose_master_daily_overlay,
)
from rehab_excel.native_excel import apply_write_plan_to_copy  # noqa: E402
from rehab_excel.outpatient_presentation import OutpatientPresentationError  # noqa: E402
from rehab_excel.outpatient_schedule_source import (  # noqa: E402
    OutpatientScheduleSourceError,
    read_unified_base_schedule,
)
from rehab_excel.patient_centric_preview import (  # noqa: E402
    PatientCentricPreviewError,
    build_patient_centric_preview_plan,
)
from rehab_excel.patient_registry_source import (  # noqa: E402
    PatientRegistrySourceError,
    read_patient_registry,
)
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.student_registry import StudentRegistryError, read_students  # noqa: E402


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _is_legacy_student_name(name: str) -> bool:
    compact = _norm(name).replace(" ", "")
    return compact.startswith("ΦΟΙΤ") or compact.startswith("FOIT")


def _read_daily_input_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputReadError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid DAILY_INPUT date: {raw!r}")


def _replacement_id_factory(history_rows, session_by_id):
    used = {str(row.replacement_id).strip() for row in history_rows}

    def allocate(session_id: str) -> str:
        target = session_by_id[session_id]
        base = f"daily-global:{target.session_id}:{target.session_date.isoformat()}"
        if base not in used:
            used.add(base)
            return base
        version = 2
        while f"{base}:v{version}" in used:
            version += 1
        value = f"{base}:v{version}"
        used.add(value)
        return value

    return allocate


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Resolve all therapist-absence displaced sessions as one global assignment, "
            "then build a reversible Excel preview with authoritative replacement facts."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "DAILY_INPUT_PREVIEW.xlsm",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "GLOBAL_DAILY_REPLACEMENTS_PREVIEW.xlsm",
    )
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    input_book = args.input.resolve()
    output = args.output.resolve()
    if not input_book.exists():
        print(f"SAFETY STOP: input workbook not found: {input_book}")
        return 2

    output_created = False
    try:
        input_contract = assert_application_contract(input_book)

        target_date = _read_daily_input_date(input_book)
        patients = read_patient_registry(input_book)
        students = tuple(read_students(input_book))
        settings = read_settings(input_book)
        base_entries = read_unified_base_schedule(input_book)
        sessions = materialize_sessions_for_date(base_entries, target_date)
        daily = read_daily_input(
            input_book,
            patients=patients,
            sessions=sessions,
            therapist_names=settings.therapist_names,
        )
        policy_book = load_policy_book(REPO_ROOT / "Config" / "provider_policy.json")
        service_hours = load_service_hours(REPO_ROOT / "Config" / "service_hours.json")
        adjacency_penalty = build_reclining_adjacency_penalty(base_entries)
        service_id_by_session = build_service_id_by_session(base_entries, target_date)

        existing_replacements = tuple(
            read_daily_replacement_facts(
                input_book,
                sessions=sessions,
                target_date=target_date,
            )
        )
        history_rows = tuple(
            read_daily_replacement_fact_rows(
                input_book,
                target_date=target_date,
                include_restored=True,
            )
        )

        therapist_absences = tuple(
            absence
            for absence in daily.absences
            if absence.absence_kind == AbsenceKind.THERAPIST
        )
        if not therapist_absences:
            raise DailyInputReadError("DAILY_INPUT has no therapist absence rows")

        therapists = tuple(
            Therapist(therapist_id=name, display_name=name)
            for name in settings.therapist_names
            if not _is_legacy_student_name(name)
        )
        queue = build_therapist_absence_replacement_queue(
            sessions=sessions,
            absences=daily.absences,
            cancellations=daily.cancellations,
            replacements=existing_replacements,
            therapists=therapists,
            patients=patients,
            timeslots=settings.standard_timeslots,
            students=students,
            policy_book=policy_book,
            base_entries=base_entries,
        )

        session_by_id = {session.session_id: session for session in sessions}
        target_sessions = tuple(session_by_id[item.session_id] for item in queue.items)
        if not target_sessions:
            raise DailyInputReadError(
                "No unresolved therapist-absence sessions remain after existing daily facts"
            )

        result = build_global_replacement_assignments(
            target_sessions,
            therapists=therapists,
            sessions=sessions,
            absences=daily.absences,
            patients=patients,
            replacements=existing_replacements,
            cancellations=daily.cancellations,
            timeslots=settings.standard_timeslots,
            students=students,
            adjacency_penalty=adjacency_penalty,
            policy_book=policy_book,
            base_entries=base_entries,
            service_hours=service_hours,
            service_id_by_session=service_id_by_session,
            replacement_id_factory=_replacement_id_factory(history_rows, session_by_id),
            reason="DAILY_INPUT global therapist absence plan",
        )
        if result.unresolved_request_ids:
            unresolved = ", ".join(result.unresolved_request_ids)
            raise DailyInputReadError(
                "Global replacement plan is incomplete; no preview will be written. "
                f"Unresolved sessions: {unresolved}"
            )

        all_replacements = existing_replacements + result.assignments
        states = build_daily_session_states(
            sessions,
            absences=daily.absences,
            replacements=all_replacements,
            cancellations=daily.cancellations,
            target_date=target_date,
        )
        statistics = calculate_daily_treatment_statistics(states)
        if not statistics.is_finalized:
            raise DailyInputReadError(
                "Global daily resolution is not finalized: "
                f"{statistics.unresolved_sessions} orphan session(s) remain"
            )

        preview = build_patient_centric_preview_plan(
            input_book,
            base_entries=base_entries,
            states=states,
            sessions=sessions,
            patients=patients,
            students=students,
            rebuild_multi_member_groups=True,
        )
        write_plan = compose_outpatient_daily_plan(
            preview.write_plan,
            states=states,
            patients=patients,
        )
        master_overlay = build_master_daily_overlay_plan(
            input_book,
            base_entries=base_entries,
            states=states,
            patients=patients,
        )
        write_plan = compose_master_daily_overlay(write_plan, master_overlay)
        report = apply_write_plan_to_copy(write_plan, output, overwrite=args.overwrite)
        output_created = True

        persist_global_replacement_result(
            output,
            result,
            sessions=sessions,
            existing_active_replacements=existing_replacements,
        )
        write_daily_statistics(output, target_date, statistics)

        ledger_records = build_global_replacement_ledger_records(
            result,
            sessions=sessions,
            patients=patients,
            bindings=preview.bindings,
        )
        install_daily_change_ledger_preview(output, input_book, ledger_records)
        output_contract = assert_application_contract(output)

    except (
        ApplicationContractError,
        DailyInputReadError,
        DailyReplacementFactError,
        DailyStatisticsWriteError,
        PatientCentricPreviewError,
        PatientRegistrySourceError,
        StudentRegistryError,
        OutpatientScheduleSourceError,
        OutpatientPresentationError,
        DailyPreviewCompositionError,
        MasterDailyOverlayError,
        GlobalReplacementLifecycleError,
        GlobalReplacementLedgerError,
        DailyChangeLedgerPreviewError,
        ValueError,
        KeyError,
    ) as exc:
        if output_created and output.exists():
            try:
                output.unlink()
            except OSError:
                pass
        print(f"SAFETY STOP: {exc}")
        return 2

    print("GLOBAL DAILY REPLACEMENT PREVIEW OK")
    print(f"Date: {target_date.isoformat()}")
    print(f"Displaced sessions solved together: {len(target_sessions)}")
    print(f"New replacement assignments: {len(result.assignments)}")
    print(f"Existing active replacements retained: {len(existing_replacements)}")
    print(f"Unresolved sessions: {len(result.unresolved_request_ids)}")
    print(f"Preview: {report.output_path}")
    print(f"Input unchanged: {report.source_unchanged}")
    print("Application contract input: OK")
    print("Application contract output: OK")
    print(f"Required app components: {len(input_contract.required_components)}")
    print(f"Output VBA component count: {len(output_contract.component_names)}")
    for assignment in result.assignments:
        print(
            f"  {assignment.patient_id}: {assignment.original_therapist_id} -> "
            f"{assignment.replacement_therapist_id} "
            f"{assignment.replacement_time.strftime('%H:%M')}"
        )
    return 0 if report.source_unchanged else 3


if __name__ == "__main__":
    raise SystemExit(main())
