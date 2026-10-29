from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.day_finalization import build_finalized_day_snapshot  # noqa: E402
from rehab_core.day_history_events import build_day_finalized_event  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.therapist_day_history import build_therapist_day_history_rows  # noqa: E402
from rehab_excel.application_contract import assert_application_contract  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_reader import read_daily_input  # noqa: E402
from rehab_excel.daily_replacement_facts import read_daily_replacement_facts  # noqa: E402
from rehab_excel.day_history_event_store import read_day_history_events  # noqa: E402
from rehab_excel.finalized_day_snapshots import read_finalized_day_snapshots  # noqa: E402
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.therapist_day_history_store import read_therapist_day_history_rows  # noqa: E402


class CloseDayPreparationError(RuntimeError):
    pass


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def _daily_input_date(path: Path) -> date:
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise CloseDayPreparationError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise CloseDayPreparationError(f"Invalid DAILY_INPUT date: {raw!r}")


def _is_legacy_student_name(value: str) -> bool:
    compact = str(value).strip().casefold().replace(" ", "")
    return compact.startswith("φοιτ") or compact.startswith("foit")


def _escape(value: object | None) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("\\", "\\\\")
        .replace("\t", "\\t")
        .replace("\r", "\\r")
        .replace("\n", "\\n")
    )


def _line(kind: str, values: tuple[object | None, ...]) -> str:
    return "\t".join([kind, *(_escape(value) for value in values)])


def prepare_close_day_payload(
    workbook_path: str | Path,
    payload_path: str | Path,
    *,
    actor_or_source: str,
    approved_exception_count: int = 0,
) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    payload = Path(payload_path).resolve()
    if not source.exists():
        raise CloseDayPreparationError(f"Workbook not found: {source}")
    if source.suffix.casefold() != ".xlsm":
        raise CloseDayPreparationError("Workbook must be an .xlsm file")
    if approved_exception_count < 0:
        raise CloseDayPreparationError("approved_exception_count cannot be negative")
    actor = actor_or_source.strip()
    if not actor:
        raise CloseDayPreparationError("actor_or_source cannot be blank")

    source_hash = file_sha256(source)
    contract = assert_application_contract(source)
    service_date = _daily_input_date(source)

    existing_current = read_finalized_day_snapshots(
        source,
        service_date=service_date,
        current_only=True,
    )
    if existing_current:
        raise CloseDayPreparationError(
            f"Service date {service_date.isoformat()} is already finalized"
        )

    patients = read_patient_registry(source)
    settings = read_settings(source)
    base_entries = read_unified_base_schedule(source)
    sessions = materialize_sessions_for_date(base_entries, service_date)
    daily = read_daily_input(
        source,
        patients=patients,
        sessions=sessions,
        therapist_names=settings.therapist_names,
    )
    replacements = read_daily_replacement_facts(
        source,
        sessions=sessions,
        target_date=service_date,
    )
    states = build_daily_session_states(
        sessions,
        absences=daily.absences,
        replacements=replacements,
        cancellations=daily.cancellations,
        target_date=service_date,
    )
    finalized_at = datetime.now().astimezone()
    snapshot = build_finalized_day_snapshot(
        states,
        snapshot_id=f"finalized:{service_date.isoformat()}:v1",
        service_date=service_date,
        finalized_at=finalized_at,
        planned_session_count=len(sessions),
        approved_exception_count=approved_exception_count,
    )
    therapist_ids = tuple(
        name
        for name in settings.therapist_names
        if str(name).strip() and not _is_legacy_student_name(name)
    )
    therapist_rows = build_therapist_day_history_rows(
        states,
        snapshot_id=snapshot.snapshot_id,
        service_date=service_date,
        therapist_ids=therapist_ids,
        capacity_timeslots_by_therapist={name: 6 for name in therapist_ids},
    )
    event = build_day_finalized_event(
        event_id=f"day-finalized:{service_date.isoformat()}:v{snapshot.version}",
        service_date=service_date,
        recorded_at=finalized_at,
        actor_or_source=actor,
        snapshot_id=snapshot.snapshot_id,
        summary=(
            f"Day finalized: {snapshot.final_session_count} performed sessions, "
            f"{snapshot.replacement_count} replacements, "
            f"{snapshot.treatment_suspension_count} suspensions"
        ),
    )

    if read_therapist_day_history_rows(source, snapshot_id=snapshot.snapshot_id):
        raise CloseDayPreparationError("Therapist/day history already exists for snapshot")
    if any(item.event_id == event.event_id for item in read_day_history_events(source)):
        raise CloseDayPreparationError("DAY_FINALIZED event already exists")

    snapshot_values: tuple[object | None, ...] = (
        snapshot.snapshot_id,
        snapshot.service_date.isoformat(),
        snapshot.version,
        snapshot.status.value,
        snapshot.finalized_at.isoformat(),
        snapshot.supersedes_snapshot_id,
        snapshot.unresolved_sessions,
        snapshot.planned_session_count,
        snapshot.final_session_count,
        snapshot.patient_absence_count,
        snapshot.cancellation_count,
        snapshot.no_show_count,
        snapshot.replacement_count,
        snapshot.treatment_suspension_count,
        snapshot.approved_exception_count,
    )
    event_values: tuple[object | None, ...] = (
        event.event_id,
        event.service_date.isoformat(),
        event.event_type.value,
        event.recorded_at.isoformat(),
        event.actor_or_source,
        event.therapist_id,
        event.patient_id,
        event.session_id,
        event.reason_code,
        event.related_event_id,
        event.related_snapshot_id,
        event.summary,
    )

    lines = [_line("SNAPSHOT", snapshot_values)]
    for row in therapist_rows:
        lines.append(
            _line(
                "THERAPIST",
                (
                    row.snapshot_id,
                    row.service_date.isoformat(),
                    row.therapist_id,
                    row.capacity_timeslots,
                    row.planned_occupied_timeslots,
                    row.planned_sessions,
                    row.final_occupied_timeslots,
                    row.final_active_sessions,
                    row.replacements_received,
                    row.released_by_cancellation_or_no_show,
                    row.treatment_suspensions_affecting_load,
                    row.approved_shared_or_over_capacity_events,
                ),
            )
        )
    lines.append(_line("EVENT", event_values))
    payload.parent.mkdir(parents=True, exist_ok=True)
    payload.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")

    if file_sha256(source) != source_hash:
        try:
            payload.unlink()
        except OSError:
            pass
        raise CloseDayPreparationError("Workbook changed during close-day validation")

    return {
        "ok": True,
        "payload_path": str(payload),
        "source_sha256_before": source_hash,
        "service_date": service_date.isoformat(),
        "snapshot_id": snapshot.snapshot_id,
        "event_id": event.event_id,
        "final_sessions": snapshot.final_session_count,
        "therapist_history_rows": len(therapist_rows),
        "unresolved_sessions": snapshot.unresolved_sessions,
        "application_contract_components": len(contract.component_names),
        "source_unchanged": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare a validated in-place close-day payload.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--payload", required=True, type=Path)
    parser.add_argument("--response", type=Path)
    parser.add_argument("--actor-or-source", default="close-day-action")
    parser.add_argument("--approved-exception-count", type=int, default=0)
    args = parser.parse_args()

    try:
        result = prepare_close_day_payload(
            args.input,
            args.payload,
            actor_or_source=args.actor_or_source,
            approved_exception_count=args.approved_exception_count,
        )
    except Exception as exc:
        try:
            if args.payload.exists():
                args.payload.unlink()
        except OSError:
            pass
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2

    _write_response(args.response, result)
    print("CLOSE DAY IN-PLACE PAYLOAD VERIFIED")
    print(f"Date: {result['service_date']}")
    print(f"Snapshot: {result['snapshot_id']}")
    print(f"Therapist/day history rows: {result['therapist_history_rows']}")
    print(f"Event: {result['event_id']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
