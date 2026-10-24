from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.daily_statistics import calculate_daily_treatment_statistics  # noqa: E402
from rehab_core.models import SessionCancellationKind  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_replacement_facts import (  # noqa: E402
    DailyReplacementFactError,
    read_daily_replacement_facts,
)
from rehab_excel.outpatient_schedule_source import (  # noqa: E402
    OutpatientScheduleSourceError,
    read_unified_base_schedule,
)
from rehab_excel.patient_registry_source import (  # noqa: E402
    PatientRegistrySourceError,
    read_patient_registry,
)


STATISTICS_SHEET = "DAILY_STATISTICS"


class DailyFactVerificationError(RuntimeError):
    pass


def _stored_statistics(path: Path) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.casefold() == ".xlsm")
    try:
        if STATISTICS_SHEET not in wb.sheetnames:
            raise DailyFactVerificationError(f"Workbook has no {STATISTICS_SHEET} sheet")
        ws = wb[STATISTICS_SHEET]
        values: dict[str, str] = {}
        for row in range(4, ws.max_row + 1):
            key = str(ws.cell(row, 1).value or "").strip()
            if not key:
                continue
            value = ws.cell(row, 3).value
            values[key] = str(value if value is not None else "").strip()
        return values
    finally:
        wb.close()


def _as_int(values: dict[str, str], key: str) -> int:
    if key not in values:
        raise DailyFactVerificationError(f"{STATISTICS_SHEET} has no {key!r} row")
    try:
        return int(float(values[key]))
    except ValueError as exc:
        raise DailyFactVerificationError(
            f"{STATISTICS_SHEET} {key!r} is not numeric: {values[key]!r}"
        ) from exc


def verify_daily_fact_reconstruction(
    workbook_path: str | Path,
    *,
    expected_replacements: int | None = None,
    expected_suspended: int | None = None,
    expected_unresolved: int | None = 0,
) -> tuple[int, int, int, bool]:
    path = Path(workbook_path).resolve()
    if not path.exists():
        raise DailyFactVerificationError(f"Workbook not found: {path}")

    patients = read_patient_registry(path)
    base_entries = read_unified_base_schedule(path)

    # DAILY_INPUT owns the concrete operational date. Read B2 first without
    # duplicating its date parser by materializing candidate sessions for each
    # date is not viable, so inspect the stored date value here and normalize it.
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.casefold() == ".xlsm")
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyFactVerificationError("Workbook has no DAILY_INPUT sheet")
        raw_date = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()

    from datetime import date, datetime

    if isinstance(raw_date, datetime):
        target_date = raw_date.date()
    elif isinstance(raw_date, date):
        target_date = raw_date
    else:
        text = str(raw_date or "").strip()
        target_date = None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                target_date = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                pass
        if target_date is None:
            raise DailyFactVerificationError(f"Invalid DAILY_INPUT date: {raw_date!r}")

    sessions = materialize_sessions_for_date(base_entries, target_date)
    daily = read_daily_input(path, patients=patients, sessions=sessions)
    replacements = read_daily_replacement_facts(
        path,
        sessions=sessions,
        target_date=target_date,
    )
    states = build_daily_session_states(
        sessions,
        absences=daily.absences,
        replacements=replacements,
        cancellations=daily.cancellations,
        target_date=target_date,
    )
    statistics = calculate_daily_treatment_statistics(states)
    suspended = sum(
        1
        for cancellation in daily.cancellations
        if cancellation.kind == SessionCancellationKind.TREATMENT_SUSPENDED
    )

    if expected_replacements is not None and len(replacements) != expected_replacements:
        raise DailyFactVerificationError(
            f"Expected {expected_replacements} persisted replacements, found {len(replacements)}"
        )
    if expected_suspended is not None and suspended != expected_suspended:
        raise DailyFactVerificationError(
            f"Expected {expected_suspended} treatment suspensions, found {suspended}"
        )
    if expected_unresolved is not None and statistics.unresolved_sessions != expected_unresolved:
        raise DailyFactVerificationError(
            f"Expected {expected_unresolved} unresolved sessions, found {statistics.unresolved_sessions}"
        )

    stored = _stored_statistics(path)
    expected_stored = {
        "scheduled": statistics.scheduled_sessions,
        "performed": statistics.performed_sessions,
        "replaced": statistics.replaced_sessions,
        "patient_absent": statistics.patient_absent_sessions,
        "department_postponed": statistics.department_postponed_sessions,
        "patient_no_show": statistics.patient_no_show_sessions,
        "treatment_suspended": statistics.treatment_suspended_sessions,
        "unresolved": statistics.unresolved_sessions,
    }
    for key, reconstructed in expected_stored.items():
        saved = _as_int(stored, key)
        if saved != reconstructed:
            raise DailyFactVerificationError(
                f"Stored {key}={saved} does not match reconstructed {reconstructed}"
            )
    stored_finalized = stored.get("finalized", "").strip().upper()
    reconstructed_finalized = "ΝΑΙ" if statistics.is_finalized else "ΟΧΙ"
    if stored_finalized != reconstructed_finalized:
        raise DailyFactVerificationError(
            f"Stored finalized={stored_finalized!r} does not match reconstructed "
            f"{reconstructed_finalized!r}"
        )

    return (
        len(replacements),
        suspended,
        statistics.unresolved_sessions,
        statistics.is_finalized,
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Reconstruct one finalized day from persisted operational facts and compare statistics."
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--expected-replacements", type=int)
    parser.add_argument("--expected-suspended", type=int)
    parser.add_argument("--expected-unresolved", type=int, default=0)
    args = parser.parse_args()

    try:
        replacements, suspended, unresolved, finalized = verify_daily_fact_reconstruction(
            args.workbook,
            expected_replacements=args.expected_replacements,
            expected_suspended=args.expected_suspended,
            expected_unresolved=args.expected_unresolved,
        )
    except (
        DailyFactVerificationError,
        DailyInputReadError,
        DailyReplacementFactError,
        OutpatientScheduleSourceError,
        PatientRegistrySourceError,
        ValueError,
        KeyError,
    ) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("DAILY FACT RECONSTRUCTION OK")
    print(f"Workbook: {args.workbook.resolve()}")
    print(f"Persisted replacements reread: {replacements}")
    print(f"Treatment suspensions reread: {suspended}")
    print(f"Reconstructed unresolved: {unresolved}")
    print(f"Reconstructed finalized: {'ΝΑΙ' if finalized else 'ΟΧΙ'}")
    print("Stored DAILY_STATISTICS matches reconstructed daily state: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
