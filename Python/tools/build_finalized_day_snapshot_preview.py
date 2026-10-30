from __future__ import annotations

import argparse
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.day_finalization import (  # noqa: E402
    DayFinalizationError,
    FinalizedDaySnapshot,
    build_finalized_day_snapshot,
)
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.therapist_day_history import (  # noqa: E402
    TherapistDayHistoryError,
    build_therapist_day_history_rows,
)
from rehab_excel.application_contract import (  # noqa: E402
    ApplicationContractError,
    assert_application_contract,
)
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_replacement_facts import (  # noqa: E402
    DailyReplacementFactError,
    read_daily_replacement_facts,
)
from rehab_excel.finalized_day_snapshots import (  # noqa: E402
    FinalizedDaySnapshotStoreError,
    append_finalized_day_snapshot,
    read_finalized_day_snapshots,
)
from rehab_excel.outpatient_schedule_source import (  # noqa: E402
    OutpatientScheduleSourceError,
    read_unified_base_schedule,
)
from rehab_excel.patient_registry_source import (  # noqa: E402
    PatientRegistrySourceError,
    read_patient_registry,
)
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.therapist_day_history_store import (  # noqa: E402
    TherapistDayHistoryStoreError,
    append_therapist_day_history_rows,
    read_therapist_day_history_rows,
)


class FinalizedDayPreviewError(RuntimeError):
    pass


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _daily_input_date(path: Path) -> date:
    wb = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise FinalizedDayPreviewError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()

    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise FinalizedDayPreviewError(f"Invalid DAILY_INPUT date: {raw!r}")


def _is_legacy_student_name(value: str) -> bool:
    compact = str(value).strip().casefold().replace(" ", "")
    return compact.startswith("φοιτ") or compact.startswith("foit")


def build_finalized_day_snapshot_preview(
    input_path: str | Path,
    output_path: str | Path,
    *,
    approved_exception_count: int = 0,
    overwrite: bool = False,
) -> tuple[FinalizedDaySnapshot, int, int]:
    source = Path(input_path).resolve()
    output = Path(output_path).resolve()
    if not source.exists():
        raise FinalizedDayPreviewError(f"Input workbook not found: {source}")
    if source == output:
        raise FinalizedDayPreviewError("Output must be different from input workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise FinalizedDayPreviewError("Input and output must both be .xlsm files")
    if output.exists() and not overwrite:
        raise FinalizedDayPreviewError(f"Output already exists: {output}")
    if approved_exception_count < 0:
        raise FinalizedDayPreviewError("approved_exception_count cannot be negative")

    source_before = _sha256(source)
    source_contract = assert_application_contract(source)

    service_date = _daily_input_date(source)
    patients = read_patient_registry(source)
    settings = read_settings(source)
    base_entries = read_unified_base_schedule(source)
    sessions = materialize_sessions_for_date(base_entries, service_date)
    daily = read_daily_input(
        source,
        patients=patients,
        sessions=sessions,
        therapist_names=settings.therapist_names,
    )
    replacements = read_daily_replacement_facts(
        source,
        sessions=sessions,
        target_date=service_date,
    )
    states = build_daily_session_states(
        sessions,
        absences=daily.absences,
        replacements=replacements,
        cancellations=daily.cancellations,
        target_date=service_date,
    )

    snapshot = build_finalized_day_snapshot(
        states,
        snapshot_id=f"finalized:{service_date.isoformat()}:v1",
        service_date=service_date,
        finalized_at=datetime.now().astimezone(),
        planned_session_count=len(sessions),
        approved_exception_count=approved_exception_count,
    )
    therapist_ids = tuple(
        name
        for name in settings.therapist_names
        if str(name).strip() and not _is_legacy_student_name(name)
    )
    therapist_history = build_therapist_day_history_rows(
        states,
        snapshot_id=snapshot.snapshot_id,
        service_date=service_date,
        therapist_ids=therapist_ids,
        capacity_timeslots_by_therapist={name: 6 for name in therapist_ids},
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        try:
            output.unlink()
        except PermissionError as exc:
            raise FinalizedDayPreviewError(
                f"Output workbook is open or locked by Excel: {output}"
            ) from exc
    shutil.copy2(source, output)

    try:
        written = append_finalized_day_snapshot(output, snapshot)
        if not written:
            raise FinalizedDayPreviewError("Snapshot append unexpectedly reported no change")
        reread = read_finalized_day_snapshots(
            output,
            service_date=service_date,
            current_only=True,
        )
        if reread != (snapshot,):
            raise FinalizedDayPreviewError(
                "Persisted finalized snapshot did not read back exactly as requested"
            )

        history_written = append_therapist_day_history_rows(output, therapist_history)
        if history_written != len(therapist_history):
            raise FinalizedDayPreviewError(
                "Therapist/day history append count did not match requested rows"
            )
        history_reread = read_therapist_day_history_rows(
            output,
            snapshot_id=snapshot.snapshot_id,
            service_date=service_date,
        )
        if history_reread != therapist_history:
            raise FinalizedDayPreviewError(
                "Persisted therapist/day history did not read back exactly as requested"
            )

        output_contract = assert_application_contract(output)
    except Exception:
        try:
            output.unlink()
        except Exception:
            pass
        raise

    source_after = _sha256(source)
    if source_before != source_after:
        try:
            output.unlink()
        except Exception:
            pass
        raise FinalizedDayPreviewError("Input workbook changed during preview build")

    return (
        snapshot,
        len(source_contract.component_names),
        len(output_contract.component_names),
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Build a safe .xlsm copy containing one verified CURRENT_FINAL day snapshot "
            "and therapist/day historical workload rows."
        )
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--approved-exception-count", type=int, default=0)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    try:
        snapshot, source_components, output_components = build_finalized_day_snapshot_preview(
            args.input,
            args.output,
            approved_exception_count=args.approved_exception_count,
            overwrite=args.overwrite,
        )
    except (
        FinalizedDayPreviewError,
        DayFinalizationError,
        FinalizedDaySnapshotStoreError,
        TherapistDayHistoryError,
        TherapistDayHistoryStoreError,
        DailyInputReadError,
        DailyReplacementFactError,
        OutpatientScheduleSourceError,
        PatientRegistrySourceError,
        ApplicationContractError,
        ValueError,
        KeyError,
    ) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    history = read_therapist_day_history_rows(
        args.output.resolve(),
        snapshot_id=snapshot.snapshot_id,
        service_date=snapshot.service_date,
    )
    print("FINALIZED DAY SNAPSHOT PREVIEW VERIFIED")
    print(f"Date: {snapshot.service_date.isoformat()}")
    print(f"Snapshot: {snapshot.snapshot_id}")
    print(f"Version/status: {snapshot.version} / {snapshot.status.value}")
    print(f"Final sessions: {snapshot.final_session_count}")
    print(f"Replacements: {snapshot.replacement_count}")
    print(f"Patient absences: {snapshot.patient_absence_count}")
    print(f"No-shows: {snapshot.no_show_count}")
    print(f"Treatment suspensions: {snapshot.treatment_suspension_count}")
    print(f"Approved exceptions: {snapshot.approved_exception_count}")
    print(f"Unresolved sessions: {snapshot.unresolved_sessions}")
    print(f"Therapist/day history rows: {len(history)}")
    print("Input unchanged: True")
    print(f"Application contract source: OK ({source_components} components)")
    print(f"Application contract output: OK ({output_components} components)")
    print(f"Output: {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
