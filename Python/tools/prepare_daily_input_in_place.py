from __future__ import annotations

import argparse
from datetime import date, datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_excel.application_contract import assert_application_contract  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_preview_composer import compose_outpatient_daily_plan  # noqa: E402
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_centric_preview import build_patient_centric_preview_plan  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.writeback import CellPatch, validate_write_plan  # noqa: E402


class DailyInputInPlacePreparationError(RuntimeError):
    pass


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def _target_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputInPlacePreparationError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputInPlacePreparationError(f"Invalid DAILY_INPUT date: {raw!r}")


def _escape(value: object | None) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("\\", "\\\\")
        .replace("\t", "\\t")
        .replace("\r", "\\r")
        .replace("\n", "\\n")
    )


def _bool_token(value: bool | None) -> str:
    if value is None:
        return ""
    return "1" if value else "0"


def _text_runs_token(patch: CellPatch) -> str:
    return ";".join(
        ":".join(
            (
                str(run.start),
                str(run.length),
                "1" if run.strike_through else "0",
                "1" if run.italic else "0",
                run.font_role or "",
            )
        )
        for run in patch.text_runs
    )


def _patch_line(patch: CellPatch) -> str:
    fields = (
        patch.sheet,
        patch.cell,
        patch.intent.value,
        patch.value,
        "1" if patch.strike_through else "0",
        patch.font_role,
        patch.fill_role,
        patch.border_role,
        _bool_token(patch.wrap_text),
        patch.min_font_size,
        patch.source_tag,
        _text_runs_token(patch),
    )
    return "\t".join(["PATCH", *(_escape(value) for value in fields)])


def prepare_daily_input_payload(
    workbook_path: str | Path,
    payload_path: str | Path,
) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    payload = Path(payload_path).resolve()
    if not source.exists():
        raise DailyInputInPlacePreparationError(f"Workbook not found: {source}")
    if source.suffix.casefold() != ".xlsm":
        raise DailyInputInPlacePreparationError("Workbook must be an .xlsm file")

    source_hash = file_sha256(source)
    contract = assert_application_contract(source)
    target_date = _target_date(source)
    patients = read_patient_registry(source)
    settings = read_settings(source)
    base_entries = read_unified_base_schedule(source)
    sessions = materialize_sessions_for_date(base_entries, target_date)
    daily = read_daily_input(
        source,
        patients=patients,
        sessions=sessions,
        therapist_names=settings.therapist_names,
    )
    if daily.target_date != target_date:
        raise DailyInputReadError("DAILY_INPUT date changed during read")

    states = build_daily_session_states(
        sessions,
        absences=daily.absences,
        replacements=(),
        cancellations=daily.cancellations,
        target_date=target_date,
    )
    changed = tuple(state for state in states if state.status.value != "active")
    if not changed:
        raise DailyInputReadError(
            "DAILY_INPUT rows did not match any scheduled session on the selected date"
        )
    replacement_sessions_needed = sum(
        1 for state in states if state.status.value == "therapist_absent"
    )

    preview = build_patient_centric_preview_plan(
        source,
        base_entries=base_entries,
        states=states,
        sessions=sessions,
        patients=patients,
        rebuild_multi_member_groups=True,
    )
    plan = compose_outpatient_daily_plan(
        preview.write_plan,
        states=states,
        patients=patients,
    )
    validate_write_plan(plan)

    payload.parent.mkdir(parents=True, exist_ok=True)
    payload.write_text(
        "\n".join(_patch_line(patch) for patch in plan.patches) + "\n",
        encoding="utf-8-sig",
    )

    if file_sha256(source) != source_hash:
        try:
            payload.unlink()
        except OSError:
            pass
        raise DailyInputInPlacePreparationError("Workbook changed during DAILY_INPUT validation")

    return {
        "ok": True,
        "payload_path": str(payload),
        "source_sha256_before": source_hash,
        "date": target_date.isoformat(),
        "therapist_rows_read": daily.therapist_rows_used,
        "patient_rows_read": daily.patient_rows_used,
        "operational_absences": len(daily.absences),
        "session_cancellations": len(daily.cancellations),
        "changed_sessions": len(preview.bindings),
        "replacement_sessions_needed": replacement_sessions_needed,
        "patch_count": len(plan.patches),
        "application_contract_components": len(contract.component_names),
        "source_unchanged": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare a validated in-place DAILY_INPUT write plan.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--payload", required=True, type=Path)
    parser.add_argument("--response", type=Path)
    args = parser.parse_args()

    try:
        result = prepare_daily_input_payload(args.input, args.payload)
    except Exception as exc:
        try:
            if args.payload.exists():
                args.payload.unlink()
        except OSError:
            pass
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2

    _write_response(args.response, result)
    print("DAILY INPUT IN-PLACE PAYLOAD VERIFIED")
    print(f"Date: {result['date']}")
    print(f"Patches: {result['patch_count']}")
    print(f"Replacement sessions needed: {result['replacement_sessions_needed']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
