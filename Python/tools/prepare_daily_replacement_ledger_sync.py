from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_replacement_facts import read_daily_replacement_fact_rows  # noqa: E402
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402

HEADERS = (
    "change_id",
    "date",
    "patient_id",
    "patient_name",
    "original_provider",
    "original_time",
    "effective_provider",
    "effective_time",
    "reason",
    "status",
)


class DailyReplacementLedgerSyncError(RuntimeError):
    pass


def _target_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyReplacementLedgerSyncError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyReplacementLedgerSyncError(f"Invalid DAILY_INPUT date: {raw!r}")


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8-sig")


def prepare_ledger_rows(workbook_path: str | Path, rows_path: str | Path) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    output = Path(rows_path).resolve()
    if not source.exists():
        raise DailyReplacementLedgerSyncError(f"Workbook not found: {source}")
    source_hash = file_sha256(source)

    target_date = _target_date(source)
    patients = tuple(read_patient_registry(source))
    patient_by_id = {patient.patient_id: patient for patient in patients}
    base_entries = tuple(read_unified_base_schedule(source))
    sessions = tuple(materialize_sessions_for_date(base_entries, target_date))
    session_by_id = {session.session_id: session for session in sessions}
    facts = tuple(
        read_daily_replacement_fact_rows(
            source,
            target_date=target_date,
            include_restored=True,
        )
    )

    rows: list[dict[str, str]] = []
    for fact in facts:
        target = session_by_id.get(fact.target_session_id)
        if target is None:
            raise DailyReplacementLedgerSyncError(
                f"Replacement fact target session is missing: {fact.target_session_id!r}"
            )
        if target.patient_id != fact.patient_id:
            raise DailyReplacementLedgerSyncError(
                f"Replacement fact patient mismatch for {fact.replacement_id!r}"
            )
        patient = patient_by_id.get(fact.patient_id)
        patient_name = patient.display_name if patient is not None else fact.patient_id
        rows.append(
            {
                "change_id": fact.replacement_id,
                "date": fact.replacement_date.isoformat(),
                "patient_id": fact.patient_id,
                "patient_name": patient_name,
                "original_provider": fact.original_therapist_id,
                "original_time": target.start_time.strftime("%H:%M"),
                "effective_provider": fact.replacement_therapist_id,
                "effective_time": fact.replacement_time.strftime("%H:%M"),
                "reason": fact.reason,
                "status": "ACTIVE" if fact.is_active else "RESTORED",
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    if file_sha256(source) != source_hash:
        output.unlink(missing_ok=True)
        raise DailyReplacementLedgerSyncError("Workbook changed during ledger preparation")

    return {
        "ok": True,
        "date": target_date.isoformat(),
        "row_count": len(rows),
        "source_unchanged": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare DAILY_CHANGE_LEDGER rows from authoritative replacement facts.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--rows", required=True, type=Path)
    parser.add_argument("--response", required=True, type=Path)
    args = parser.parse_args()
    try:
        result = prepare_ledger_rows(args.input, args.rows)
    except Exception as exc:
        args.rows.unlink(missing_ok=True)
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2
    _write_response(args.response, result)
    print("DAILY REPLACEMENT LEDGER PREP OK")
    print(f"Rows: {result['row_count']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
