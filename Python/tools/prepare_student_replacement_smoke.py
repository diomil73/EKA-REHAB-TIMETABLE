from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.models import AbsenceKind  # noqa: E402
from rehab_excel.application_contract import (  # noqa: E402
    ApplicationContractError,
    assert_application_contract,
)
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_patients import audit_daily_input_patients  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_input_sheet import (  # noqa: E402
    DailyInputSheetError,
    build_daily_input_spec,
    create_daily_input_preview,
)
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule  # noqa: E402
from rehab_excel.patient_registry_source import read_patient_registry  # noqa: E402
from rehab_excel.reader import read_settings  # noqa: E402
from rehab_excel.student_registry import (  # noqa: E402
    STUDENT_REGISTRY_SHEET,
    StudentRegistryError,
    read_students,
)
from rehab_excel.student_registry_schema import (  # noqa: E402
    StudentRegistrySchemaError,
    Win32ComStudentRegistrySchemaBackend,
    student_registry_schema_present,
)


SMOKE_STUDENT_ID = "0000-STUDENT-SMOKE"
SMOKE_STUDENT_NAME = "0000 STUDENT SMOKE"


class StudentReplacementSmokeError(RuntimeError):
    pass


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value.strip())
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Date must be YYYY-MM-DD") from exc


def _daily_input_therapist_names(path: Path) -> tuple[str, ...]:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            return ()
        ws = wb["DAILY_INPUT"]
        names: list[str] = []
        max_row = max(7, ws.max_row)
        for row in range(7, max_row + 1):
            value = str(ws.cell(row=row, column=1).value or "").strip()
            if value:
                names.append(value)
        return tuple(dict.fromkeys(names))
    finally:
        wb.close()


def choose_student_smoke_target(
    sessions,
    *,
    allowed_therapists: tuple[str, ...],
    max_displaced: int = 5,
):
    """Choose one real therapist absence small enough for one student to absorb.

    Prefer a provider whose affected sessions are all non-robotic, then the
    smallest displaced load. ``allowed_therapists`` must already be restricted
    to providers that can actually be marked absent in the DAILY_INPUT sheet.
    """
    if max_displaced < 1:
        raise ValueError("max_displaced must be positive")

    allowed = set(allowed_therapists)
    grouped = defaultdict(list)
    for session in sessions:
        if session.therapist_id in allowed:
            grouped[session.therapist_id].append(session)

    candidates = []
    for therapist_id, values in grouped.items():
        affected = tuple(
            sorted(values, key=lambda item: (item.start_time, item.patient_id, item.session_id))
        )
        if not affected or len(affected) > max_displaced:
            continue
        candidates.append(
            (
                int(any(session.robotic for session in affected)),
                len(affected),
                therapist_id,
                affected,
            )
        )

    if not candidates:
        raise ValueError(
            f"No DAILY_INPUT therapist has between 1 and {max_displaced} session(s) on the smoke date"
        )

    _robotic_flag, _count, therapist_id, affected = min(
        candidates, key=lambda item: (item[0], item[1], item[2])
    )
    return therapist_id, affected


def _write_smoke_rows(
    workbook_path: Path,
    *,
    target_date: date,
    absent_therapist_id: str,
) -> None:
    if sys.platform != "win32":
        raise StudentReplacementSmokeError("Windows + Microsoft Excel are required")
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise StudentReplacementSmokeError("pywin32 is required") from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(
            str(workbook_path.resolve()), UpdateLinks=0, ReadOnly=False
        )

        students_ws = workbook.Worksheets(STUDENT_REGISTRY_SHEET)
        last_row = students_ws.Cells(students_ws.Rows.Count, 1).End(-4162).Row
        next_row = 2 if last_row < 2 else last_row + 1
        students_ws.Cells(next_row, 1).Value = SMOKE_STUDENT_ID
        students_ws.Cells(next_row, 2).Value = SMOKE_STUDENT_NAME
        students_ws.Cells(next_row, 3).Value = target_date.isoformat()
        students_ws.Cells(next_row, 4).Value = target_date.isoformat()
        students_ws.Cells(next_row, 5).ClearContents()
        students_ws.Cells(next_row, 6).Value = "ΝΑΙ"
        students_ws.Cells(next_row, 7).Value = "ΝΑΙ"

        daily_ws = workbook.Worksheets("DAILY_INPUT")
        last_daily_row = max(7, daily_ws.UsedRange.Rows.Count)
        absence_row = None
        for row in range(7, last_daily_row + 1):
            value = str(daily_ws.Cells(row, 1).Value2 or "").strip()
            if value == absent_therapist_id:
                absence_row = row
                break
        if absence_row is None:
            raise StudentReplacementSmokeError(
                f"Therapist {absent_therapist_id!r} not found in DAILY_INPUT"
            )

        daily_ws.Cells(absence_row, 2).Value = "ΝΑΙ"
        daily_ws.Range(
            daily_ws.Cells(absence_row, 3), daily_ws.Cells(absence_row, 4)
        ).ClearContents()
        daily_ws.Cells(absence_row, 5).Value = "Student replacement smoke"
        daily_ws.Cells(absence_row, 6).Value = (
            f"Expect {SMOKE_STUDENT_ID} to be eligible as STUDENT provider"
        )
        workbook.Save()
    except StudentReplacementSmokeError:
        raise
    except Exception as exc:
        raise StudentReplacementSmokeError(f"Excel smoke preparation failed: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Build a full-app smoke copy that makes an authoritative STUDENT provider "
            "available for one therapist-absence replacement scenario."
        )
    )
    parser.add_argument("--date", type=_parse_date, required=True)
    parser.add_argument(
        "--source",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "OUTPATIENT_RECURRING_SMOKE.xlsm",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "STUDENT_REPLACEMENT_INPUT.xlsm",
    )
    parser.add_argument("--max-displaced", type=int, default=5)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    source = args.source.resolve()
    output = args.output.resolve()
    if not source.exists():
        print(f"SAFETY STOP: source workbook not found: {source}")
        return 2

    source_hash = file_sha256(source)
    output_created = False
    try:
        source_contract = assert_application_contract(source)
        patient_audit = audit_daily_input_patients(source)
        if not patient_audit.ok:
            raise StudentReplacementSmokeError("Patient identity audit failed")

        settings = read_settings(source)
        daily_input_names = _daily_input_therapist_names(source)
        allowed_smoke_therapists = tuple(
            name for name in settings.therapist_names if name in set(daily_input_names)
        )
        if not allowed_smoke_therapists:
            raise StudentReplacementSmokeError(
                "DAILY_INPUT contains no therapist names that match SETTINGS"
            )

        base_entries = read_unified_base_schedule(source)
        sessions = materialize_sessions_for_date(base_entries, args.date)
        absent_therapist_id, affected = choose_student_smoke_target(
            sessions,
            allowed_therapists=allowed_smoke_therapists,
            max_displaced=args.max_displaced,
        )
        patients = read_patient_registry(source)
        patient_by_id = {patient.patient_id: patient.display_name for patient in patients}

        spec = build_daily_input_spec(
            target_date=args.date,
            therapist_names=settings.therapist_names,
            patient_names=patient_audit.dropdown_names,
            patient_statuses=settings.patient_statuses,
            timeslots=settings.standard_timeslots,
        )
        preview, _ = create_daily_input_preview(
            source, output, spec, overwrite=args.overwrite
        )
        output_created = True

        Win32ComStudentRegistrySchemaBackend().ensure_schema(preview)
        if not student_registry_schema_present(preview):
            raise StudentReplacementSmokeError("STUDENTS schema verification failed")

        _write_smoke_rows(
            preview,
            target_date=args.date,
            absent_therapist_id=absent_therapist_id,
        )

        smoke_students = read_students(preview)
        smoke_student = next(
            (student for student in smoke_students if student.student_id == SMOKE_STUDENT_ID),
            None,
        )
        if smoke_student is None:
            raise StudentReplacementSmokeError("Synthetic smoke student did not read back")
        if not smoke_student.is_active(args.date):
            raise StudentReplacementSmokeError("Synthetic smoke student is not active")
        if not smoke_student.replacement_capable:
            raise StudentReplacementSmokeError("Synthetic smoke student is not replacement-capable")

        daily = read_daily_input(
            preview,
            patients=patients,
            sessions=sessions,
            therapist_names=settings.therapist_names,
        )
        if not any(
            absence.absence_kind == AbsenceKind.THERAPIST
            and absence.subject_id == absent_therapist_id
            for absence in daily.absences
        ):
            raise StudentReplacementSmokeError("Therapist absence did not read back")

        output_contract = assert_application_contract(preview)
    except (
        ValueError,
        ApplicationContractError,
        DailyInputSheetError,
        DailyInputReadError,
        StudentRegistryError,
        StudentRegistrySchemaError,
        StudentReplacementSmokeError,
    ) as exc:
        if output_created:
            try:
                output.unlink(missing_ok=True)
            except OSError:
                pass
        print(f"SAFETY STOP: {exc}")
        return 2

    source_unchanged = file_sha256(source) == source_hash
    print("STUDENT REPLACEMENT SMOKE INPUT OK")
    print(f"Date: {args.date.isoformat()}")
    print(f"Smoke student: {SMOKE_STUDENT_ID} | {SMOKE_STUDENT_NAME}")
    print(f"Absent therapist: {absent_therapist_id}")
    print(f"Displaced sessions: {len(affected)}")
    for session in affected:
        patient_name = patient_by_id.get(session.patient_id, session.patient_id)
        print(f"  {session.start_time.strftime('%H:%M')} | {patient_name}")
    print(f"Preview: {output}")
    print(f"Source unchanged: {source_unchanged}")
    print(f"Application contract source: OK ({len(source_contract.component_names)} components)")
    print(f"Application contract output: OK ({len(output_contract.component_names)} components)")
    print("NEXT: run apply_global_daily_replacements.py on this smoke input.")
    return 0 if source_unchanged else 3


if __name__ == "__main__":
    raise SystemExit(main())
