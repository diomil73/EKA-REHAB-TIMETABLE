from __future__ import annotations

import argparse
import sys
import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from openpyxl import load_workbook  # noqa: E402

from rehab_config.provider_policy_store import load_policy_book  # noqa: E402
from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.models import (  # noqa: E402
    AbsenceKind,
    ReplacementAssignment,
    Therapist,
)
from rehab_core.replacement_policy import (  # noqa: E402
    find_policy_replacement_options,
    validate_policy_replacement_time,
)
from rehab_core.replacements import create_replacement_assignment  # noqa: E402
from rehab_core.therapist_absence_queue import (  # noqa: E402
    build_therapist_absence_replacement_queue,
)
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.native_excel import apply_write_plan_to_copy  # noqa: E402
from rehab_excel.patient_centric_preview import (  # noqa: E402
    PatientCentricPreviewError,
    build_patient_centric_preview_plan,
)
from rehab_excel.reader import read_base_schedule, read_patients, read_settings  # noqa: E402


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _is_legacy_student_name(name: str) -> bool:
    compact = _norm(name).replace(" ", "")
    return compact.startswith("ΦΟΙΤ") or compact.startswith("FOIT")


def _read_daily_input_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise DailyInputReadError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise DailyInputReadError(f"Invalid DAILY_INPUT date: {raw!r}")


def _fmt_times(slots: tuple[time, ...]) -> str:
    return ", ".join(slot.strftime("%H:%M") for slot in slots)


def _capacity_text(option) -> str:
    if option.capacity_limit is None or option.capacity_remaining is None:
        return ""
    used = option.capacity_limit - option.capacity_remaining
    return f" | slots {used}/{option.capacity_limit}"


def _has_vba(path: Path) -> bool:
    with ZipFile(path) as archive:
        return "xl/vbaProject.bin" in archive.namelist()


def _ask_rank(option_count: int) -> int | None:
    while True:
        raw = input(
            f"Choose candidate [1-{option_count}], 0 to leave unresolved, or Q to cancel: "
        ).strip()
        if raw.upper() == "Q":
            raise KeyboardInterrupt
        try:
            rank = int(raw)
        except ValueError:
            print("Enter a number or Q.")
            continue
        if rank == 0:
            return None
        if 1 <= rank <= option_count:
            return rank
        print(f"Choose 1-{option_count}, 0, or Q.")


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Read therapist absences from DAILY_INPUT, choose replacements one by one, "
            "recalculate load/availability after every accepted choice, apply provider "
            "policy overrides, and create a new patient-centric Excel preview."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "DAILY_INPUT_PREVIEW.xlsm",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "DAILY_REPLACEMENTS_PREVIEW.xlsm",
    )
    parser.add_argument("--limit", type=int, default=5)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    input_book = args.input.resolve()
    output = args.output.resolve()
    if not input_book.exists():
        print(f"SAFETY STOP: input workbook not found: {input_book}")
        return 2

    try:
        target_date = _read_daily_input_date(input_book)
        patients = read_patients(input_book)
        settings = read_settings(input_book)
        base_entries = read_base_schedule(input_book)
        sessions = materialize_sessions_for_date(base_entries, target_date)
        daily = read_daily_input(input_book, patients=patients, sessions=sessions)
        policy_book = load_policy_book(REPO_ROOT / "Config" / "provider_policy.json")

        therapist_absences = tuple(
            absence
            for absence in daily.absences
            if absence.absence_kind == AbsenceKind.THERAPIST
        )
        if not therapist_absences:
            raise DailyInputReadError("DAILY_INPUT has no therapist absence rows")

        therapists = tuple(
            Therapist(therapist_id=name, display_name=name)
            for name in settings.therapist_names
            if not _is_legacy_student_name(name)
        )
        initial_queue = build_therapist_absence_replacement_queue(
            sessions=sessions,
            absences=daily.absences,
            therapists=therapists,
            patients=patients,
            timeslots=settings.standard_timeslots,
        )
        if not initial_queue.items:
            raise DailyInputReadError(
                "No session needs replacement after applying patient/therapist absences"
            )

        session_by_id = {session.session_id: session for session in sessions}
        patient_by_id = {patient.patient_id: patient for patient in patients}
        pending_ids = [item.session_id for item in initial_queue.items]
        replacements: list[ReplacementAssignment] = []
        unresolved: list[str] = []

        print("SEQUENTIAL THERAPIST ABSENCE REPLACEMENTS")
        print(f"Date: {target_date.isoformat()}")
        print(f"Sessions needing replacement: {len(pending_ids)}")
        if initial_queue.skipped_patient_absent:
            print(
                "Patient-absent sessions excluded: "
                f"{initial_queue.skipped_patient_absent}"
            )
        print(
            "IMPORTANT: ranking is recalculated after every accepted replacement, "
            "so load and occupied timeslots update immediately."
        )
        print(
            "PROVIDER POLICY: replacement exclusions, temporary max load, blocked times, "
            "and Manager/Acting Manager open times are active."
        )

        for index, session_id in enumerate(pending_ids, start=1):
            target = session_by_id[session_id]
            patient = patient_by_id.get(target.patient_id)
            patient_name = patient.display_name if patient is not None else target.patient_id

            options = find_policy_replacement_options(
                target_session=target,
                therapists=therapists,
                sessions=sessions,
                policy_book=policy_book,
                absences=daily.absences,
                patients=patients,
                replacements=replacements,
                requested_time=target.start_time,
                timeslots=settings.standard_timeslots,
            )

            print(
                f"\n{index}/{len(pending_ids)}. {patient_name} | "
                f"{target.therapist_id} {target.start_time.strftime('%H:%M')}"
            )
            if target.robotic:
                print(
                    "   ROBOTIC SESSION: only authoritative robotic-capable providers "
                    "may be selected."
                )
            visible = options[: max(args.limit, 1)]
            if not visible:
                print(
                    "   No feasible provider/timeslot remains after capacity, availability, "
                    "and provider-policy checks."
                )
                unresolved.append(session_id)
                continue

            for rank, option in enumerate(visible, start=1):
                exact = "exact" if option.exact_time_available else "alternative"
                print(
                    f"   {rank}. {option.display_name} | load {option.active_sessions}"
                    f"{_capacity_text(option)} | "
                    f"recommended {option.recommended_time.strftime('%H:%M')} ({exact}) | "
                    f"free: {_fmt_times(option.available_timeslots)}"
                )

            try:
                rank = _ask_rank(len(visible))
            except KeyboardInterrupt:
                print("\nCANCELLED: no preview was written.")
                return 0
            if rank is None:
                unresolved.append(session_id)
                print("   Left unresolved.")
                continue

            option = visible[rank - 1]
            validate_policy_replacement_time(
                provider_id=option.provider_id,
                target_date=target.session_date,
                replacement_time=option.recommended_time,
                policy_book=policy_book,
            )
            replacement = create_replacement_assignment(
                replacement_id=f"daily-input-{len(replacements) + 1}",
                target_session=target,
                replacement_therapist_id=option.provider_id,
                therapists=therapists,
                sessions=sessions,
                absences=daily.absences,
                replacements=replacements,
                replacement_time=option.recommended_time,
                reason="DAILY_INPUT therapist absence",
                patients=patients,
            )
            replacements.append(replacement)
            print(
                f"   ACCEPTED: {option.display_name} "
                f"{option.recommended_time.strftime('%H:%M')}"
            )

        states = build_daily_session_states(
            sessions,
            absences=daily.absences,
            replacements=replacements,
            target_date=target_date,
        )
        preview = build_patient_centric_preview_plan(
            input_book,
            base_entries=base_entries,
            states=states,
            sessions=sessions,
            patients=patients,
            rebuild_multi_member_groups=True,
        )
        report = apply_write_plan_to_copy(
            preview.write_plan,
            output,
            overwrite=args.overwrite,
        )
    except (DailyInputReadError, PatientCentricPreviewError, ValueError, KeyError) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    vba_preserved = _has_vba(input_book) == _has_vba(output) and _has_vba(output)
    print("\nDAILY REPLACEMENTS PREVIEW OK")
    print(f"Input: {report.source_path}")
    print(f"Preview: {report.output_path}")
    print(f"Accepted replacements: {len(replacements)}")
    print(f"Unresolved sessions: {len(unresolved)}")
    for replacement in replacements:
        patient = patient_by_id.get(replacement.patient_id)
        patient_name = patient.display_name if patient is not None else replacement.patient_id
        print(
            f"  {patient_name}: {replacement.original_therapist_id} -> "
            f"{replacement.replacement_therapist_id} "
            f"{replacement.replacement_time.strftime('%H:%M')}"
        )
    if unresolved:
        print("Unresolved session ids: " + ", ".join(unresolved))
    print(f"Input unchanged: {report.source_unchanged}")
    print(f"VBA preserved: {vba_preserved}")
    print(
        "NEXT: open only DAILY_REPLACEMENTS_PREVIEW.xlsm and inspect all affected "
        "original/destination cells."
    )
    return 0 if report.source_unchanged and vba_preserved else 3


if __name__ == "__main__":
    raise SystemExit(main())
