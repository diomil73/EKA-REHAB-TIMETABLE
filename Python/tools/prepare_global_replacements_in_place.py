from __future__ import annotations

import argparse
from datetime import date, datetime
import json
from pathlib import Path
import sys
import unicodedata

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_config.provider_policy_store import load_policy_book  # noqa: E402
from rehab_config.service_hours_store import load_service_hours  # noqa: E402
from rehab_core.base_schedule import materialize_sessions_for_date  # noqa: E402
from rehab_core.daily_state import build_daily_session_states  # noqa: E402
from rehab_core.daily_statistics import calculate_daily_treatment_statistics  # noqa: E402
from rehab_core.double_booking_review import build_double_booking_review_items  # noqa: E402
from rehab_core.global_replacement_end_to_end import build_global_replacement_assignments  # noqa: E402
from rehab_core.models import AbsenceKind, Therapist  # noqa: E402
from rehab_core.reclining_adjacency import build_reclining_adjacency_penalty  # noqa: E402
from rehab_core.service_replacement_options import build_service_id_by_session  # noqa: E402
from rehab_core.therapist_absence_queue import build_therapist_absence_replacement_queue  # noqa: E402
from rehab_excel.application_contract import assert_application_contract  # noqa: E402
from rehab_excel.authoritative_commit import file_sha256  # noqa: E402
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input  # noqa: E402
from rehab_excel.daily_preview_composer import compose_outpatient_daily_plan  # noqa: E402
from rehab_excel.daily_replacement_facts import (
    build_daily_replacement_fact_rows,
    read_daily_replacement_fact_rows,
    read_daily_replacement_facts,
)
from rehab_excel.daily_suspension_resolution import (
    SUSPENSION_REASON,
    SUSPENSION_STATUS,
    build_unresolved_suspensions,
)
from rehab_excel.master_daily_overlay import build_master_daily_overlay_plan, compose_master_daily_overlay
from rehab_excel.outpatient_schedule_source import read_unified_base_schedule
from rehab_excel.patient_centric_preview import build_patient_centric_preview_plan
from rehab_excel.patient_registry_source import read_patient_registry
from rehab_excel.reader import read_settings
from rehab_excel.student_registry import read_students
from rehab_excel.writeback import CellPatch, validate_write_plan


class GlobalReplacementPreparationError(RuntimeError):
    pass


def _norm(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", str(value).strip().upper())
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def _is_legacy_student_name(name: str) -> bool:
    compact = _norm(name).replace(" ", "")
    return compact.startswith("ΦΟΙΤ") or compact.startswith("FOIT")


def _target_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise GlobalReplacementPreparationError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise GlobalReplacementPreparationError(f"Invalid DAILY_INPUT date: {raw!r}")


def _replacement_id_factory(history_rows, session_by_id):
    used = {str(row.replacement_id).strip() for row in history_rows}
    def allocate(session_id: str) -> str:
        target = session_by_id[session_id]
        base = f"daily-global:{target.session_id}:{target.session_date.isoformat()}"
        if base not in used:
            used.add(base)
            return base
        version = 2
        while f"{base}:v{version}" in used:
            version += 1
        value = f"{base}:v{version}"
        used.add(value)
        return value
    return allocate


def _escape(value: object | None) -> str:
    text = "" if value is None else str(value)
    return text.replace("\\", "\\\\").replace("\t", "\\t").replace("\r", "\\r").replace("\n", "\\n")


def _line(kind: str, values: tuple[object | None, ...]) -> str:
    return "\t".join([kind, *(_escape(value) for value in values)])


def _bool_token(value: bool | None) -> str:
    if value is None:
        return ""
    return "1" if value else "0"


def _text_runs_token(patch: CellPatch) -> str:
    return ";".join(": ".join(()) for _ in ()) if False else ";".join(
        ":".join((str(run.start), str(run.length), "1" if run.strike_through else "0", "1" if run.italic else "0", run.font_role or ""))
        for run in patch.text_runs
    )


def _patch_line(patch: CellPatch) -> str:
    min_font_size = None if patch.text_runs else patch.min_font_size
    return _line("PATCH", (patch.sheet, patch.cell, patch.intent.value, patch.value, "1" if patch.strike_through else "0", patch.font_role, patch.fill_role, patch.border_role, _bool_token(patch.wrap_text), min_font_size, patch.source_tag, _text_runs_token(patch)))


def _patient_display(patient) -> str:
    patient_type = str(getattr(getattr(patient, "patient_type", None), "value", "")).casefold()
    return f"ΕΞ | {patient.display_name}" if patient_type == "outpatient" else patient.display_name


def _write_response(path: Path | None, payload: dict[str, object]) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def prepare_global_replacements_payload(workbook_path: str | Path, payload_path: str | Path) -> dict[str, object]:
    source = Path(workbook_path).resolve()
    payload = Path(payload_path).resolve()
    if not source.exists():
        raise GlobalReplacementPreparationError(f"Workbook not found: {source}")
    if source.suffix.casefold() != ".xlsm":
        raise GlobalReplacementPreparationError("Workbook must be an .xlsm file")

    source_hash = file_sha256(source)
    contract = assert_application_contract(source)
    target_date = _target_date(source)
    patients = read_patient_registry(source)
    patient_by_id = {patient.patient_id: patient for patient in patients}
    students = tuple(read_students(source))
    settings = read_settings(source)
    base_entries = read_unified_base_schedule(source)
    sessions = materialize_sessions_for_date(base_entries, target_date)
    session_by_id = {session.session_id: session for session in sessions}
    daily = read_daily_input(source, patients=patients, sessions=sessions, therapist_names=settings.therapist_names)
    if not any(item.absence_kind == AbsenceKind.THERAPIST for item in daily.absences):
        raise DailyInputReadError("DAILY_INPUT has no therapist absence rows")

    policy_book = load_policy_book(REPO_ROOT / "Config" / "provider_policy.json")
    service_hours = load_service_hours(REPO_ROOT / "Config" / "service_hours.json")
    adjacency_penalty = build_reclining_adjacency_penalty(base_entries)
    service_id_by_session = build_service_id_by_session(base_entries, target_date)
    existing = tuple(read_daily_replacement_facts(source, sessions=sessions, target_date=target_date))
    history = tuple(read_daily_replacement_fact_rows(source, target_date=target_date, include_restored=True))
    therapists = tuple(Therapist(therapist_id=name, display_name=name) for name in settings.therapist_names if not _is_legacy_student_name(name))
    queue = build_therapist_absence_replacement_queue(sessions=sessions, absences=daily.absences, cancellations=daily.cancellations, replacements=existing, therapists=therapists, patients=patients, timeslots=settings.standard_timeslots, students=students, policy_book=policy_book, base_entries=base_entries)
    targets = tuple(session_by_id[item.session_id] for item in queue.items)
    if not targets:
        payload.parent.mkdir(parents=True, exist_ok=True)
        payload.write_text("", encoding="utf-8-sig")
        return {"ok": True, "mode": "already_resolved", "date": target_date.isoformat(), "normal_assignments": 0, "review_required": False, "review_items": 0, "suspensions": 0, "patch_count": 0, "application_contract_components": len(contract.component_names), "source_unchanged": True}

    result = build_global_replacement_assignments(targets, therapists=therapists, sessions=sessions, absences=daily.absences, patients=patients, replacements=existing, cancellations=daily.cancellations, timeslots=settings.standard_timeslots, students=students, adjacency_penalty=adjacency_penalty, policy_book=policy_book, base_entries=base_entries, service_hours=service_hours, service_id_by_session=service_id_by_session, replacement_id_factory=_replacement_id_factory(history, session_by_id), reason="DAILY_INPUT global therapist absence plan")
    review_items = build_double_booking_review_items(result, sessions=sessions, therapists=therapists, patients=patients, absences=daily.absences, replacements=existing, cancellations=daily.cancellations, timeslots=settings.standard_timeslots, policy_book=policy_book, base_entries=base_entries)
    reviewable = tuple(item for item in review_items if item.has_options)

    replacement_rows = build_daily_replacement_fact_rows(result.assignments, sessions=sessions)
    lines: list[str] = []
    for row in replacement_rows:
        lines.append(_line("REPLACEMENT", (row.replacement_id, row.target_session_id, row.patient_id, row.original_therapist_id, row.replacement_therapist_id, row.replacement_date.isoformat(), row.replacement_time.strftime("%H:%M"), row.reason, row.provider_kind.value, row.status)))

    if reviewable:
        mode = "review_required"
        suspensions = ()
        patch_count = 0
    else:
        mode = "resolved"
        suspensions = build_unresolved_suspensions(result.unresolved_request_ids, sessions=sessions, target_date=target_date)
        for suspension in suspensions:
            session = session_by_id[suspension.target_session_id]
            patient = patient_by_id[session.patient_id]
            lines.append(_line("SUSPENSION", (suspension.cancellation_id, suspension.target_session_id, _patient_display(patient), session.start_time.strftime("%H:%M"), SUSPENSION_STATUS, SUSPENSION_REASON)))
        all_replacements = existing + result.assignments
        all_cancellations = daily.cancellations + suspensions
        states = build_daily_session_states(sessions, absences=daily.absences, replacements=all_replacements, cancellations=all_cancellations, target_date=target_date)
        statistics = calculate_daily_treatment_statistics(states)
        if not statistics.is_finalized:
            raise GlobalReplacementPreparationError(f"Global replacement plan remains unresolved: {statistics.unresolved_sessions}")
        preview = build_patient_centric_preview_plan(source, base_entries=base_entries, states=states, sessions=sessions, patients=patients, students=students, rebuild_multi_member_groups=True)
        plan = compose_outpatient_daily_plan(preview.write_plan, states=states, patients=patients)
        master_overlay = build_master_daily_overlay_plan(source, base_entries=base_entries, states=states, patients=patients)
        plan = compose_master_daily_overlay(plan, master_overlay)
        validate_write_plan(plan)
        for patch in plan.patches:
            lines.append(_patch_line(patch))
        patch_count = len(plan.patches)

    payload.parent.mkdir(parents=True, exist_ok=True)
    payload.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8-sig")
    if file_sha256(source) != source_hash:
        try:
            payload.unlink()
        except OSError:
            pass
        raise GlobalReplacementPreparationError("Workbook changed during replacement validation")

    return {"ok": True, "mode": mode, "date": target_date.isoformat(), "normal_assignments": len(result.assignments), "review_required": bool(reviewable), "review_items": len(reviewable), "suspensions": len(suspensions), "patch_count": patch_count, "application_contract_components": len(contract.component_names), "source_unchanged": True}


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare in-place global therapist-absence resolution.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--payload", required=True, type=Path)
    parser.add_argument("--response", type=Path)
    args = parser.parse_args()
    try:
        result = prepare_global_replacements_payload(args.input, args.payload)
    except Exception as exc:
        try:
            if args.payload.exists(): args.payload.unlink()
        except OSError:
            pass
        _write_response(args.response, {"ok": False, "error": str(exc)})
        print(f"SAFETY STOP: {exc}")
        return 2
    _write_response(args.response, result)
    print("GLOBAL REPLACEMENT IN-PLACE PAYLOAD VERIFIED")
    print(f"Date: {result['date']}")
    print(f"Mode: {result['mode']}")
    print(f"Normal assignments: {result['normal_assignments']}")
    print(f"Review required: {result['review_required']}")
    print(f"Suspensions: {result['suspensions']}")
    print(f"Patches: {result['patch_count']}")
    print("Input unchanged: True")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
