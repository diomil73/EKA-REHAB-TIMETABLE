from __future__ import annotations

import argparse
from datetime import date
from hashlib import sha256
from pathlib import Path
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from rehab_excel.master_projection_writer import (  # noqa: E402
    MasterProjectionWriteError,
    refresh_master_projection_in_place,
)
from rehab_excel.master_toolbar_vba import (  # noqa: E402
    MasterToolbarError,
    finalize_user_navigation,
)
from rehab_excel.master_visual_style import (  # noqa: E402
    MasterVisualStyleError,
    apply_master_visual_style_in_place,
)


REPO_ROOT = Path(__file__).resolve().parents[2]
TOOLS_DIR = REPO_ROOT / "Python" / "tools"

STEPS = (
    ("registration menu + forms", "build_registration_menu_preview.py"),
    ("patient registration bridge", "build_patient_registration_bridge_preview.py"),
    ("outpatient schedule form", "build_outpatient_schedule_form_preview.py"),
    ("DAILY_INPUT sheet", "preview_daily_input_sheet.py"),
)


class UnifiedBuildError(RuntimeError):
    pass


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value.strip())
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Date must be YYYY-MM-DD") from exc


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _has_vba(path: Path) -> bool:
    try:
        with ZipFile(path) as archive:
            return "xl/vbaProject.bin" in archive.namelist()
    except Exception:
        return False


def _run_step(label: str, args: list[str]) -> None:
    completed = subprocess.run(
        args,
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        errors="replace",
    )
    if completed.returncode == 0:
        print(f"[OK] {label}")
        return

    detail = (completed.stderr or completed.stdout or "").strip()
    raise UnifiedBuildError(
        f"{label} failed with exit code {completed.returncode}"
        + (f":\n{detail}" if detail else "")
    )


def _verify_final_workbook(path: Path) -> None:
    if not path.exists():
        raise UnifiedBuildError("Unified output workbook was not created")
    if not _has_vba(path):
        raise UnifiedBuildError("Unified output does not contain a VBA project")

    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise UnifiedBuildError("Unified output has no DAILY_INPUT sheet")
    finally:
        wb.close()


def build_unified_preview(
    source: Path,
    output: Path,
    *,
    target_date: date,
    overwrite: bool = False,
) -> Path:
    source = source.resolve()
    output = output.resolve()

    if not source.exists():
        raise UnifiedBuildError(f"Source workbook not found: {source}")
    if source == output:
        raise UnifiedBuildError("Output must be different from source workbook")
    if source.suffix.casefold() != ".xlsm" or output.suffix.casefold() != ".xlsm":
        raise UnifiedBuildError("Source and output must both be .xlsm files")
    if output.exists() and not overwrite:
        raise UnifiedBuildError(f"Output already exists: {output}")

    source_before = _sha256(source)
    output.parent.mkdir(parents=True, exist_ok=True)

    if output.exists():
        try:
            output.unlink()
        except PermissionError as exc:
            raise UnifiedBuildError(
                f"Output workbook is open or locked by Excel: {output}"
            ) from exc

    with tempfile.TemporaryDirectory(prefix="eka_unified_", dir=output.parent) as temp_dir:
        temp = Path(temp_dir)
        stage1 = temp / "01_menu.xlsm"
        stage2 = temp / "02_patient_bridge.xlsm"
        stage3 = temp / "03_outpatient_form.xlsm"

        _run_step(
            "registration menu + forms",
            [
                sys.executable,
                str(TOOLS_DIR / "build_registration_menu_preview.py"),
                "--source",
                str(source),
                "--output",
                str(stage1),
                "--overwrite",
            ],
        )
        _run_step(
            "patient registration bridge",
            [
                sys.executable,
                str(TOOLS_DIR / "build_patient_registration_bridge_preview.py"),
                "--source",
                str(stage1),
                "--output",
                str(stage2),
                "--overwrite",
            ],
        )
        _run_step(
            "outpatient schedule form",
            [
                sys.executable,
                str(TOOLS_DIR / "build_outpatient_schedule_form_preview.py"),
                "--source",
                str(stage2),
                "--output",
                str(stage3),
                "--overwrite",
            ],
        )
        _run_step(
            "DAILY_INPUT sheet",
            [
                sys.executable,
                str(TOOLS_DIR / "preview_daily_input_sheet.py"),
                "--date",
                target_date.isoformat(),
                "--source",
                str(stage3),
                "--output",
                str(output),
                "--overwrite",
            ],
        )

        # MASTER content must be rebuilt before navigation and visual decoration.
        # Styling alone is never allowed to reorder or move patient rows.
        try:
            projected = refresh_master_projection_in_place(output)
            print(f"[OK] MASTER projection refresh ({projected} patients)")
        except MasterProjectionWriteError as exc:
            raise UnifiedBuildError(f"MASTER projection refresh failed: {exc}") from exc

        try:
            finalize_user_navigation(output)
            print("[OK] final user navigation")
        except MasterToolbarError as exc:
            raise UnifiedBuildError(f"final user navigation failed: {exc}") from exc

        try:
            apply_master_visual_style_in_place(output)
            print("[OK] MASTER visual baseline")
        except MasterVisualStyleError as exc:
            raise UnifiedBuildError(f"MASTER visual baseline failed: {exc}") from exc

    if _sha256(source) != source_before:
        output.unlink(missing_ok=True)
        raise UnifiedBuildError("Original source workbook changed during unified build")

    try:
        _verify_final_workbook(output)
    except Exception:
        output.unlink(missing_ok=True)
        raise

    return output


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Build one safe XLSM containing the central menu, registration forms, "
            "patient bridge, outpatient schedule form, DAILY_INPUT action and DAILY_INPUT sheet."
        )
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=REPO_ROOT / "Rehab_Center_System_v27_1.xlsm",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "UNIFIED_APP_SMOKE.xlsm",
    )
    parser.add_argument("--date", type=_parse_date, required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    try:
        output = build_unified_preview(
            args.source,
            args.output,
            target_date=args.date,
            overwrite=args.overwrite,
        )
    except UnifiedBuildError as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("UNIFIED PREVIEW OK")
    print(f"Source: {args.source.resolve()}")
    print(f"Preview: {output}")
    print(f"DAILY_INPUT date: {args.date.isoformat()}")
    print("Source unchanged: True")
    print("VBA project present: True")
    print("DAILY_INPUT sheet present: True")
    print("NEXT: open only the unified preview and run ShowRegistrationMenu.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
