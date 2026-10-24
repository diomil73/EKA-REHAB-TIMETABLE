from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from openpyxl import load_workbook  # noqa: E402

from rehab_core.models import ReplacementProviderKind  # noqa: E402
from rehab_excel.application_contract import (  # noqa: E402
    ApplicationContractError,
    assert_application_contract,
)
from rehab_excel.daily_replacement_facts import (  # noqa: E402
    DailyReplacementFactError,
    read_daily_replacement_fact_rows,
)
from rehab_excel.student_registry import StudentRegistryError, read_students  # noqa: E402


SMOKE_STUDENT_ID = "0000-STUDENT-SMOKE"


class StudentReplacementSmokeVerificationError(RuntimeError):
    pass


def _read_daily_input_date(path: Path) -> date:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            raise StudentReplacementSmokeVerificationError("Workbook has no DAILY_INPUT sheet")
        raw = wb["DAILY_INPUT"]["B2"].value
    finally:
        wb.close()

    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise StudentReplacementSmokeVerificationError(
        f"Invalid DAILY_INPUT date: {raw!r}"
    )


def _therapist_daily_contains(path: Path, text: str) -> bool:
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    try:
        if "THERAPIST_DAILY" not in wb.sheetnames:
            raise StudentReplacementSmokeVerificationError(
                "Workbook has no THERAPIST_DAILY sheet"
            )
        ws = wb["THERAPIST_DAILY"]
        needle = text.strip().casefold()
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is None:
                    continue
                if needle in str(value).strip().casefold():
                    return True
        return False
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Verify that the student replacement smoke produced an authoritative "
            "STUDENT replacement fact and visible THERAPIST_DAILY projection."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=REPO_ROOT / "Excel" / "previews" / "STUDENT_REPLACEMENT_RESULT.xlsm",
    )
    parser.add_argument("--student-id", default=SMOKE_STUDENT_ID)
    args = parser.parse_args()

    workbook_path = args.input.resolve()
    if not workbook_path.exists():
        print(f"SAFETY STOP: workbook not found: {workbook_path}")
        return 2

    try:
        contract = assert_application_contract(workbook_path)
        target_date = _read_daily_input_date(workbook_path)
        students = read_students(workbook_path)
        student = next(
            (item for item in students if item.student_id == args.student_id),
            None,
        )
        if student is None:
            raise StudentReplacementSmokeVerificationError(
                f"Student {args.student_id!r} is missing from STUDENTS"
            )

        rows = read_daily_replacement_fact_rows(
            workbook_path,
            target_date=target_date,
            include_restored=False,
        )
        matches = tuple(
            row
            for row in rows
            if row.replacement_therapist_id == args.student_id
            and row.provider_kind == ReplacementProviderKind.STUDENT
            and str(row.status).strip().upper() == "ACTIVE"
        )
        if not matches:
            raise StudentReplacementSmokeVerificationError(
                "No ACTIVE STUDENT replacement fact was found for the smoke student"
            )

        if not _therapist_daily_contains(workbook_path, student.display_name):
            raise StudentReplacementSmokeVerificationError(
                "Student replacement fact exists, but THERAPIST_DAILY does not show the student"
            )
    except (
        ApplicationContractError,
        DailyReplacementFactError,
        StudentRegistryError,
        StudentReplacementSmokeVerificationError,
        ValueError,
    ) as exc:
        print(f"SAFETY STOP: {exc}")
        return 2

    print("STUDENT REPLACEMENT SMOKE VERIFIED")
    print(f"Date: {target_date.isoformat()}")
    print(f"Student: {student.student_id} | {student.display_name}")
    print(f"ACTIVE student replacement facts: {len(matches)}")
    for row in matches:
        print(
            f"  {row.patient_id}: {row.original_therapist_id} -> "
            f"{row.replacement_therapist_id} {row.replacement_time.strftime('%H:%M')}"
        )
    print("THERAPIST_DAILY projection: OK")
    print(f"Application contract: OK ({len(contract.component_names)} components)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
