from __future__ import annotations

from pathlib import Path
import sys

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from openpyxl import load_workbook


def main() -> int:
    path = REPO_ROOT / "Excel" / "previews" / "DAILY_INPUT_PREVIEW.xlsm"
    if not path.exists():
        print(f"SAFETY STOP: workbook not found: {path}")
        return 2
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=True)
    try:
        if "DAILY_INPUT" not in wb.sheetnames:
            print("SAFETY STOP: workbook has no DAILY_INPUT sheet")
            return 2
        ws = wb["DAILY_INPUT"]
        print("DAILY INPUT THERAPIST DIAGNOSTIC")
        print(f"Workbook: {path}")
        print(f"Date B2: {ws['B2'].value!r}")
        print("Therapist input rows 7-16:")
        found = 0
        for row in range(7, 17):
            values = tuple(ws.cell(row, col).value for col in range(1, 7))
            if any(value not in (None, "") for value in values):
                found += 1
                print(f"  row {row}: {values!r}")
        if found == 0:
            print("  NONE")
            print("RESULT: no therapist entry was saved in DAILY_INPUT_PREVIEW.xlsm rows 7-16.")
        else:
            print(f"RESULT: {found} non-empty therapist row(s) found.")
        return 0
    finally:
        wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
