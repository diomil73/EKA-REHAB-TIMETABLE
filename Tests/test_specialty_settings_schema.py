from datetime import time
from pathlib import Path
import sys

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "Python"))

from rehab_excel.reader import read_settings


def _write_settings_workbook(path: Path, *, columns: int) -> None:
    wb = Workbook()

    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    planner.append(["PatientID", "Ασθενής"])

    settings = wb.create_sheet("SETTINGS")
    headers = [
        "THERAPISTS_FTH",
        "HOURS_STANDARD",
        "HOURS_RECLINED",
        "DAY_COMBINATIONS",
        "THERAPIES",
        "PATIENT_STATUS",
        "YES_NO",
        "ROOMS",
        "PSYCHOLOGISTS",
        "OCCUPATIONAL_THERAPISTS",
        "SPEECH_THERAPISTS",
    ][:columns]
    settings.append(headers)

    values = [
        "Φυσικοθεραπευτής Α",
        time(9, 15),
        time(9, 0),
        "Καθ/να",
        "ΦΘ",
        "ΠΑΡΩΝ",
        "ΝΑΙ",
        "101",
        "Ψυχολόγος Α",
        "Εργοθεραπευτής Α",
        "Λογοθεραπευτής Α",
    ][:columns]
    settings.append(values)
    wb.save(path)


def test_legacy_eight_column_settings_keep_specialty_registries_empty(tmp_path):
    path = tmp_path / "legacy.xlsx"
    _write_settings_workbook(path, columns=8)

    settings = read_settings(path)

    assert settings.psychologist_names == ()
    assert settings.occupational_therapist_names == ()
    assert settings.speech_therapist_names == ()


def test_psychology_only_schema_keeps_new_registries_empty(tmp_path):
    path = tmp_path / "psych_only.xlsx"
    _write_settings_workbook(path, columns=9)

    settings = read_settings(path)

    assert settings.psychologist_names == ("Ψυχολόγος Α",)
    assert settings.occupational_therapist_names == ()
    assert settings.speech_therapist_names == ()


def test_extended_specialty_registries_are_read_from_j_and_k(tmp_path):
    path = tmp_path / "extended.xlsx"
    _write_settings_workbook(path, columns=11)

    settings = read_settings(path)

    assert settings.psychologist_names == ("Ψυχολόγος Α",)
    assert settings.occupational_therapist_names == ("Εργοθεραπευτής Α",)
    assert settings.speech_therapist_names == ("Λογοθεραπευτής Α",)
