from __future__ import annotations

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from rehab_core.daily_state import DailySessionStatus, build_daily_session_states
from rehab_core.models import Patient, Session
from rehab_excel.daily_input_reader import read_daily_input


DAY = date(2026, 10, 26)


def _book(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_INPUT"
    ws["A2"] = "Ημερομηνία"
    ws["B2"] = "26/10/2026"
    for col, value in enumerate(
        ("Θεραπευτής", "Όλη ημέρα", "Από", "Έως", "Αιτία", "Σχόλιο"), 1
    ):
        ws.cell(6, col, value)
    for col, value in enumerate(
        ("Ασθενής", "Όλη ημέρα", "Ώρα", "Κατάσταση", "Σχόλιο"), 1
    ):
        ws.cell(20, col, value)
    wb.save(path)
    wb.close()


def test_all_day_patient_absence_affects_every_materialized_session_and_is_reversible(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΓΙΑΝΝΙΑ"
    ws["B21"] = "ΝΑΙ"
    ws["D21"] = "ΑΠΩΝ"
    ws["E21"] = "Διακομιστήριο"
    wb.save(path)
    wb.close()

    patient = Patient("P1", "ΓΙΑΝΝΙΑ")
    sessions = (
        Session("S1", "P1", "Σάκκου", DAY, time(8, 30), treatment="ΦΘ"),
        Session("S2", "P1", "Άλλος", DAY, time(9, 15), treatment="Εργο"),
        Session("S3", "P1", "Τρίτος", DAY, time(10, 0), treatment="ΕΦΑ"),
    )
    daily = read_daily_input(path, patients=(patient,), sessions=sessions)

    assert len(daily.absences) == 1
    assert daily.absences[0].start_time is None
    assert daily.absences[0].end_time is None
    assert "Διακομιστήριο" in (daily.absences[0].reason or "")
    assert len(daily.source_changes) == 1
    change = daily.source_changes[0]
    assert change.source_row == 21
    assert change.change_type == "patient_absence"
    assert change.all_day is True
    assert change.change_id == "daily-input:2026-10-26:patient:21"

    states = build_daily_session_states(
        sessions,
        absences=daily.absences,
        target_date=DAY,
    )
    assert {state.status for state in states} == {DailySessionStatus.PATIENT_ABSENT}
    assert {state.session_id for state in states} == {"S1", "S2", "S3"}


def test_restore_reader_can_ignore_exact_source_row_and_allow_no_remaining_changes(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΓΙΑΝΝΙΑ"
    ws["B21"] = "ΝΑΙ"
    ws["D21"] = "ΑΠΩΝ"
    wb.save(path)
    wb.close()

    patient = Patient("P1", "ΓΙΑΝΝΙΑ")
    sessions = (Session("S1", "P1", "Σάκκου", DAY, time(8, 30)),)
    daily = read_daily_input(
        path,
        patients=(patient,),
        sessions=sessions,
        ignore_rows=(21,),
        allow_empty=True,
    )
    assert daily.absences == ()
    assert daily.cancellations == ()
    assert daily.source_changes == ()
