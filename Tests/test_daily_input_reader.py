from __future__ import annotations

from datetime import date, time
from pathlib import Path
import sys

from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.models import AbsenceKind, Patient, Session  # noqa: E402
from rehab_excel.daily_input_reader import read_daily_input  # noqa: E402


def _book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_INPUT"
    ws["A2"] = "Ημερομηνία"
    ws["B2"] = "18/09/2026"
    for c, value in enumerate(("Θεραπευτής", "Όλη ημέρα", "Από", "Έως", "Αιτία", "Σχόλιο"), 1):
        ws.cell(6, c, value)
    for c, value in enumerate(("Ασθενής", "Όλη ημέρα", "Ώρα", "Κατάσταση", "Σχόλιο"), 1):
        ws.cell(20, c, value)
    wb.save(path)
    wb.close()


def _fixtures():
    patients = [Patient("P1", "ΠΕΤΙΡΟΠΟΥΛΟΣ")]
    sessions = [Session("S1", "P1", "Γαβράς", date(2026, 9, 18), time(8, 30))]
    return patients, sessions


def test_reads_patient_timeslot_absence(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΠΕΤΙΡΟΠΟΥΛΟΣ"
    ws["B21"] = "ΟΧΙ"
    ws["C21"] = "08:30"
    ws["D21"] = "ΕΚΤΟΣ ΚΛΙΝΙΚΗΣ"
    ws["E21"] = "TEST"
    wb.save(path)
    wb.close()

    patients, sessions = _fixtures()
    result = read_daily_input(path, patients=patients, sessions=sessions)
    assert result.target_date == date(2026, 9, 18)
    assert result.patient_rows_used == 1
    assert len(result.absences) == 1
    absence = result.absences[0]
    assert absence.absence_kind == AbsenceKind.PATIENT
    assert absence.subject_id == "P1"
    assert absence.start_time == time(8, 30)
    assert absence.end_time == time(8, 31)
    assert "TEST" in (absence.reason or "")


def test_reads_excel_serial_time(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΠΕΤΙΡΟΠΟΥΛΟΣ"
    ws["C21"] = 8.5 / 24.0
    ws["D21"] = "ΑΠΩΝ"
    wb.save(path)
    wb.close()

    patients, sessions = _fixtures()
    result = read_daily_input(path, patients=patients, sessions=sessions)
    assert result.absences[0].start_time == time(8, 30)


def test_single_therapist_slot_does_not_require_fake_range(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A7"] = "Γαβράς"
    ws["B7"] = "ΟΧΙ"
    ws["C7"] = "08:30"
    ws["E7"] = "Απουσία"
    wb.save(path)
    wb.close()

    patients, sessions = _fixtures()
    result = read_daily_input(path, patients=patients, sessions=sessions)
    absence = result.absences[0]
    assert absence.absence_kind == AbsenceKind.THERAPIST
    assert absence.start_time == time(8, 30)
    assert absence.end_time == time(8, 31)
