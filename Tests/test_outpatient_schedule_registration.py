from __future__ import annotations

from datetime import time
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
PYTHON_ROOT = REPO_ROOT / "Python"
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

from rehab_core.models import Patient, PatientType  # noqa: E402
from rehab_excel.outpatient_schedule_registration import (  # noqa: E402
    OutpatientScheduleRequest,
    OutpatientScheduleWriteError,
    create_outpatient_schedule_preview,
    validate_outpatient_schedule_request,
)


def outpatient(patient_id: str = "P-OUT") -> Patient:
    return Patient(
        patient_id=patient_id,
        display_name="Εξωτερικός Ασθενής",
        patient_type=PatientType.OUTPATIENT,
    )


def test_request_requires_registered_outpatient():
    request = OutpatientScheduleRequest(
        patient_id="P-IN",
        treatment="ΦΘ",
        start_time=time(9, 0),
        day_pattern="Δε-Τε-Πα",
        therapist_id="T1",
    )

    with pytest.raises(OutpatientScheduleWriteError, match="not registered as outpatient"):
        validate_outpatient_schedule_request(
            request,
            patients=[Patient("P-IN", "Εσωτερικός")],
        )


def test_request_rejects_robotic_treatment():
    request = OutpatientScheduleRequest(
        patient_id="P-OUT",
        treatment="Ρομποτικό",
        start_time=time(9, 0),
        day_pattern="Δε-Τε-Πα",
        therapist_id="T1",
    )

    with pytest.raises(OutpatientScheduleWriteError, match="cannot have robotic"):
        validate_outpatient_schedule_request(request, patients=[outpatient()])


def _baseline(path: Path, *, second_outpatient: bool = False) -> None:
    wb = Workbook()
    patients = wb.active
    patients.title = "PATIENTS"
    for col, value in enumerate(
        ("PatientID", "Δωμάτιο", "Ασθενής", "Λοιμώδης", "Κατάσταση", "PatientType", "HospitalMRN"),
        start=1,
    ):
        patients.cell(1, col, value)
    patients.append(["P-OUT", "", "Εξωτερικός Ασθενής", "", "", "Εξωτερικός", "MRN1"])
    if second_outpatient:
        patients.append(["P-OTHER", "", "Άλλος Εξωτερικός", "", "", "Εξωτερικός", "MRN2"])

    schedule = wb.create_sheet("OUTPATIENT_SCHEDULE")
    schedule.append(["PatientID", "Ασθενής", "Θεραπεία", "Ώρα", "Ημέρες", "Θεραπευτής"])
    schedule.append(["P-OUT", "Εξωτερικός Ασθενής", "ΦΘ", time(9, 0), "Δε-Τε-Πα", "T1"])
    wb.save(path)
    wb.close()


class FakeBackend:
    def upsert(self, workbook_path, request, *, patient_name):
        wb = load_workbook(workbook_path, keep_vba=True)
        try:
            ws = wb["OUTPATIENT_SCHEDULE"]
            if request.target_base_entry_id:
                row = int(request.target_base_entry_id.split(":", 1)[1])
                created = False
            else:
                row = ws.max_row + 1
                created = True
            ws.cell(row, 1, request.patient_id.strip())
            ws.cell(row, 2, patient_name)
            ws.cell(row, 3, request.treatment.strip())
            ws.cell(row, 4, request.start_time)
            ws.cell(row, 5, request.day_pattern.strip())
            ws.cell(row, 6, (request.therapist_id or "").strip())
            wb.save(workbook_path)
            return row, created
        finally:
            wb.close()


def test_preview_append_is_verified_and_source_stays_unchanged(tmp_path):
    source = tmp_path / "source.xlsm"
    output = tmp_path / "preview.xlsm"
    _baseline(source)
    before = source.read_bytes()

    report = create_outpatient_schedule_preview(
        source,
        output,
        OutpatientScheduleRequest(
            patient_id="P-OUT",
            treatment="Πισίνα",
            start_time=time(10, 0),
            day_pattern="Τρ-Πε",
            therapist_id="T2",
        ),
        backend=FakeBackend(),
    )

    assert report.created is True
    assert report.base_entry_id == "outpatient:3"
    assert report.verified_in_output is True
    assert report.source_unchanged is True
    assert source.read_bytes() == before


def test_update_cannot_move_authoritative_row_to_another_patient(tmp_path):
    source = tmp_path / "source.xlsm"
    output = tmp_path / "preview.xlsm"
    _baseline(source, second_outpatient=True)

    with pytest.raises(OutpatientScheduleWriteError, match="Cannot move"):
        create_outpatient_schedule_preview(
            source,
            output,
            OutpatientScheduleRequest(
                patient_id="P-OTHER",
                treatment="ΦΘ",
                start_time=time(11, 0),
                day_pattern="Δε-Πα",
                therapist_id="T2",
                target_base_entry_id="outpatient:2",
            ),
            backend=FakeBackend(),
        )
