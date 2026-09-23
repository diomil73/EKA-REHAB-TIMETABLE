from datetime import date, time
from pathlib import Path
import sys

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "Python"))

from rehab_core.availability import is_patient_available
from rehab_core.base_schedule import materialize_sessions_for_date
from rehab_excel.reader import read_base_schedule, read_settings


MONDAY = date(2026, 9, 21)


def _write_workbook(path: Path, *, with_psychology: bool) -> None:
    wb = Workbook()

    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    patients.append(["P1", "101", "ΔΟΚΙΜΗ ΑΣΘΕΝΗ", "ΟΧΙ", "ΠΑΡΩΝ"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    headers = [
        "PatientID",
        "Θάλαμος",
        "Ασθενής",
        "Μολυσματικός",
        "Κατάσταση",
        "ΦΘ_Ώρα",
        "ΦΘ_Ημέρες",
        "ΦΘ_Θεραπευτής",
        "Σχόλια 1",
        "Σχόλια 2",
    ]
    if with_psychology:
        headers.extend(["Ψυχ_Ώρα", "Ψυχ_Ημέρες", "Ψυχ_Ψυχολόγος"])
    planner.append(headers)

    row = ["P1", "101", "ΔΟΚΙΜΗ ΑΣΘΕΝΗ", "ΟΧΙ", "ΠΑΡΩΝ", None, None, None, None, None]
    if with_psychology:
        row.extend([time(10, 0), "Καθ/να", "Ψυχολόγος Α"])
    planner.append(row)

    settings = wb.create_sheet("SETTINGS")
    setting_headers = [
        "THERAPISTS_FTH",
        "HOURS_STANDARD",
        "HOURS_RECLINED",
        "DAY_COMBINATIONS",
        "THERAPIES",
        "PATIENT_STATUS",
        "YES_NO",
        "ROOMS",
    ]
    if with_psychology:
        setting_headers.append("PSYCHOLOGISTS")
    settings.append(setting_headers)
    setting_row = ["Θεραπευτής Α", time(9, 15), time(9, 0), "Καθ/να", "ΦΘ", "ΠΑΡΩΝ", "ΝΑΙ", "101"]
    if with_psychology:
        setting_row.append("Ψυχολόγος Α")
    settings.append(setting_row)

    wb.save(path)


def test_old_eight_column_settings_remain_compatible(tmp_path):
    path = tmp_path / "old_schema.xlsx"
    _write_workbook(path, with_psychology=False)

    settings = read_settings(path)

    assert settings.therapist_names == ("Θεραπευτής Α",)
    assert settings.psychologist_names == ()


def test_psychology_registry_and_patient_planner_block_are_imported(tmp_path):
    path = tmp_path / "psychology_schema.xlsx"
    _write_workbook(path, with_psychology=True)

    settings = read_settings(path)
    entries = read_base_schedule(path)

    assert settings.psychologist_names == ("Ψυχολόγος Α",)
    psychology = [entry for entry in entries if entry.treatment == "Ψυχ"]
    assert len(psychology) == 1
    assert psychology[0].patient_id == "P1"
    assert psychology[0].start_time == time(10, 0)
    assert psychology[0].day_pattern == "Καθ/να"
    assert psychology[0].therapist_id == "Ψυχολόγος Α"


def test_psychology_session_blocks_same_patient_time(tmp_path):
    path = tmp_path / "psychology_conflict.xlsx"
    _write_workbook(path, with_psychology=True)

    entries = read_base_schedule(path)
    sessions = materialize_sessions_for_date(entries, MONDAY)

    psychology = [session for session in sessions if session.treatment == "Ψυχ"]
    assert len(psychology) == 1
    assert is_patient_available(
        "P1",
        MONDAY,
        time(10, 0),
        sessions,
    ) is False
    assert is_patient_available(
        "P1",
        MONDAY,
        time(10, 45),
        sessions,
    ) is True
