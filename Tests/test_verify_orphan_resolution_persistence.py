from pathlib import Path

import pytest
from openpyxl import Workbook

from rehab_excel.daily_statistics_writer import SHEET_NAME
from tools.verify_orphan_resolution_persistence import read_statistics_values


def test_read_statistics_values_reads_key_value_rows(tmp_path: Path):
    path = tmp_path / "stats.xlsm"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws["A4"] = "treatment_suspended"
    ws["C4"] = 2
    ws["A5"] = "unresolved"
    ws["C5"] = 0
    ws["A6"] = "finalized"
    ws["C6"] = "ΝΑΙ"
    wb.save(path)

    values = read_statistics_values(path)

    assert values == {
        "treatment_suspended": 2,
        "unresolved": 0,
        "finalized": "ΝΑΙ",
    }


def test_read_statistics_values_requires_statistics_sheet(tmp_path: Path):
    path = tmp_path / "missing.xlsm"
    wb = Workbook()
    wb.save(path)

    with pytest.raises(RuntimeError, match="DAILY_STATISTICS"):
        read_statistics_values(path)
