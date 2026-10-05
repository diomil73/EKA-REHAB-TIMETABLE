from datetime import time
from pathlib import Path

from openpyxl import Workbook

from rehab_excel.reader import read_base_schedule


def _write_workbook(path: Path, *, with_provider_columns: bool) -> None:
    wb = Workbook()

    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    patients.append(["P1", "101", "ΔΟΚΙΜΗ ΑΣΘΕΝΗ", "ΟΧΙ", "ΠΑΡΩΝ"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    headers = [
        "PatientID",
        "Θάλαμος",
        "Ασθενής",
        "Μολυσματικός",
        "Κατάσταση",
        "Εργο_Ώρα",
        "Εργο_Ημέρες",
        "Λογο_Ώρα",
        "Λογο_Ημέρες",
        "Ψυχ_Ώρα",
        "Ψυχ_Ημέρες",
        "Ψυχ_Ψυχολόγος",
    ]
    if with_provider_columns:
        headers.extend(["Εργο_Θεραπευτής", "Λογο_Θεραπευτής"])
    planner.append(headers)

    row = [
        "P1",
        "101",
        "ΔΟΚΙΜΗ ΑΣΘΕΝΗ",
        "ΟΧΙ",
        "ΠΑΡΩΝ",
        time(9, 15),
        "Δε-Τε-Πα",
        time(10, 0),
        "Τρ-Πε",
        time(10, 45),
        "Καθ/να",
        "Ψυχολόγος Α",
    ]
    if with_provider_columns:
        row.extend(["Εργοθεραπευτής Α", "Λογοθεραπευτής Α"])
    planner.append(row)

    settings = wb.create_sheet("SETTINGS")
    settings.append([
        "THERAPISTS_FTH",
        "HOURS_STANDARD",
        "HOURS_RECLINED",
        "DAY_COMBINATIONS",
        "THERAPIES",
        "PATIENT_STATUS",
        "YES_NO",
        "ROOMS",
        "PSYCHOLOGISTS",
    ])
    settings.append([
        "Θεραπευτής Α",
        time(9, 15),
        time(9, 0),
        "Καθ/να",
        "ΦΘ",
        "ΠΑΡΩΝ",
        "ΝΑΙ",
        "101",
        "Ψυχολόγος Α",
    ])

    wb.save(path)


def _by_treatment(path: Path):
    return {entry.treatment: entry for entry in read_base_schedule(path)}


def test_optional_ot_and_speech_provider_columns_are_imported(tmp_path):
    path = tmp_path / "specialty_providers.xlsx"
    _write_workbook(path, with_provider_columns=True)

    entries = _by_treatment(path)

    assert entries["Εργο"].therapist_id == "Εργοθεραπευτής Α"
    assert entries["Λογο"].therapist_id == "Λογοθεραπευτής Α"
    assert entries["Ψυχ"].therapist_id == "Ψυχολόγος Α"


def test_old_ot_and_speech_schema_remains_backward_compatible(tmp_path):
    path = tmp_path / "legacy_specialties.xlsx"
    _write_workbook(path, with_provider_columns=False)

    entries = _by_treatment(path)

    assert entries["Εργο"].therapist_id is None
    assert entries["Λογο"].therapist_id is None
    assert entries["Ψυχ"].therapist_id == "Ψυχολόγος Α"
