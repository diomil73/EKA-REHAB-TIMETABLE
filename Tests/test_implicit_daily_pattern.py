from __future__ import annotations

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook

from rehab_core.base_schedule import materialize_sessions_for_date
from rehab_excel.reader import IMPLICIT_DAILY_PATTERN, audit_workbook, read_base_schedule


def _fixture(tmp_path: Path) -> Path:
    path = tmp_path / "fixture.xlsx"
    wb = Workbook()
    patients = wb.active
    patients.title = "PATIENTS"
    patients.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    patients.append([1, "A1", "ΑΣΘΕΝΗΣ Α", "Ο", "ΠΑΡΩΝ"])

    planner = wb.create_sheet("PATIENT_PLANNER")
    planner.append(["PatientID", "Ασθενής", "ΦΘ_Ώρα", "ΦΘ_Ημέρες", "ΦΘ_Θεραπευτής"])
    planner.append([1, "ΑΣΘΕΝΗΣ Α", time(8, 30), None, "ΘΕΡΑΠΕΥΤΗΣ Α"])

    settings = wb.create_sheet("SETTINGS")
    settings.append(["Therapists", "Times", "Reclined", "Days", "Therapies", "Statuses", "YesNo", "Rooms"])
    settings.append(["ΘΕΡΑΠΕΥΤΗΣ Α", time(8, 30), None, "Καθ/να", "ΦΘ", "ΠΑΡΩΝ", "Ν", "A1"])
    wb.save(path)
    return path


def test_blank_day_pattern_is_imported_as_daily(tmp_path: Path):
    path = _fixture(tmp_path)
    entries = read_base_schedule(path)
    assert len(entries) == 1
    assert entries[0].day_pattern == IMPLICIT_DAILY_PATTERN == "Καθ/να"


def test_implicit_daily_materializes_weekday_but_not_weekend(tmp_path: Path):
    path = _fixture(tmp_path)
    entries = read_base_schedule(path)
    assert len(materialize_sessions_for_date(entries, date(2026, 9, 18))) == 1  # Fri
    assert len(materialize_sessions_for_date(entries, date(2026, 9, 19))) == 0  # Sat


def test_audit_reports_explicit_daily_assumption_instead_of_missing_day(tmp_path: Path):
    path = _fixture(tmp_path)
    audit = audit_workbook(path)
    codes = {issue.code: issue for issue in audit.issues}
    assert "timed_entries_missing_day_pattern" not in codes
    assert codes["timed_entries_assumed_daily"].count == 1
    assert audit.base_entry_count == 1
