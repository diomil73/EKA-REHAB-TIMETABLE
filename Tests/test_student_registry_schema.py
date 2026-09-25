from datetime import date
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "Python"))

from rehab_excel.source_contract import imported_sheets, is_authoritative
from rehab_excel.student_registry import (
    STUDENT_REGISTRY_HEADERS,
    StudentRegistryError,
    read_students,
)


def test_old_workbook_without_students_sheet_remains_compatible(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.active.title = "PATIENTS"
    wb.save(path)

    assert read_students(path) == []


def test_students_registry_imports_identity_placement_and_capabilities(tmp_path):
    path = tmp_path / "students.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "STUDENTS"
    ws.append(STUDENT_REGISTRY_HEADERS)
    ws.append(
        [
            "STU-1",
            "Μαρία Παπαδοπούλου",
            1,
            date(2026, 9, 1),
            date(2026, 12, 20),
            "Πέτσιος",
            "ΝΑΙ",
            "ΟΧΙ",
        ]
    )
    wb.save(path)

    students = read_students(path)

    assert len(students) == 1
    student = students[0]
    assert student.student_id == "STU-1"
    assert student.display_name == "Μαρία Παπαδοπούλου"
    assert student.student_number == 1
    assert student.placement_start == date(2026, 9, 1)
    assert student.placement_end == date(2026, 12, 20)
    assert student.supervisor_therapist_id == "Πέτσιος"
    assert student.replacement_capable is True
    assert student.robotic_capable is False
    assert student.max_daily_timeslots == 5


def test_blank_capabilities_use_confirmed_student_defaults(tmp_path):
    path = tmp_path / "defaults.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "STUDENTS"
    ws.append(STUDENT_REGISTRY_HEADERS)
    ws.append(
        [
            "STU-2",
            "Νίκος Δοκιμή",
            2,
            "01/10/2026",
            "31/12/2026",
            None,
            None,
            None,
        ]
    )
    wb.save(path)

    student = read_students(path)[0]

    assert student.replacement_capable is True
    assert student.robotic_capable is False


def test_duplicate_student_identity_is_rejected(tmp_path):
    path = tmp_path / "duplicate.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "STUDENTS"
    ws.append(STUDENT_REGISTRY_HEADERS)
    ws.append(["STU-1", "Α", 1, "01/09/2026", "30/09/2026", None, None, None])
    ws.append(["stu-1", "Β", 2, "01/09/2026", "30/09/2026", None, None, None])
    wb.save(path)

    with pytest.raises(StudentRegistryError, match="duplicate StudentID"):
        read_students(path)


def test_malformed_students_headers_are_rejected(tmp_path):
    path = tmp_path / "bad_headers.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "STUDENTS"
    ws.append(["Wrong"] + list(STUDENT_REGISTRY_HEADERS[1:]))
    wb.save(path)

    with pytest.raises(StudentRegistryError, match="headers"):
        read_students(path)


def test_students_is_authoritative_imported_source():
    assert is_authoritative("STUDENTS") is True
    assert "STUDENTS" in imported_sheets()
