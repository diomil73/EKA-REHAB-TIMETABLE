from datetime import date, time

from openpyxl import Workbook, load_workbook

from rehab_core.models import Patient, Session
from rehab_excel.daily_input_reader import read_daily_input


TODAY = date(2026, 10, 30)


def _book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_INPUT"
    ws["B2"] = "30/10/2026"
    for c, value in enumerate(("Θεραπευτής", "Όλη ημέρα", "Από", "Έως", "Αιτία", "Σχόλιο"), 1):
        ws.cell(6, c, value)
    for c, value in enumerate(("Ασθενής", "Όλη ημέρα", "Ώρα", "Κατάσταση", "Σχόλιο"), 1):
        ws.cell(20, c, value)
    wb.save(path)
    wb.close()


def _fixtures():
    patients = [Patient("P1", "ΑΣΘΕΝΗΣ 1")]
    sessions = [Session("S1", "P1", "T1", TODAY, time(9, 0))]
    return patients, sessions


def test_reader_emits_source_change_for_patient_absence(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΑΣΘΕΝΗΣ 1"
    ws["B21"] = "ΟΧΙ"
    ws["C21"] = "09:00"
    ws["D21"] = "ΑΠΩΝ"
    wb.save(path)
    wb.close()

    patients, sessions = _fixtures()
    result = read_daily_input(path, patients=patients, sessions=sessions)

    assert len(result.source_changes) == 1
    change = result.source_changes[0]
    assert change.source_row == 21
    assert change.change_type == "patient_absence"
    assert change.subject_id == "P1"
    assert change.start_time == time(9, 0)


def test_reader_can_ignore_already_committed_source_row(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΑΣΘΕΝΗΣ 1"
    ws["B21"] = "ΟΧΙ"
    ws["C21"] = "09:00"
    ws["D21"] = "ΑΠΩΝ"
    wb.save(path)
    wb.close()

    patients, sessions = _fixtures()
    result = read_daily_input(
        path,
        patients=patients,
        sessions=sessions,
        ignore_rows=(21,),
        allow_empty=True,
    )

    assert result.absences == ()
    assert result.cancellations == ()
    assert result.source_changes == ()
    assert result.patient_rows_used == 0


def test_reader_allow_empty_returns_empty_result(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    patients, sessions = _fixtures()

    result = read_daily_input(
        path,
        patients=patients,
        sessions=sessions,
        allow_empty=True,
    )

    assert result.target_date == TODAY
    assert result.absences == ()
    assert result.cancellations == ()
    assert result.source_changes == ()
