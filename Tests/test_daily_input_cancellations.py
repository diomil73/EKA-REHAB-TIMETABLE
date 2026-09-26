from __future__ import annotations

from datetime import date, time

import pytest
from openpyxl import Workbook, load_workbook

from rehab_core.models import (
    AbsenceKind,
    DailyAbsence,
    DailySessionCancellation,
    Patient,
    Session,
    SessionCancellationKind,
    Therapist,
)
from rehab_core.therapist_absence_queue import build_therapist_absence_replacement_queue
from rehab_excel.daily_input_reader import DailyInputReadError, read_daily_input


DAY = date(2026, 9, 18)


def _book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_INPUT"
    ws["A2"] = "Ημερομηνία"
    ws["B2"] = "18/09/2026"
    for column, value in enumerate(
        ("Θεραπευτής", "Όλη ημέρα", "Από", "Έως", "Αιτία", "Σχόλιο"),
        1,
    ):
        ws.cell(6, column, value)
    for column, value in enumerate(
        ("Ασθενής", "Όλη ημέρα", "Ώρα", "Κατάσταση", "Σχόλιο"),
        1,
    ):
        ws.cell(20, column, value)
    wb.save(path)
    wb.close()


def _patient() -> Patient:
    return Patient("P1", "ΕΞΩΤΕΡΙΚΟΣ")


def _session(session_id: str = "S1", at: time = time(8, 30)) -> Session:
    return Session(
        session_id=session_id,
        patient_id="P1",
        therapist_id="T1",
        session_date=DAY,
        start_time=at,
        treatment="ΦΘ",
    )


def _write_patient_row(path, *, status: str, at="08:30", all_day="ΟΧΙ", comment=""):
    wb = load_workbook(path)
    ws = wb["DAILY_INPUT"]
    ws["A21"] = "ΕΞΩΤΕΡΙΚΟΣ"
    ws["B21"] = all_day
    ws["C21"] = at
    ws["D21"] = status
    ws["E21"] = comment
    wb.save(path)
    wb.close()


def test_reads_department_postponement_as_session_cancellation(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    _write_patient_row(
        path,
        status="ΑΝΑΒΟΛΗ ΤΜΗΜΑΤΟΣ",
        comment="μη διαθέσιμος χώρος",
    )

    result = read_daily_input(path, patients=[_patient()], sessions=[_session()])

    assert result.absences == ()
    assert len(result.cancellations) == 1
    cancellation = result.cancellations[0]
    assert cancellation.target_session_id == "S1"
    assert cancellation.cancellation_date == DAY
    assert cancellation.kind == SessionCancellationKind.DEPARTMENT_POSTPONED
    assert "μη διαθέσιμος χώρος" in (cancellation.reason or "")


def test_reads_patient_no_show_as_session_cancellation(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    _write_patient_row(path, status="ΔΕΝ ΠΡΟΣΗΛΘΕ")

    result = read_daily_input(path, patients=[_patient()], sessions=[_session()])

    assert result.absences == ()
    assert result.cancellations[0].kind == SessionCancellationKind.PATIENT_NO_SHOW


def test_session_cancellation_rejects_all_day_flag(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    _write_patient_row(path, status="ΔΕΝ ΠΡΟΣΗΛΘΕ", all_day="ΝΑΙ")

    with pytest.raises(DailyInputReadError, match="must target one session"):
        read_daily_input(path, patients=[_patient()], sessions=[_session()])


def test_session_cancellation_requires_exactly_one_matching_session(tmp_path):
    path = tmp_path / "input.xlsm"
    _book(path)
    _write_patient_row(path, status="ΑΝΑΒΟΛΗ ΤΜΗΜΑΤΟΣ")

    sessions = [_session("S1"), _session("S2")]
    with pytest.raises(DailyInputReadError, match="Expected exactly one session"):
        read_daily_input(path, patients=[_patient()], sessions=sessions)


def test_cancelled_session_is_not_sent_to_therapist_replacement_queue():
    session = _session()
    cancellation = DailySessionCancellation(
        cancellation_id="C1",
        target_session_id=session.session_id,
        cancellation_date=DAY,
        kind=SessionCancellationKind.PATIENT_NO_SHOW,
        reason="no show",
    )
    queue = build_therapist_absence_replacement_queue(
        sessions=[session],
        absences=[DailyAbsence(AbsenceKind.THERAPIST, "T1", DAY)],
        cancellations=[cancellation],
        therapists=[Therapist("T1", "T1"), Therapist("T2", "T2")],
        patients=[_patient()],
        timeslots=(time(8, 30), time(9, 15)),
    )

    assert queue.items == ()
    assert queue.skipped_cancelled == 1
