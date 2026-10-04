from pathlib import Path

from openpyxl import Workbook, load_workbook

from rehab_core.models import Patient
from rehab_excel.master_projection import MasterProjectionRow
from rehab_excel.master_projection_writer import (
    MasterProjectionWriteError,
    create_master_projection_preview,
)


class _FakeBackend:
    def apply(self, workbook_path, projection):
        wb = load_workbook(workbook_path, keep_vba=False)
        try:
            ws = wb["MASTER_SCHEDULE"]
            for item in projection:
                row = item.target_row
                ws.cell(row, 1).value = f"=PATIENT_PLANNER!D{item.planner_row}"
                ws.cell(row, 2).value = f"=PATIENT_PLANNER!B{item.planner_row}"
                ws.cell(row, 3).value = f"=PATIENT_PLANNER!C{item.planner_row}"
                ws.cell(row, 4).value = f"=PATIENT_PLANNER!F{item.planner_row}"
                ws.cell(row, 5).value = f"=PATIENT_PLANNER!I{item.planner_row}"
                ws.cell(row, 6).value = f"=PATIENT_PLANNER!N{item.planner_row}"
                ws.cell(row, 7).value = f"=PATIENT_PLANNER!L{item.planner_row}"
                ws.cell(row, 8).value = f"=PATIENT_PLANNER!P{item.planner_row}"
                ws.cell(row, 9).value = f"=PATIENT_PLANNER!R{item.planner_row}"
                ws.cell(row, 10).value = f"=PATIENT_PLANNER!T{item.planner_row}"
                ws.cell(row, 11).value = f"=PATIENT_PLANNER!E{item.planner_row}"
            tail_start = max((item.target_row for item in projection), default=1) + 1
            for row in range(tail_start, ws.max_row + 1):
                for col in range(1, 12):
                    ws.cell(row, col).value = None
            wb.save(workbook_path)
        finally:
            wb.close()


def _book(tmp_path: Path) -> Path:
    path = tmp_path / "source.xlsm"
    wb = Workbook()
    patients = wb.active
    patients.title = "PATIENTS"
    patients.append([
        "PatientID",
        "Θάλαμος",
        "Ασθενής",
        "Μολυσματικός",
        "Κατάσταση",
        "PatientType",
    ])
    patients.append([1, "A09", "ΧΑΣΙΚΟΣ", "Ο", "παρών", "Εσωτερικός"])
    patients.append([2, "A02", "ΒΑΡΒΑΡΑΣ", "Ν", "παρών", "Εσωτερικός"])
    patients.append([3, "A09", "ΤΕΣΤ", "Ν", "παρών", "Εσωτερικός"])

    settings = wb.create_sheet("SETTINGS")
    settings["H1"] = "ROOMS"
    settings["H2"] = "A01"
    settings["H3"] = "A02"
    settings["H4"] = "A09"

    planner = wb.create_sheet("PATIENT_PLANNER")
    planner.append([
        "PatientID",
        "Θάλαμος",
        "Ασθενής",
        "Μολυσματικός",
        "Κατάσταση",
        "ΦΘ_Ώρα",
        "ΦΘ_Ημέρες",
        "ΦΘ_Θεραπευτής",
    ])
    for row in range(2, 5):
        planner.cell(row, 1).value = f"=PATIENTS!A{row}"
        planner.cell(row, 2).value = f"=PATIENTS!B{row}"
        planner.cell(row, 3).value = f"=PATIENTS!C{row}"
        planner.cell(row, 4).value = f"=PATIENTS!D{row}"
        planner.cell(row, 5).value = f"=PATIENTS!E{row}"

    master = wb.create_sheet("MASTER_SCHEDULE")
    master.append([
        "Μολυσματικός",
        "Θάλαμος",
        "Ασθενής",
        "ΦΘ",
        "Ρομποτικό",
        "Πισίνα",
        "Ανακλινόμενο",
        "Εργο",
        "Λογο",
        "ΕΦΑ",
        "Κατάσταση",
    ])
    for row in range(2, 6):
        master.cell(row, 1).value = f"=PATIENT_PLANNER!D{row}"
        master.cell(row, 2).value = f"=PATIENT_PLANNER!B{row}"
        master.cell(row, 3).value = f"=PATIENT_PLANNER!C{row}"
        master.cell(row, 4).value = f"=PATIENT_PLANNER!F{row}"
        master.cell(row, 5).value = f"=PATIENT_PLANNER!I{row}"
        master.cell(row, 6).value = f"=PATIENT_PLANNER!N{row}"
        master.cell(row, 7).value = f"=PATIENT_PLANNER!L{row}"
        master.cell(row, 8).value = f"=PATIENT_PLANNER!P{row}"
        master.cell(row, 9).value = f"=PATIENT_PLANNER!R{row}"
        master.cell(row, 10).value = f"=PATIENT_PLANNER!T{row}"
        master.cell(row, 11).value = f"=PATIENT_PLANNER!E{row}"

    wb.save(path)
    wb.close()
    return path


def test_create_master_projection_preview_preserves_source_and_verifies(tmp_path, monkeypatch):
    source = _book(tmp_path)
    output = tmp_path / "preview.xlsm"

    monkeypatch.setattr(
        "rehab_excel.master_projection_writer.build_master_projection",
        lambda _path: (
            MasterProjectionRow(2, 3, "2", "A02", "ΒΑΡΒΑΡΑΣ", True, "παρών"),
            MasterProjectionRow(3, 4, "3", "A09", "ΤΕΣΤ", True, "παρών"),
            MasterProjectionRow(4, 2, "1", "A09", "ΧΑΣΙΚΟΣ", False, "παρών"),
        ),
    )

    before = source.read_bytes()
    report = create_master_projection_preview(
        source,
        output,
        backend=_FakeBackend(),
        overwrite=True,
    )

    assert report.projected_rows == 3
    assert report.source_unchanged is True
    assert report.verified_in_output is True
    assert source.read_bytes() == before
    assert output.exists()


def test_create_master_projection_preview_rejects_same_source_and_output(tmp_path):
    source = _book(tmp_path)

    try:
        create_master_projection_preview(source, source)
    except MasterProjectionWriteError as exc:
        assert "differ from source" in str(exc)
    else:
        raise AssertionError("Expected same source/output to be rejected")



def test_master_writer_verification_supports_shifted_header(tmp_path, monkeypatch):
    source = _book(tmp_path)
    wb = load_workbook(source)
    try:
        ws = wb["MASTER_SCHEDULE"]
        ws.insert_rows(1, amount=3)
        wb.save(source)
    finally:
        wb.close()

    output = tmp_path / "shifted_preview.xlsm"
    monkeypatch.setattr(
        "rehab_excel.master_projection_writer.build_master_projection",
        lambda _path: (
            MasterProjectionRow(5, 3, "2", "A02", "ΒΑΡΒΑΡΑΣ", True, "παρών"),
            MasterProjectionRow(6, 2, "1", "A09", "ΧΑΣΙΚΟΣ", False, "παρών"),
        ),
    )

    report = create_master_projection_preview(
        source,
        output,
        backend=_FakeBackend(),
        overwrite=True,
    )

    assert report.verified_in_output is True
