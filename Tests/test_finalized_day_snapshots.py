from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from rehab_core.day_finalization import FinalizedDaySnapshot, FinalizedSnapshotStatus
from rehab_excel.finalized_day_snapshots import (
    HEADERS,
    SHEET_NAME,
    FinalizedDaySnapshotStoreError,
    read_finalized_day_snapshots,
)


DAY = date(2026, 10, 26)
FINALIZED_AT = datetime(2026, 10, 26, 15, 30, tzinfo=timezone.utc)


def _snapshot(
    snapshot_id: str = "snap-2026-10-26-v1",
    *,
    status: FinalizedSnapshotStatus = FinalizedSnapshotStatus.CURRENT_FINAL,
    version: int = 1,
) -> FinalizedDaySnapshot:
    return FinalizedDaySnapshot(
        snapshot_id=snapshot_id,
        service_date=DAY,
        version=version,
        status=status,
        finalized_at=FINALIZED_AT,
        supersedes_snapshot_id=None,
        unresolved_sessions=0,
        planned_session_count=5,
        final_session_count=3,
        patient_absence_count=1,
        cancellation_count=1,
        no_show_count=0,
        replacement_count=1,
        treatment_suspension_count=1,
        approved_exception_count=0,
    )


def _append_row(ws, snapshot: FinalizedDaySnapshot) -> None:
    ws.append(
        [
            snapshot.snapshot_id,
            snapshot.service_date.isoformat(),
            snapshot.version,
            snapshot.status.value,
            snapshot.finalized_at.isoformat(),
            snapshot.supersedes_snapshot_id or "",
            snapshot.unresolved_sessions,
            snapshot.planned_session_count,
            snapshot.final_session_count,
            snapshot.patient_absence_count,
            snapshot.cancellation_count,
            snapshot.no_show_count,
            snapshot.replacement_count,
            snapshot.treatment_suspension_count,
            snapshot.approved_exception_count,
        ]
    )


def _workbook(tmp_path: Path, *snapshots: FinalizedDaySnapshot) -> Path:
    path = tmp_path / "history.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for snapshot in snapshots:
        _append_row(ws, snapshot)
    visible = wb.create_sheet("VISIBLE")
    wb.active = wb.sheetnames.index(visible.title)
    ws.sheet_state = "veryHidden"
    ws.protection.sheet = True
    wb.save(path)
    wb.close()
    return path


def test_reader_round_trips_finalized_snapshot_header(tmp_path: Path):
    expected = _snapshot()
    path = _workbook(tmp_path, expected)

    actual = read_finalized_day_snapshots(path)

    assert actual == (expected,)


def test_reader_filters_current_snapshot_for_service_date(tmp_path: Path):
    superseded = FinalizedDaySnapshot(
        **{
            **_snapshot("snap-old", status=FinalizedSnapshotStatus.SUPERSEDED).__dict__,
            "supersedes_snapshot_id": None,
        }
    )
    current = FinalizedDaySnapshot(
        **{
            **_snapshot("snap-current", version=2).__dict__,
            "supersedes_snapshot_id": "snap-old",
        }
    )
    path = _workbook(tmp_path, superseded, current)

    actual = read_finalized_day_snapshots(
        path,
        service_date=DAY,
        current_only=True,
    )

    assert actual == (current,)


def test_reader_rejects_two_current_final_snapshots_for_same_date(tmp_path: Path):
    path = _workbook(tmp_path, _snapshot("snap-a"), _snapshot("snap-b", version=2))

    with pytest.raises(FinalizedDaySnapshotStoreError, match="Multiple CURRENT_FINAL"):
        read_finalized_day_snapshots(path)


def test_reader_rejects_nonzero_unresolved_in_finalized_history(tmp_path: Path):
    invalid = FinalizedDaySnapshot(
        **{
            **_snapshot().__dict__,
            "unresolved_sessions": 1,
        }
    )
    path = _workbook(tmp_path, invalid)

    with pytest.raises(FinalizedDaySnapshotStoreError, match="unresolved"):
        read_finalized_day_snapshots(path)


def test_reader_rejects_schema_drift(tmp_path: Path):
    path = tmp_path / "bad-schema.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(("WrongHeader",) + HEADERS[1:])
    wb.save(path)
    wb.close()

    with pytest.raises(FinalizedDaySnapshotStoreError, match="headers"):
        read_finalized_day_snapshots(path)


def test_native_writer_does_not_depend_on_com_number_format_assignment():
    source = Path("Python/rehab_excel/finalized_day_snapshots.py").read_text(encoding="utf-8")

    assert ".NumberFormat" not in source
    assert "ws.Cells(next_row, col).Value2 = value" in source
    assert "Finalized snapshot COM read-back did not match requested snapshot" in source
