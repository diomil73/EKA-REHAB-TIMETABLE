from datetime import date, time

import pytest
from openpyxl import Workbook

from rehab_core.models import ReplacementAssignment, ReplacementProviderKind, Session
from rehab_excel.daily_replacement_facts import (
    ACTIVE_STATUS,
    HEADERS,
    SHEET_NAME,
    DailyReplacementFactError,
    build_daily_replacement_fact_rows,
    read_daily_replacement_facts,
)


DAY = date(2026, 10, 19)


def _session(
    session_id: str = "S1",
    *,
    patient_id: str = "P1",
    therapist_id: str = "Σάρρας",
    at: time = time(9, 15),
) -> Session:
    return Session(
        session_id=session_id,
        patient_id=patient_id,
        therapist_id=therapist_id,
        session_date=DAY,
        start_time=at,
        treatment="ΦΘ",
    )


def _replacement(
    *,
    replacement_id: str = "R1",
    target_session_id: str = "S1",
    patient_id: str = "P1",
    original_therapist_id: str = "Σάρρας",
    replacement_therapist_id: str = "Πέτσιος",
    replacement_time: time = time(10, 0),
) -> ReplacementAssignment:
    return ReplacementAssignment(
        replacement_id=replacement_id,
        target_session_id=target_session_id,
        patient_id=patient_id,
        original_therapist_id=original_therapist_id,
        replacement_therapist_id=replacement_therapist_id,
        replacement_date=DAY,
        replacement_time=replacement_time,
        reason="DAILY_INPUT therapist absence",
        replacement_provider_kind=ReplacementProviderKind.THERAPIST,
    )


def _write_fact_sheet(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(1, col).value = header
    for row_number, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row_number, col).value = value
    wb.save(path)
    wb.close()


def test_build_rows_preserves_exact_replacement_semantics():
    session = _session()
    replacement = _replacement()

    rows = build_daily_replacement_fact_rows([replacement], sessions=[session])

    assert len(rows) == 1
    row = rows[0]
    assert row.replacement_id == "R1"
    assert row.target_session_id == "S1"
    assert row.patient_id == "P1"
    assert row.original_therapist_id == "Σάρρας"
    assert row.replacement_therapist_id == "Πέτσιος"
    assert row.replacement_date == DAY
    assert row.replacement_time == time(10, 0)
    assert row.provider_kind == ReplacementProviderKind.THERAPIST
    assert row.status == ACTIVE_STATUS


def test_build_rows_rejects_patient_mismatch():
    with pytest.raises(DailyReplacementFactError, match="does not match session patient"):
        build_daily_replacement_fact_rows(
            [_replacement(patient_id="OTHER")],
            sessions=[_session()],
        )


def test_build_rows_rejects_original_therapist_mismatch():
    with pytest.raises(DailyReplacementFactError, match="does not match session therapist"):
        build_daily_replacement_fact_rows(
            [_replacement(original_therapist_id="OTHER")],
            sessions=[_session()],
        )


def test_build_rows_rejects_two_active_facts_for_same_target_date():
    with pytest.raises(DailyReplacementFactError, match="Multiple active replacement facts"):
        build_daily_replacement_fact_rows(
            [
                _replacement(replacement_id="R1"),
                _replacement(
                    replacement_id="R2",
                    replacement_therapist_id="Φιλιππούσης",
                    replacement_time=time(10, 45),
                ),
            ],
            sessions=[_session()],
        )


def test_reader_missing_sheet_is_backward_compatible(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()

    assert read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY) == ()


def test_reader_reconstructs_replacement_assignment(tmp_path):
    path = tmp_path / "facts.xlsx"
    _write_fact_sheet(
        path,
        [
            (
                "R1",
                "S1",
                "P1",
                "Σάρρας",
                "Πέτσιος",
                DAY,
                "10:00",
                "DAILY_INPUT therapist absence",
                "therapist",
                "ACTIVE",
            )
        ],
    )

    facts = read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY)

    assert facts == (_replacement(),)


def test_reader_defaults_blank_provider_kind_to_therapist(tmp_path):
    path = tmp_path / "facts.xlsx"
    _write_fact_sheet(
        path,
        [
            (
                "R1",
                "S1",
                "P1",
                "Σάρρας",
                "Πέτσιος",
                "2026-10-19",
                "10:00",
                "DAILY_INPUT therapist absence",
                "",
                "ACTIVE",
            )
        ],
    )

    facts = read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY)

    assert facts[0].replacement_provider_kind == ReplacementProviderKind.THERAPIST


def test_reader_filters_other_dates_before_session_validation(tmp_path):
    path = tmp_path / "facts.xlsx"
    _write_fact_sheet(
        path,
        [
            (
                "OTHER-DAY",
                "UNKNOWN",
                "P-X",
                "T-X",
                "T-Y",
                "2026-10-20",
                "10:00",
                "other day",
                "therapist",
                "ACTIVE",
            ),
            (
                "R1",
                "S1",
                "P1",
                "Σάρρας",
                "Πέτσιος",
                "2026-10-19",
                "10:00",
                "DAILY_INPUT therapist absence",
                "therapist",
                "ACTIVE",
            ),
        ],
    )

    facts = read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY)

    assert [fact.replacement_id for fact in facts] == ["R1"]


def test_reader_rejects_conflicting_duplicate_target_date(tmp_path):
    path = tmp_path / "facts.xlsx"
    _write_fact_sheet(
        path,
        [
            (
                "R1",
                "S1",
                "P1",
                "Σάρρας",
                "Πέτσιος",
                DAY,
                "10:00",
                "one",
                "therapist",
                "ACTIVE",
            ),
            (
                "R2",
                "S1",
                "P1",
                "Σάρρας",
                "Φιλιππούσης",
                DAY,
                "10:45",
                "two",
                "therapist",
                "ACTIVE",
            ),
        ],
    )

    with pytest.raises(DailyReplacementFactError, match="Multiple active replacement facts"):
        read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY)


def test_reader_rejects_header_drift(tmp_path):
    path = tmp_path / "facts.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws["A1"] = "wrong"
    wb.save(path)
    wb.close()

    with pytest.raises(DailyReplacementFactError, match="headers do not match"):
        read_daily_replacement_facts(path, sessions=[_session()], target_date=DAY)
