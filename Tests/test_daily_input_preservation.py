from openpyxl import Workbook

from rehab_excel.daily_input_preservation import (
    DailyInputPreservationError,
    assert_daily_input_preserved,
    snapshot_daily_input,
)


def _book(path, therapist="ΘΕΡΑΠΕΥΤΗΣ Α", patient="ΑΣΘΕΝΗΣ Α"):
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_INPUT"
    ws["A7"] = therapist
    ws["B7"] = "ΝΑΙ"
    ws["A21"] = patient
    ws["D21"] = "ΑΠΩΝ"
    wb.save(path)
    wb.close()


def test_daily_input_snapshot_roundtrip_preserves_therapist_and_patient_rows(tmp_path):
    source = tmp_path / "source.xlsm"
    output = tmp_path / "output.xlsm"
    _book(source)
    _book(output)

    snapshot = snapshot_daily_input(source)
    assert_daily_input_preserved(snapshot, output)


def test_daily_input_preservation_detects_lost_therapist_row(tmp_path):
    source = tmp_path / "source.xlsm"
    output = tmp_path / "output.xlsm"
    _book(source)
    _book(output, therapist="")

    snapshot = snapshot_daily_input(source)

    try:
        assert_daily_input_preserved(snapshot, output)
    except DailyInputPreservationError as exc:
        assert "A7" in str(exc)
    else:
        raise AssertionError("Expected lost therapist row to be detected")
