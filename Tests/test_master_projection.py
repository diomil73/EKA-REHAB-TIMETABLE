from pathlib import Path

from openpyxl import Workbook

from rehab_core.models import Patient, PatientType
from rehab_excel.master_projection import (
    build_master_projection,
    remap_master_formula,
)


def _workbook(tmp_path: Path) -> Path:
    xlsx = tmp_path / "book.xlsx"
    path = tmp_path / "book.xlsm"
    wb = Workbook()
    ws = wb.active
    ws.title = "PATIENTS"
    ws.append(["PatientID", "Θάλαμος", "Ασθενής", "Μολυσματικός", "Κατάσταση"])
    ws.append([1, "A09", "ΧΑΣΙΚΟΣ", "Ο", "παρών"])
    ws.append([2, "A02", "ΒΑΡΒΑΡΑΣ", "Ν", "παρών"])
    ws.append([3, "A09", "ΤΕΣΤ", "Ν", "παρών"])
    ws.append([4, "", "OUT", "", ""])
    wb.create_sheet("SETTINGS")
    wb.save(xlsx)
    path.write_bytes(xlsx.read_bytes())
    return path


def test_projection_sorts_rooms_naturally_then_patient_name(tmp_path):
    path = _workbook(tmp_path)
    patients = [
        Patient("1", "ΧΑΣΙΚΟΣ", room="A09"),
        Patient("2", "ΒΑΡΒΑΡΑΣ", room="A02", infectious=True),
        Patient("3", "ΤΕΣΤ", room="A09", infectious=True),
        Patient(
            "4",
            "OUT",
            patient_type=PatientType.OUTPATIENT,
        ),
    ]

    rows = build_master_projection(
        path,
        patients=patients,
        room_order=("A09", "A02", "B01", "A01"),
    )

    # SETTINGS order is intentionally stale; clinical natural order wins.
    # Row 2 is deliberately reserved for the Α' ΚΛΙΝΙΚΗ banner.
    assert [(row.target_row, row.planner_row, row.patient_id) for row in rows] == [
        (3, 3, "2"),
        (4, 4, "3"),
        (5, 2, "1"),
    ]
    assert all(row.patient_id != "4" for row in rows)


def test_projection_never_interleaves_b_clinic_inside_a_clinic(tmp_path):
    path = _workbook(tmp_path)
    patients = [
        Patient("1", "A11 PATIENT", room="A11"),
        Patient("2", "B08 PATIENT", room="B08"),
        Patient("3", "A05 PATIENT", room="A05"),
    ]

    rows = build_master_projection(
        path,
        patients=patients,
        room_order=("A11", "B08", "A05"),
    )

    assert [(row.room, row.target_row) for row in rows] == [
        ("A05", 3),
        ("A11", 4),
        ("B08", 6),
    ]
    # Rows 2 and 5 are the Α' and Β' clinic banner rows.


def test_projection_keeps_patient_registry_rows_as_formula_targets(tmp_path):
    path = _workbook(tmp_path)
    patients = [
        Patient("3", "ΤΕΣΤ", room="A09", infectious=True),
        Patient("1", "ΧΑΣΙΚΟΣ", room="A09"),
    ]

    rows = build_master_projection(
        path,
        patients=patients,
        room_order=("A09",),
    )

    assert rows[0].display_name == "ΤΕΣΤ"
    assert rows[0].target_row == 3
    assert rows[0].planner_row == 4
    assert rows[1].display_name == "ΧΑΣΙΚΟΣ"
    assert rows[1].target_row == 4
    assert rows[1].planner_row == 2


def test_projection_reserves_exactly_one_banner_row_at_clinic_transition(tmp_path):
    path = _workbook(tmp_path)
    patients = [
        Patient("1", "ΧΑΣΙΚΟΣ", room="A09"),
        Patient("2", "ΒΑΡΒΑΡΑΣ", room="A02"),
        Patient("3", "ΤΕΣΤ", room="B01"),
    ]

    rows = build_master_projection(
        path,
        patients=patients,
        room_order=("A02", "A09", "B01"),
    )

    assert [(row.room, row.target_row) for row in rows] == [
        ("A02", 3),
        ("A09", 4),
        ("B01", 6),
    ]
    # Rows 2 and 5 are reserved for Α' and Β' clinic labels respectively.


def test_remap_master_formula_redirects_all_planner_references():
    formula = (
        '=IF(PATIENT_PLANNER!F26="","",'
        'PATIENT_PLANNER!F26&" "&PATIENT_PLANNER!G26)'
    )

    remapped = remap_master_formula(formula, 100)

    assert "PATIENT_PLANNER!F100" in remapped
    assert "PATIENT_PLANNER!G100" in remapped
    assert "PATIENT_PLANNER!F26" not in remapped


def test_remap_master_formula_handles_absolute_row_reference():
    assert (
        remap_master_formula("=PATIENT_PLANNER!$D$2", 100)
        == "=PATIENT_PLANNER!$D$100"
    )


def test_projection_starts_below_shifted_master_header(tmp_path):
    path = _workbook(tmp_path)

    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        ws = wb.create_sheet("MASTER_SCHEDULE")
        ws.append([
            "Μολυσματικός",
            "Θάλαμος",
            "Ασθενής",
            "ΦΘ",
            "Ρομποτικό",
            "Πισίνα",
            "Ανακλινόμενο",
            "Εργο",
            "Λογο",
            "ΕΦΑ",
            "Κατάσταση",
        ])
        ws.insert_rows(1, amount=3)
        wb.save(path)
    finally:
        wb.close()

    patients = [
        Patient("1", "ΧΑΣΙΚΟΣ", room="A09"),
        Patient("2", "ΒΑΡΒΑΡΑΣ", room="A02", infectious=True),
    ]

    rows = build_master_projection(
        path,
        patients=patients,
        room_order=("A02", "A09"),
    )

    # Header is row 4, row 5 is reserved for Α' ΚΛΙΝΙΚΗ.
    assert [row.target_row for row in rows] == [6, 7]
