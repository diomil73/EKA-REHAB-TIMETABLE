from datetime import date, time

import pytest
from openpyxl import Workbook

from rehab_core.models import Session
from tools.apply_daily_input_replacements import _replacement_id
from tools.verify_daily_fact_reconstruction import (
    DailyFactVerificationError,
    _as_int,
    _stored_statistics,
)


def test_replacement_id_is_stable_per_session_and_date():
    session = Session(
        session_id="S-123",
        patient_id="P1",
        therapist_id="T1",
        session_date=date(2026, 10, 19),
        start_time=time(9, 15),
    )

    assert _replacement_id(session, []) == "daily-input:S-123:2026-10-19"


def test_stored_statistics_reads_key_value_rows(tmp_path):
    path = tmp_path / "stats.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY_STATISTICS"
    ws["A4"] = "scheduled"
    ws["C4"] = 80
    ws["A5"] = "replaced"
    ws["C5"] = 5
    ws["A6"] = "finalized"
    ws["C6"] = "ΝΑΙ"
    wb.save(path)
    wb.close()

    values = _stored_statistics(path)

    assert values == {"scheduled": "80", "replaced": "5", "finalized": "ΝΑΙ"}
    assert _as_int(values, "scheduled") == 80
    assert _as_int(values, "replaced") == 5


def test_as_int_rejects_missing_metric():
    with pytest.raises(DailyFactVerificationError, match="has no 'replaced' row"):
        _as_int({"scheduled": "80"}, "replaced")


def test_as_int_rejects_non_numeric_metric():
    with pytest.raises(DailyFactVerificationError, match="is not numeric"):
        _as_int({"replaced": "five"}, "replaced")
