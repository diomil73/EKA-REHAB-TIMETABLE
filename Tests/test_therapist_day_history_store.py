from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from rehab_core.therapist_day_history import TherapistDayHistoryRow
from rehab_excel.therapist_day_history_store import (
    HEADERS,
    SHEET_NAME,
    TherapistDayHistoryStoreError,
    read_therapist_day_history_rows,
)


DAY = date(2026, 10, 26)


def _row(therapist_id: str = "T1", *, snapshot_id: str = "snap-v1") -> TherapistDayHistoryRow:
    return TherapistDayHistoryRow(
        snapshot_id=snapshot_id,
        service_date=DAY,
        therapist_id=therapist_id,
        capacity_timeslots=6,
        planned_occupied_timeslots=5,
        planned_sessions=6,
        final_occupied_timeslots=4,
        final_active_sessions=5,
        replacements_received=1,
        released_by_cancellation_or_no_show=1,
        treatment_suspensions_affecting_load=0,
        approved_shared_or_over_capacity_events=1,
    )


def _workbook(tmp_path: Path, *rows: TherapistDayHistoryRow) -> Path:
    path = tmp_path / "history.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for row in rows:
        ws.append(
            [
                row.snapshot_id,
                row.service_date.isoformat(),
                row.therapist_id,
                row.capacity_timeslots,
                row.planned_occupied_timeslots,
                row.planned_sessions,
                row.final_occupied_timeslots,
                row.final_active_sessions,
                row.replacements_received,
                row.released_by_cancellation_or_no_show,
                row.treatment_suspensions_affecting_load,
                row.approved_shared_or_over_capacity_events,
            ]
        )
    visible = wb.create_sheet("VISIBLE")
    wb.active = wb.sheetnames.index(visible.title)
    ws.sheet_state = "veryHidden"
    ws.protection.sheet = True
    wb.save(path)
    wb.close()
    return path


def test_reader_round_trips_therapist_day_history(tmp_path: Path):
    expected = (_row("T1"), _row("T2"))
    path = _workbook(tmp_path, *expected)
    assert read_therapist_day_history_rows(path) == expected


def test_reader_filters_by_snapshot_and_date(tmp_path: Path):
    current = _row("T1", snapshot_id="snap-v1")
    other = _row("T2", snapshot_id="snap-v2")
    path = _workbook(tmp_path, current, other)
    assert read_therapist_day_history_rows(
        path, snapshot_id="snap-v1", service_date=DAY
    ) == (current,)


def test_reader_rejects_duplicate_snapshot_therapist_key(tmp_path: Path):
    row = _row("T1")
    path = _workbook(tmp_path, row, row)
    with pytest.raises(TherapistDayHistoryStoreError, match="Duplicate"):
        read_therapist_day_history_rows(path)


def test_reader_rejects_schema_drift(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(("WrongHeader",) + HEADERS[1:])
    wb.save(path)
    wb.close()
    with pytest.raises(TherapistDayHistoryStoreError, match="headers"):
        read_therapist_day_history_rows(path)


def test_store_source_does_not_depend_on_com_number_format_assignment():
    source = Path("Python/rehab_excel/therapist_day_history_store.py").read_text(
        encoding="utf-8"
    )
    executable_lines = [
        line for line in source.splitlines() if not line.lstrip().startswith("#")
    ]
    assert not any(".NumberFormat" in line for line in executable_lines)
    assert "ws.Cells(target_row, col).Value2 = value" in source
    assert "COM read-back mismatch" in source
