from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from rehab_core.day_history_events import DayHistoryEvent, DayHistoryEventType
from rehab_excel.day_history_event_store import (
    HEADERS,
    SHEET_NAME,
    DayHistoryEventStoreError,
    read_day_history_events,
)


DAY = date(2026, 10, 26)
STAMP = datetime(2026, 10, 26, 16, 0, tzinfo=timezone.utc)


def _event(event_id: str = "event-1") -> DayHistoryEvent:
    return DayHistoryEvent(
        event_id=event_id,
        service_date=DAY,
        event_type=DayHistoryEventType.DAY_FINALIZED,
        recorded_at=STAMP,
        actor_or_source="close-day-action",
        related_snapshot_id="finalized:2026-10-26:v1",
        summary="Day finalized",
    )


def _workbook(tmp_path: Path, *events: DayHistoryEvent) -> Path:
    path = tmp_path / "events.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for event in events:
        ws.append(
            [
                event.event_id,
                event.service_date.isoformat(),
                event.event_type.value,
                event.recorded_at.isoformat(),
                event.actor_or_source,
                event.therapist_id or "",
                event.patient_id or "",
                event.session_id or "",
                event.reason_code or "",
                event.related_event_id or "",
                event.related_snapshot_id or "",
                event.summary,
            ]
        )
    visible = wb.create_sheet("VISIBLE")
    wb.active = wb.sheetnames.index(visible.title)
    ws.sheet_state = "veryHidden"
    ws.protection.sheet = True
    wb.save(path)
    wb.close()
    return path


def test_reader_round_trips_day_history_event(tmp_path: Path):
    expected = _event()
    path = _workbook(tmp_path, expected)

    assert read_day_history_events(path) == (expected,)


def test_reader_filters_by_date_and_type(tmp_path: Path):
    expected = _event("event-a")
    path = _workbook(tmp_path, expected)

    assert read_day_history_events(
        path,
        service_date=DAY,
        event_type=DayHistoryEventType.DAY_FINALIZED,
    ) == (expected,)


def test_reader_rejects_duplicate_event_id(tmp_path: Path):
    event = _event()
    path = _workbook(tmp_path, event, event)

    with pytest.raises(DayHistoryEventStoreError, match="Duplicate EventId"):
        read_day_history_events(path)


def test_reader_rejects_schema_drift(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(("WrongHeader",) + HEADERS[1:])
    wb.save(path)
    wb.close()

    with pytest.raises(DayHistoryEventStoreError, match="headers"):
        read_day_history_events(path)


def test_store_source_uses_value2_and_semantic_readback_without_number_format():
    source = Path("Python/rehab_excel/day_history_event_store.py").read_text(
        encoding="utf-8"
    )
    executable_lines = [
        line for line in source.splitlines() if not line.lstrip().startswith("#")
    ]

    assert not any(".NumberFormat" in line for line in executable_lines)
    assert "ws.Cells(target_row, col).Value2 = value" in source
    assert "COM read-back mismatch" in source
